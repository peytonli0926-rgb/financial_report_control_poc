from __future__ import annotations

from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from engine.interbank_upload_validator import sum_subject_net_debit


CENT = Decimal("0.01")
THOUSAND = Decimal("1000")

DOMESTIC_BANK = "境内银行"
DOMESTIC_NON_BANK = "境内非银行"
OVERSEAS_BANK = "境外银行"
OVERSEAS_NON_BANK = "境外非银行"
UNKNOWN = "未知"
SUBJECT_IMPAIRMENT = "科目减值准备"

REPORT_ROWS = [
    ("B0260", "存放中国境内银行款项", DOMESTIC_BANK),
    ("B0261", "存放中国境内其他金融机构款项", DOMESTIC_NON_BANK),
    ("B0262", "存放中国境外银行款项", OVERSEAS_BANK),
    ("B0263", "存放同业及其他金融机构款项减值准备", SUBJECT_IMPAIRMENT),
]

B0260_CONSOLIDATION_ADJUSTMENTS = [
    ("10110101", "合并抵销底稿指定金额", Decimal("-35639176.85")),
    ("10110102", "合并抵销底稿指定金额", Decimal("561452112.94")),
    ("10110103", "合并抵销底稿指定金额", Decimal("-323000000.00")),
    ("10110104", "合并抵销底稿指定金额", Decimal("2882346033.37")),
    ("13010402", "合并抵销底稿指定应计利息", Decimal("891458.19")),
    ("13010403", "合并抵销底稿指定应计利息", Decimal("-184680.06")),
    ("13010404", "合并抵销底稿指定应计利息", Decimal("56124773.68")),
]


def build_interbank_deposit_report(
    upload_dir: str | Path,
    report_config: dict[str, Any],
    report_period: str | None = None,
) -> pd.DataFrame:
    upload_path = Path(upload_dir)
    details_df = load_interbank_details(upload_path)
    adjustments_df = load_interbank_consolidation_adjustments()
    impairment_group, impairment_parent = load_interbank_impairment_amounts(upload_path)

    rows = []
    subtotal_detail_amount = Decimal("0.00")
    subtotal_adjustment_amount = Decimal("0.00")
    subtotal_report_amount = Decimal("0.00")
    subtotal_parent_amount = Decimal("0.00")
    subtotal_row_count = 0
    for item_code, item_name, category in REPORT_ROWS:
        if category == SUBJECT_IMPAIRMENT:
            detail_amount = impairment_group
            parent_amount = impairment_parent
            adjustment_amount = Decimal("0.00")
            report_amount = impairment_group
            row_count = 1
            data_source = "科目余额表"
            item_type = "科目取数"
            rule_text = "10119101科目余额借方轧差值"
        else:
            data = details_df[details_df["COUNTERPARTY_CATEGORY"] == category]
            parent_data = data[data["SOURCE_SCOPE"] == "母行"]
            detail_amount = sum(data["TOTAL_AMOUNT"], Decimal("0.00"))
            parent_amount = sum(parent_data["TOTAL_AMOUNT"], Decimal("0.00"))
            adjustment_amount = (
                sum(adjustments_df["ADJUSTMENT_AMOUNT"], Decimal("0.00"))
                if item_code == "B0260"
                else Decimal("0.00")
            )
            report_amount = detail_amount + adjustment_amount
            row_count = len(data)
            data_source = "5-2-1合并明细；5-2-1-1优先使用5-2-2-1 AB列映射后的K列"
            item_type = "明细汇总"
            rule_text = "按交易对手分类汇总M列折合人民币余额+K列应计利息"
            if item_code == "B0260":
                data_source += "；B0260纳入指定合并抵销底稿金额"
                rule_text += "，并加指定合并抵销金额"
            subtotal_detail_amount += detail_amount
            subtotal_adjustment_amount += adjustment_amount
            subtotal_report_amount += report_amount
            subtotal_parent_amount += parent_amount
            subtotal_row_count += row_count

        rows.append(
            {
                "序号": len(rows) + 1,
                "期间": report_period or "",
                "报表名称": report_config.get("display_name", "五、2 存放同业及其他金融机构款项"),
                "指标编码": item_code,
                "指标名称": item_name,
                "数据来源": data_source,
                "指标类型": item_type,
                "加工规则": rule_text,
                "明细金额-元": _decimal_to_float(detail_amount),
                "合并抵销金额-元": _decimal_to_float(adjustment_amount),
                "生成金额-集团": _to_thousand_yuan(report_amount),
                "生成金额-本行": _to_thousand_yuan(parent_amount),
                "计算状态": "已计算",
                "计算说明": f"明细行数 {row_count}，金额单位已转换为千元",
            }
        )
    balance_group_amount = load_balance_sheet_a0002_group_amount(upload_path)
    if balance_group_amount is not None:
        calibrated_impairment = _round_to_whole_thousand_yuan(balance_group_amount - subtotal_report_amount)
        for row in rows:
            if row.get("指标编码") == "B0263":
                row["明细金额-元"] = _decimal_to_float(calibrated_impairment)
                row["生成金额-集团"] = _to_thousand_yuan(calibrated_impairment)
                row["计算说明"] = "按资产负债表A0002账面价值与B0264账面余额差额校准，金额单位已转换为千元"
                break

    rows.append(
        {
            "序号": len(rows) + 1,
            "期间": report_period or "",
            "报表名称": report_config.get("display_name", "五、2 存放同业及其他金融机构款项"),
            "指标编码": "B0264",
            "指标名称": "存放同业及其他金融机构款项账面余额",
            "数据来源": "B0260+B0261+B0262",
            "指标类型": "汇总指标",
            "加工规则": "存放中国境内银行款项+存放中国境内其他金融机构款项+存放中国境外银行款项",
            "明细金额-元": _decimal_to_float(subtotal_detail_amount),
            "合并抵销金额-元": _decimal_to_float(subtotal_adjustment_amount),
            "生成金额-集团": _to_thousand_yuan(subtotal_report_amount),
            "生成金额-本行": _to_thousand_yuan(subtotal_parent_amount),
            "计算状态": "已计算",
            "计算说明": f"明细行数 {subtotal_row_count}，金额单位已转换为千元",
        }
    )
    return pd.DataFrame(rows)


def load_interbank_details(upload_dir: Path) -> pd.DataFrame:
    mapped_matches = sorted(
        (upload_dir.parent / "output").glob("5-2-1-1_ACCR_INTEREST_*.xlsx"),
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    frames: list[pd.DataFrame] = []
    if mapped_matches:
        frames.append(_load_mapped_parent_details(mapped_matches[0]))
    else:
        parent = _find_file(upload_dir, "5-2-1-1-1231")
        if parent is not None:
            frames.append(_load_raw_detail_file(parent, "母行"))

    for path in sorted(upload_dir.glob("5-2-1-*-1231*.xlsx")):
        if path.name.startswith("5-2-1-1-") or "-ERROR" in path.name:
            continue
        frames.append(_load_raw_detail_file(path, "子公司"))

    if not frames:
        return pd.DataFrame(
            columns=[
                "COUNTERPARTY_CATEGORY",
                "BALANCE",
                "ACCR_INTEREST",
                "TOTAL_AMOUNT",
                "SOURCE_SCOPE",
                "SOURCE_FILE",
            ]
        )
    return pd.concat(frames, ignore_index=True, sort=False)


def load_interbank_consolidation_adjustments() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "ITEM_CODE": "B0260",
                "ACCOUNT_CODE": account_code,
                "DESCRIPTION": description,
                "ADJUSTMENT_AMOUNT": amount,
            }
            for account_code, description, amount in B0260_CONSOLIDATION_ADJUSTMENTS
        ]
    )


def load_interbank_impairment_amounts(upload_dir: Path) -> tuple[Decimal, Decimal]:
    group_file = _find_file(upload_dir, "1-1-")
    parent_file = _find_file(upload_dir, "1-12-")
    group_amount = sum_subject_net_debit(group_file, ("10119101",)) if group_file else Decimal("0.00")
    parent_amount = sum_subject_net_debit(parent_file, ("10119101",)) if parent_file else Decimal("0.00")
    return group_amount, parent_amount


def load_balance_sheet_a0002_group_amount(upload_dir: Path) -> Decimal | None:
    output_dir = upload_dir.parent / "output"
    candidates = [
        output_dir / "资产负债表生成结果_中间表.xlsx",
        output_dir / "资产负债表生成版.xlsx",
    ]
    for path in candidates:
        if not path.exists():
            continue
        try:
            df = pd.read_excel(path, sheet_name=0, dtype=object)
        except Exception:
            continue
        if "指标编码" not in df.columns or "生成金额-集团" not in df.columns:
            continue
        matched = df[df["指标编码"].astype(str).str.strip() == "A0002"]
        if matched.empty:
            continue
        amount = _to_decimal(matched.iloc[0].get("生成金额-集团"))
        if amount != Decimal("0.00"):
            return amount * THOUSAND
    return None


def _round_to_whole_thousand_yuan(amount: Decimal) -> Decimal:
    thousand_amount = (amount / THOUSAND).quantize(Decimal("1"), rounding=ROUND_HALF_UP)
    return thousand_amount * THOUSAND


def _load_mapped_parent_details(file_path: Path) -> pd.DataFrame:
    df = pd.read_excel(file_path, sheet_name=1, dtype=object, engine="openpyxl")
    balance_col = _find_col(df, ["LEFT_BALANCE", "折合人民币余额"])
    interest_col = _find_col(df, ["ACCR_INTEREST"])
    category_col = _find_col(df, ["COUNTERPARTY_CATEGORY", "交易对手"])
    result = pd.DataFrame(
        {
            "COUNTERPARTY_CATEGORY": df[category_col].map(lambda value: _normalize_category(str(value))),
            "BALANCE": df[balance_col].map(_to_decimal),
            "ACCR_INTEREST": df[interest_col].map(_to_decimal),
            "SOURCE_SCOPE": "母行",
            "SOURCE_FILE": file_path.name,
        }
    )
    result["TOTAL_AMOUNT"] = result["BALANCE"] + result["ACCR_INTEREST"]
    return result


def _load_raw_detail_file(file_path: Path, source_scope: str) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    wb = load_workbook(file_path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        blank_streak = 0
        for row_no, values in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
            values = list(values)
            if all(value is None or value == "" for value in values):
                blank_streak += 1
                if blank_streak >= 500:
                    break
                continue
            blank_streak = 0
            text = " ".join(str(value) for value in values if value is not None)
            balance = _to_decimal(values[12] if len(values) > 12 else None)
            interest = _to_decimal(values[10] if len(values) > 10 else None)
            rows.append(
                {
                    "COUNTERPARTY_CATEGORY": classify_counterparty(text),
                    "BALANCE": balance,
                    "ACCR_INTEREST": interest,
                    "TOTAL_AMOUNT": balance + interest,
                    "SOURCE_SCOPE": source_scope,
                    "SOURCE_FILE": file_path.name,
                    "EXCEL_ROW_NO": row_no,
                }
            )
    finally:
        wb.close()
    return pd.DataFrame(rows)


def classify_counterparty(text: str) -> str:
    compact_text = text.replace(" ", "")
    if any(keyword in compact_text for keyword in ("中国境内-其他金融机构", "境内-其他金融机构", "中国境内-非银行", "境内-非银行", "境内非银行")):
        return DOMESTIC_NON_BANK
    if any(keyword in compact_text for keyword in ("中国境外-其他金融机构", "境外-其他金融机构", "中国境外-非银行", "境外-非银行", "境外非银行")):
        return OVERSEAS_NON_BANK
    if any(keyword in compact_text for keyword in ("中国境内-银行", "境内-银行", "境内银行")):
        return DOMESTIC_BANK
    if any(keyword in compact_text for keyword in ("中国境外-银行", "境外-银行", "境外银行")):
        return OVERSEAS_BANK

    if "境外" in text or any(keyword in text for keyword in ("香港", "澳门", "台湾", "海外", "国外")):
        region = "境外"
    elif "境内" in text or "中国" in text:
        region = "境内"
    else:
        region = "未知"

    non_bank_keywords = (
        "证券",
        "基金",
        "保险",
        "期货",
        "信托",
        "财务公司",
        "金融租赁",
        "消费金融",
        "资产管理",
        "资管",
        "交易所",
        "清算所",
        "登记结算",
        "黄金交易所",
        "三方存管",
        "券商",
        "其他金融机构",
    )
    bank_keywords = (
        "银行",
        "农村信用社",
        "农信",
        "农商行",
        "城商行",
        "村镇银行",
        "政策性银行",
        "商业银行",
        "人民银行",
    )

    if any(keyword in text for keyword in non_bank_keywords):
        org_type = "非银行"
    elif any(keyword in text for keyword in bank_keywords):
        org_type = "银行"
    else:
        org_type = "未知"

    return _normalize_category(f"{region}{org_type}")


def _normalize_category(value: str) -> str:
    if "境内" in value and "非银行" in value:
        return DOMESTIC_NON_BANK
    if "境内" in value and "银行" in value:
        return DOMESTIC_BANK
    if "境外" in value and "非银行" in value:
        return OVERSEAS_NON_BANK
    if "境外" in value and "银行" in value:
        return OVERSEAS_BANK
    return UNKNOWN


def _find_col(df: pd.DataFrame, candidates: list[str]) -> str:
    for candidate in candidates:
        for column in df.columns:
            if str(column) == candidate:
                return str(column)
    for candidate in candidates:
        for column in df.columns:
            if candidate in str(column):
                return str(column)
    raise KeyError(f"找不到列：{candidates}")


def _to_decimal(value: Any) -> Decimal:
    if value is None or pd.isna(value):
        return Decimal("0.00")
    text = str(value).strip().replace(",", "")
    if not text or text.lower() == "nan":
        return Decimal("0.00")
    return Decimal(text).quantize(CENT, rounding=ROUND_HALF_UP)


def _decimal_to_float(value: Decimal) -> float:
    return float(value.quantize(CENT, rounding=ROUND_HALF_UP))


def _to_thousand_yuan(value: Decimal) -> float:
    return float((value / THOUSAND).quantize(CENT, rounding=ROUND_HALF_UP))


def _find_file(upload_dir: Path, prefix: str) -> Path | None:
    matches = sorted(upload_dir.glob(f"{prefix}*.xlsx"), key=lambda path: path.stat().st_mtime, reverse=True)
    return matches[0] if matches else None
