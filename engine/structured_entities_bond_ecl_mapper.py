from __future__ import annotations

from decimal import Decimal, InvalidOperation
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


OUTPUT_FILE_NAME = "6-3-1_债券投资持仓表_ECL_FINAL整合结果.xlsx"

POSITION_PREFIX = "6-3-1-20251231"
IMPAIRMENT_PREFIX = "5-7-3-1-20251231"
IMPAIRMENT_SHEET_NAME = "20251231"


def run_structured_entities_bond_ecl_mapping(
    upload_dir: str | Path,
    output_file: str | Path,
) -> Path:
    upload_path = Path(upload_dir)
    position_file = _latest_file_by_prefix(upload_path, POSITION_PREFIX)
    impairment_file = _latest_impairment_file(upload_path)
    if position_file is None:
        raise FileNotFoundError("未找到 6-3-1-20251231 债券投资持仓表，请先上传。")
    if impairment_file is None:
        raise FileNotFoundError("未找到包含 ECL_FINAL 的 5-7-3-1-20251231 减值明细表，请先上传。")
    if position_file.suffix.lower() != ".xlsx":
        raise ValueError("ECL_FINAL写回6-3-1持仓表需要目标文件为 .xlsx 格式。")

    position_df, header_row = _load_position_df(position_file)
    impairment_df = _load_impairment_df(impairment_file)
    result_df, unmatched_df, summary_df = _merge_ecl(position_df, impairment_df, position_file, impairment_file)
    _write_ecl_back_to_workbook(position_file, result_df, header_row)

    output_path = Path(output_file)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        result_df.drop(columns=["_excel_row"], errors="ignore").to_excel(writer, sheet_name="整合结果", index=False)
        unmatched_df.drop(columns=["_excel_row"], errors="ignore").to_excel(writer, sheet_name="未匹配明细", index=False)
        summary_df.to_excel(writer, sheet_name="匹配汇总", index=False)
    return output_path


def _latest_file_by_prefix(upload_dir: Path, prefix: str) -> Path | None:
    if not upload_dir.exists():
        return None
    files = sorted(
        [
            path
            for path in upload_dir.iterdir()
            if path.is_file()
            and path.name.startswith(prefix)
            and path.suffix.lower() in {".xlsx", ".xls"}
        ],
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    return files[0] if files else None


def _latest_impairment_file(upload_dir: Path) -> Path | None:
    candidates = sorted(
        [
            path
            for path in upload_dir.iterdir()
            if path.is_file()
            and path.name.startswith(IMPAIRMENT_PREFIX)
            and path.suffix.lower() in {".xlsx", ".xls"}
        ],
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    required = {"BIZ_TYPE", "SEC_ID", "START_DT", "DUBIL_MATR_DT", "CURRENT_BALNC", "ACCR_INTEREST", "ECL_FINAL"}
    for path in candidates:
        try:
            columns = pd.read_excel(path, sheet_name=IMPAIRMENT_SHEET_NAME, nrows=0).columns
        except Exception:
            continue
        if required.issubset(set(columns)):
            return path
    return None


def _load_position_df(path: Path) -> tuple[pd.DataFrame, int]:
    required = {"债券代码", "起息日", "到期日", "本金(元)", "应计利息(元)", "应收利息（元）"}
    header_index = _detect_position_header_index(path, required)
    df = pd.read_excel(path, sheet_name=0, header=header_index)
    missing = required - set(df.columns)
    if missing:
        raise ValueError(f"6-3-1 债券投资持仓表缺少列：{', '.join(sorted(missing))}")
    df = df[df["债券代码"].notna()].copy()
    df["_excel_row"] = df.index.astype(int) + header_index + 2
    return df, header_index + 1


def _detect_position_header_index(path: Path, required_columns: set[str]) -> int:
    preview = pd.read_excel(path, sheet_name=0, header=None, nrows=30)
    for index, row in preview.iterrows():
        values = {str(value).strip() for value in row.tolist() if not pd.isna(value)}
        if required_columns.issubset(values):
            return int(index)
    raise ValueError("未能识别6-3-1债券投资持仓表表头行。")


def _load_impairment_df(path: Path) -> pd.DataFrame:
    required = {"BIZ_TYPE", "SEC_ID", "START_DT", "DUBIL_MATR_DT", "CURRENT_BALNC", "ACCR_INTEREST", "ECL_FINAL"}
    df = pd.read_excel(path, sheet_name=IMPAIRMENT_SHEET_NAME, usecols=list(required | {"BUSI_PK_ID"}))
    missing = required - set(df.columns)
    if missing:
        raise ValueError(f"5-7-3-1 减值明细表缺少列：{', '.join(sorted(missing))}")
    biz_type = df["BIZ_TYPE"].map(_normalize_biz_type)
    return df[biz_type == "05"].copy()


def _merge_ecl(
    position_df: pd.DataFrame,
    impairment_df: pd.DataFrame,
    position_file: Path,
    impairment_file: Path,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    position = position_df.copy()
    impairment = impairment_df.copy()

    position["匹配_SEC_ID"] = position["债券代码"].map(_normalize_code)
    position["匹配_START_DT"] = position["起息日"].map(_normalize_date)
    position["匹配_DUBIL_MATR_DT"] = position["到期日"].map(_normalize_date)
    position["匹配_CURRENT_BALNC"] = position["本金(元)"].map(_normalize_amount)
    position["匹配_ACCR_INTEREST"] = (
        position["应收利息（元）"].map(_to_decimal) + position["应计利息(元)"].map(_to_decimal)
    ).map(_normalize_amount)
    position["匹配键"] = _join_match_key(
        position["匹配_SEC_ID"],
        position["匹配_START_DT"],
        position["匹配_DUBIL_MATR_DT"],
        position["匹配_CURRENT_BALNC"],
        position["匹配_ACCR_INTEREST"],
    )

    impairment["匹配_SEC_ID"] = impairment["SEC_ID"].map(_normalize_code)
    impairment["匹配_START_DT"] = impairment["START_DT"].map(_normalize_date)
    impairment["匹配_DUBIL_MATR_DT"] = impairment["DUBIL_MATR_DT"].map(_normalize_date)
    impairment["匹配_CURRENT_BALNC"] = impairment["CURRENT_BALNC"].map(_normalize_amount)
    impairment["匹配_ACCR_INTEREST"] = impairment["ACCR_INTEREST"].map(_normalize_amount)
    impairment["匹配键"] = _join_match_key(
        impairment["匹配_SEC_ID"],
        impairment["匹配_START_DT"],
        impairment["匹配_DUBIL_MATR_DT"],
        impairment["匹配_CURRENT_BALNC"],
        impairment["匹配_ACCR_INTEREST"],
    )

    grouped = (
        impairment.groupby("匹配键", dropna=False)
        .agg(
            ECL_FINAL=("ECL_FINAL", lambda values: pd.to_numeric(values, errors="coerce").fillna(0.0).sum()),
            减值明细匹配行数=("ECL_FINAL", "size"),
            BUSI_PK_ID=("BUSI_PK_ID", _join_unique_values if "BUSI_PK_ID" in impairment.columns else "size"),
            减值明细_SEC_ID=("SEC_ID", _join_unique_values),
            减值明细_START_DT=("START_DT", _join_unique_values),
            减值明细_DUBIL_MATR_DT=("DUBIL_MATR_DT", _join_unique_values),
            减值明细_CURRENT_BALNC=("CURRENT_BALNC", _join_unique_values),
            减值明细_ACCR_INTEREST=("ACCR_INTEREST", _join_unique_values),
        )
        .reset_index()
    )

    result = position.merge(grouped, on="匹配键", how="left")
    result["ECL_FINAL"] = pd.to_numeric(result["ECL_FINAL"], errors="coerce")
    result["匹配状态"] = result["ECL_FINAL"].notna().map({True: "已匹配", False: "未匹配"})
    result["匹配规则"] = "BIZ_TYPE=05 且 SEC_ID+START_DT+DUBIL_MATR_DT+CURRENT_BALNC+ACCR_INTEREST 精确匹配"
    result["未匹配原因"] = ""
    result.loc[result["匹配状态"] == "未匹配", "未匹配原因"] = "未在5-7-3-1减值明细BIZ_TYPE=05中按指定六字段匹配到记录"
    result["ECL_FINAL"] = result["ECL_FINAL"].fillna(0.0)

    unmatched = result[result["匹配状态"] == "未匹配"].copy()
    summary = pd.DataFrame(
        [
            ("6-3-1持仓表", position_file.name),
            ("5-7-3-1减值明细表", impairment_file.name),
            ("匹配规则", "BIZ_TYPE='05'，且 SEC_ID=债券代码、START_DT=起息日、DUBIL_MATR_DT=到期日、CURRENT_BALNC=本金(元)、ACCR_INTEREST=应收利息（元）+应计利息(元)"),
            ("6-3-1持仓行数", len(position)),
            ("5-7-3-1 BIZ_TYPE=05行数", len(impairment)),
            ("已匹配持仓行数", int((result["匹配状态"] == "已匹配").sum())),
            ("未匹配持仓行数", int((result["匹配状态"] == "未匹配").sum())),
            ("写回ECL_FINAL合计", float(result.loc[result["匹配状态"] == "已匹配", "ECL_FINAL"].sum())),
            ("重复匹配键持仓行数", int(position.duplicated("匹配键", keep=False).sum())),
            ("重复匹配键减值明细行数", int(impairment.duplicated("匹配键", keep=False).sum())),
        ],
        columns=["项目", "内容"],
    )
    return result, unmatched, summary


def _write_ecl_back_to_workbook(path: Path, result_df: pd.DataFrame, header_row: int) -> None:
    workbook = load_workbook(path)
    worksheet = workbook.worksheets[0]
    header_map = {
        _normalize_column(cell.value): cell.column
        for cell in worksheet[header_row]
        if cell.value is not None and str(cell.value).strip()
    }
    ecl_column = _ensure_column(worksheet, header_map, header_row, "ECL_FINAL")
    status_column = _ensure_column(worksheet, header_map, header_row, "ECL匹配状态")
    key_column = _ensure_column(worksheet, header_map, header_row, "ECL匹配键")
    reason_column = _ensure_column(worksheet, header_map, header_row, "ECL未匹配原因")

    for _, row in result_df.iterrows():
        excel_row = int(row["_excel_row"])
        worksheet.cell(row=excel_row, column=ecl_column, value=float(row["ECL_FINAL"]))
        worksheet.cell(row=excel_row, column=status_column, value=str(row["匹配状态"]))
        worksheet.cell(row=excel_row, column=key_column, value=str(row["匹配键"]))
        worksheet.cell(row=excel_row, column=reason_column, value=str(row["未匹配原因"]))
    workbook.save(path)


def _ensure_column(worksheet, header_map: dict[str, int], header_row: int, column_name: str) -> int:
    normalized = _normalize_column(column_name)
    if normalized in header_map:
        return header_map[normalized]
    next_column = worksheet.max_column + 1
    worksheet.cell(row=header_row, column=next_column, value=column_name)
    header_map[normalized] = next_column
    return next_column


def _normalize_biz_type(value) -> str:
    text = str(value or "").strip()
    if text.endswith(".0"):
        text = text[:-2]
    return text.zfill(2)


def _normalize_code(value) -> str:
    if pd.isna(value):
        return ""
    text = str(value).strip().replace(" ", "")
    if text.endswith(".0"):
        text = text[:-2]
    return text


def _normalize_date(value) -> str:
    parsed = pd.to_datetime(value, errors="coerce")
    if not pd.isna(parsed):
        return parsed.strftime("%Y%m%d")
    text = str(value or "").strip()
    if text.endswith(".0"):
        text = text[:-2]
    return text.replace("-", "").replace("/", "")


def _normalize_amount(value) -> str:
    try:
        decimal_value = value if isinstance(value, Decimal) else _to_decimal(value)
        return str(decimal_value.quantize(Decimal("0.01")))
    except (InvalidOperation, ValueError):
        return str(value).replace(",", "").strip()


def _to_decimal(value) -> Decimal:
    if pd.isna(value):
        return Decimal("0")
    text = str(value).replace(",", "").strip()
    if not text or text.lower() == "nan" or text == "-":
        return Decimal("0")
    return Decimal(text)


def _normalize_column(value) -> str:
    return "".join(str(value).split())


def _join_match_key(
    sec_id: pd.Series,
    start_dt: pd.Series,
    dubil_matr_dt: pd.Series,
    current_balnc: pd.Series,
    accr_interest: pd.Series,
) -> pd.Series:
    return (
        sec_id.astype(str)
        + "|"
        + start_dt.astype(str)
        + "|"
        + dubil_matr_dt.astype(str)
        + "|"
        + current_balnc.astype(str)
        + "|"
        + accr_interest.astype(str)
    )


def _join_unique_values(values: pd.Series) -> str:
    unique_values = []
    for value in values:
        if pd.isna(value):
            continue
        text = str(value)
        if text not in unique_values:
            unique_values.append(text)
    return ";".join(unique_values[:20])
