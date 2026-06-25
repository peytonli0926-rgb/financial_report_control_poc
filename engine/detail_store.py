from __future__ import annotations

import argparse
import csv
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

import duckdb
import yaml
from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parents[1]
DEFAULT_CONFIG_PATH = BASE_DIR / "config" / "detail_store.yaml"
DEFAULT_UPLOAD_DIR = BASE_DIR / "data" / "upload"


@dataclass(frozen=True)
class DetailDataset:
    key: str
    file_key: str
    table_name: str
    source_prefix: str
    sheet_name: str | int
    header_row: int | str
    columns: tuple[str, ...]
    description: str = ""


def load_detail_store_config(config_path: str | Path = DEFAULT_CONFIG_PATH) -> dict[str, Any]:
    path = Path(config_path)
    if not path.exists():
        return {"database_path": "data/cache/detail_store/detail_store.duckdb", "batch_size": 5000, "datasets": {}}
    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    store = data.get("detail_store") or {}
    store.setdefault("database_path", "data/cache/detail_store/detail_store.duckdb")
    store.setdefault("batch_size", 5000)
    store.setdefault("datasets", {})
    return store


def load_datasets(config_path: str | Path = DEFAULT_CONFIG_PATH) -> dict[str, DetailDataset]:
    config = load_detail_store_config(config_path)
    datasets: dict[str, DetailDataset] = {}
    for key, raw in (config.get("datasets") or {}).items():
        raw = raw or {}
        file_key = str(raw.get("file_key") or key)
        table_name = _safe_identifier(str(raw.get("table_name") or key))
        datasets[str(key)] = DetailDataset(
            key=str(key),
            file_key=file_key,
            table_name=table_name,
            source_prefix=str(raw.get("source_prefix") or ""),
            sheet_name=raw.get("sheet_name", 0),
            header_row=raw.get("header_row", "auto"),
            columns=tuple(str(item) for item in (raw.get("columns") or []) if str(item).strip()),
            description=str(raw.get("description") or ""),
        )
    return datasets


class DetailStore:
    def __init__(
        self,
        db_path: str | Path | None = None,
        config_path: str | Path = DEFAULT_CONFIG_PATH,
        upload_dir: str | Path = DEFAULT_UPLOAD_DIR,
    ) -> None:
        self.config_path = Path(config_path)
        self.config = load_detail_store_config(self.config_path)
        self.db_path = _resolve_project_path(db_path or self.config["database_path"])
        self.upload_dir = Path(upload_dir)
        self.batch_size = int(self.config.get("batch_size") or 5000)
        self.datasets = load_datasets(self.config_path)
        self.db_path.parent.mkdir(parents=True, exist_ok=True)

    def connect(self) -> duckdb.DuckDBPyConnection:
        connection = duckdb.connect(str(self.db_path))
        connection.execute(
            """
            create table if not exists detail_store_metadata (
                dataset_key varchar primary key,
                table_name varchar not null,
                source_path varchar not null,
                source_mtime double not null,
                source_size bigint not null,
                row_count bigint not null,
                columns_json varchar not null,
                built_at timestamp default current_timestamp
            )
            """
        )
        return connection

    def find_source_file(self, dataset: DetailDataset) -> Path:
        if not dataset.source_prefix:
            raise ValueError(f"数据集 {dataset.key} 缺少 source_prefix。")
        candidates = [
            path
            for path in self.upload_dir.iterdir()
            if path.is_file()
            and path.name.startswith(dataset.source_prefix)
            and path.suffix.lower() in {".xlsx", ".xlsm", ".csv"}
        ]
        if not candidates:
            raise FileNotFoundError(f"未找到 {dataset.source_prefix}*.xlsx 或 *.csv，请先上传对应明细文件。")
        return sorted(candidates, key=lambda item: item.stat().st_mtime, reverse=True)[0]

    def ensure_dataset(self, dataset_key: str, force: bool = False) -> dict[str, Any]:
        dataset = self.datasets.get(dataset_key)
        if dataset is None:
            raise KeyError(f"未配置明细数据集：{dataset_key}")
        source_path = self.find_source_file(dataset)
        with self.connect() as connection:
            if not force and self._is_cache_current(connection, dataset, source_path):
                return self.dataset_status(dataset_key)
            row_count, columns = self._rebuild_dataset(connection, dataset, source_path)
        return {
            "dataset_key": dataset.key,
            "table_name": dataset.table_name,
            "source_path": str(source_path),
            "row_count": row_count,
            "columns": columns,
            "rebuilt": True,
        }

    def ensure_all(self, force: bool = False) -> list[dict[str, Any]]:
        return [self.ensure_dataset(key, force=force) for key in self.datasets]

    def dataset_status(self, dataset_key: str) -> dict[str, Any]:
        dataset = self.datasets.get(dataset_key)
        if dataset is None:
            raise KeyError(f"未配置明细数据集：{dataset_key}")
        with self.connect() as connection:
            rows = connection.execute(
                """
                select dataset_key, table_name, source_path, source_mtime, source_size, row_count, columns_json, built_at
                from detail_store_metadata
                where dataset_key = ?
                """,
                [dataset_key],
            ).fetchall()
        if not rows:
            return {"dataset_key": dataset_key, "table_name": dataset.table_name, "cached": False}
        row = rows[0]
        return {
            "dataset_key": row[0],
            "table_name": row[1],
            "source_path": row[2],
            "source_mtime": row[3],
            "source_size": row[4],
            "row_count": row[5],
            "columns": json.loads(row[6]),
            "built_at": str(row[7]),
            "cached": True,
        }

    def query(self, sql: str, parameters: Iterable[Any] | None = None) -> Any:
        with self.connect() as connection:
            return connection.execute(sql, list(parameters or [])).fetchdf()

    def join_selected_fields(
        self,
        base_table: str,
        joins: list[dict[str, str]],
        select_columns: list[str],
        where: str | None = None,
        limit: int | None = None,
    ) -> Any:
        base_alias = "t0"
        select_clause = ", ".join(select_columns)
        sql_parts = [f"select {select_clause}", f"from {_safe_identifier(base_table)} {base_alias}"]
        for index, join in enumerate(joins, start=1):
            table = _safe_identifier(join["table"])
            alias = f"t{index}"
            left_key = join["left_key"]
            right_key = join["right_key"]
            sql_parts.append(
                f"left join {table} {alias} on cast({join.get('left_alias', base_alias)}.{_quote_ident(left_key)} as varchar) = cast({alias}.{_quote_ident(right_key)} as varchar)"
            )
        if where:
            sql_parts.append(f"where {where}")
        if limit:
            sql_parts.append(f"limit {int(limit)}")
        return self.query("\n".join(sql_parts))

    def _is_cache_current(
        self,
        connection: duckdb.DuckDBPyConnection,
        dataset: DetailDataset,
        source_path: Path,
    ) -> bool:
        stat = source_path.stat()
        rows = connection.execute(
            """
            select source_mtime, source_size
            from detail_store_metadata
            where dataset_key = ?
            """,
            [dataset.key],
        ).fetchall()
        if not rows:
            return False
        source_mtime, source_size = rows[0]
        return float(source_mtime) == float(stat.st_mtime) and int(source_size) == int(stat.st_size)

    def _rebuild_dataset(
        self,
        connection: duckdb.DuckDBPyConnection,
        dataset: DetailDataset,
        source_path: Path,
    ) -> tuple[int, list[str]]:
        header, rows = _iter_source_rows(source_path, dataset)
        selected_columns = _selected_columns(header, dataset.columns)
        if not selected_columns:
            raise ValueError(f"{source_path.name} 未识别到可入库字段。")

        table_name = _safe_identifier(dataset.table_name)
        connection.execute(f"drop table if exists {table_name}")
        column_defs = ", ".join(f"{_quote_ident(column)} varchar" for column in selected_columns)
        connection.execute(f"create table {table_name} ({column_defs})")

        index_by_column = {column: header.index(column) for column in selected_columns}
        batch: list[list[str | None]] = []
        row_count = 0
        insert_sql = f"insert into {table_name} values ({', '.join(['?'] * len(selected_columns))})"
        for row in rows:
            batch.append([_cell_to_text(row[index_by_column[column]]) for column in selected_columns])
            if len(batch) >= self.batch_size:
                connection.executemany(insert_sql, batch)
                row_count += len(batch)
                batch.clear()
        if batch:
            connection.executemany(insert_sql, batch)
            row_count += len(batch)

        stat = source_path.stat()
        connection.execute(
            """
            insert or replace into detail_store_metadata
            (dataset_key, table_name, source_path, source_mtime, source_size, row_count, columns_json, built_at)
            values (?, ?, ?, ?, ?, ?, ?, current_timestamp)
            """,
            [
                dataset.key,
                table_name,
                str(source_path),
                float(stat.st_mtime),
                int(stat.st_size),
                int(row_count),
                json.dumps(selected_columns, ensure_ascii=False),
            ],
        )
        return row_count, selected_columns


def _iter_source_rows(source_path: Path, dataset: DetailDataset) -> tuple[list[str], Iterable[tuple[Any, ...]]]:
    if source_path.suffix.lower() == ".csv":
        return _iter_csv_rows(source_path, dataset)
    return _iter_excel_rows(source_path, dataset)


def _iter_excel_rows(source_path: Path, dataset: DetailDataset) -> tuple[list[str], Iterable[tuple[Any, ...]]]:
    workbook = load_workbook(source_path, read_only=True, data_only=True)
    try:
        worksheet = workbook.worksheets[int(dataset.sheet_name)] if isinstance(dataset.sheet_name, int) else workbook[str(dataset.sheet_name)]
        header_row_index = _detect_header_row(worksheet, dataset.header_row)
        header_values = next(worksheet.iter_rows(min_row=header_row_index, max_row=header_row_index, values_only=True))
        header = _deduplicate_columns([_header_to_text(value, index) for index, value in enumerate(header_values, start=1)])

        def row_iter() -> Iterable[tuple[Any, ...]]:
            for row in worksheet.iter_rows(min_row=header_row_index + 1, values_only=True):
                if any(value is not None and str(value).strip() for value in row):
                    yield row

        return header, row_iter()
    except Exception:
        workbook.close()
        raise


def _iter_csv_rows(source_path: Path, dataset: DetailDataset) -> tuple[list[str], Iterable[tuple[Any, ...]]]:
    encoding = _detect_csv_encoding(source_path)
    header_row_index = _detect_csv_header_row(source_path, dataset.header_row, encoding)

    with source_path.open("r", encoding=encoding, newline="") as handle:
        reader = csv.reader(handle)
        header_values: list[Any] = []
        for row_index, row in enumerate(reader, start=1):
            if row_index == header_row_index:
                header_values = row
                break
    header = _deduplicate_columns([_header_to_text(value, index) for index, value in enumerate(header_values, start=1)])

    def row_iter() -> Iterable[tuple[Any, ...]]:
        with source_path.open("r", encoding=encoding, newline="") as handle:
            reader = csv.reader(handle)
            for row_index, row in enumerate(reader, start=1):
                if row_index <= header_row_index:
                    continue
                if any(value is not None and str(value).strip() for value in row):
                    if len(row) < len(header):
                        row = [*row, *([None] * (len(header) - len(row)))]
                    yield tuple(row)

    return header, row_iter()


def _detect_header_row(worksheet: Any, header_row: int | str) -> int:
    if isinstance(header_row, int):
        return header_row
    for row_index, row in enumerate(worksheet.iter_rows(min_row=1, max_row=30, values_only=True), start=1):
        non_empty = [value for value in row if value is not None and str(value).strip()]
        if len(non_empty) >= 2:
            return row_index
    return 1


def _detect_csv_header_row(source_path: Path, header_row: int | str, encoding: str) -> int:
    if isinstance(header_row, int):
        return header_row
    with source_path.open("r", encoding=encoding, newline="") as handle:
        reader = csv.reader(handle)
        for row_index, row in enumerate(reader, start=1):
            if row_index > 30:
                break
            non_empty = [value for value in row if value is not None and str(value).strip()]
            if len(non_empty) >= 2:
                return row_index
    return 1


def _detect_csv_encoding(source_path: Path) -> str:
    sample = source_path.read_bytes()[:65536]
    for encoding in ("utf-8-sig", "gb18030", "utf-16"):
        try:
            sample.decode(encoding)
            return encoding
        except UnicodeDecodeError:
            continue
    return "utf-8-sig"


def _selected_columns(header: list[str], configured_columns: tuple[str, ...]) -> list[str]:
    if not configured_columns:
        return header
    header_lookup = {_normalize_column_name(column): column for column in header}
    selected = []
    missing = []
    for column in configured_columns:
        matched = header_lookup.get(_normalize_column_name(column))
        if matched:
            selected.append(matched)
        else:
            missing.append(column)
    if missing:
        raise ValueError(f"Excel 缺少配置字段：{', '.join(missing)}")
    return selected


def _resolve_project_path(path: str | Path) -> Path:
    value = Path(path)
    return value if value.is_absolute() else BASE_DIR / value


def _safe_identifier(value: str) -> str:
    identifier = re.sub(r"\W+", "_", str(value).strip(), flags=re.UNICODE).strip("_").lower()
    if not identifier:
        raise ValueError("标识符不能为空。")
    if identifier[0].isdigit():
        identifier = f"t_{identifier}"
    return identifier


def _quote_ident(value: str) -> str:
    return '"' + str(value).replace('"', '""') + '"'


def _normalize_column_name(value: str) -> str:
    return re.sub(r"\s+", "", str(value or "")).lower()


def _header_to_text(value: Any, index: int) -> str:
    text = "" if value is None else str(value).strip()
    return text or f"column_{index}"


def _deduplicate_columns(columns: list[str]) -> list[str]:
    counts: dict[str, int] = {}
    result = []
    for column in columns:
        count = counts.get(column, 0)
        counts[column] = count + 1
        result.append(column if count == 0 else f"{column}_{count + 1}")
    return result


def _cell_to_text(value: Any) -> str | None:
    if value is None:
        return None
    return str(value).strip()


def main() -> None:
    parser = argparse.ArgumentParser(description="Build and query the local DuckDB detail store.")
    parser.add_argument("command", choices=["build", "status", "query"])
    parser.add_argument("--dataset", default="")
    parser.add_argument("--force", action="store_true")
    parser.add_argument("--sql", default="")
    args = parser.parse_args()

    store = DetailStore()
    if args.command == "build":
        result = store.ensure_dataset(args.dataset, force=args.force) if args.dataset else store.ensure_all(force=args.force)
        print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
    elif args.command == "status":
        result = store.dataset_status(args.dataset) if args.dataset else [store.dataset_status(key) for key in store.datasets]
        print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
    elif args.command == "query":
        if not args.sql:
            raise SystemExit("--sql is required for query")
        print(store.query(args.sql).to_string(index=False))


if __name__ == "__main__":
    main()
