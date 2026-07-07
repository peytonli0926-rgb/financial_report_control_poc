from __future__ import annotations

from decimal import Decimal
from pathlib import Path
from typing import Any

import duckdb
import pandas as pd
from openpyxl import load_workbook

from engine.report_generator import build_report_from_rules
from engine.rule_parser import find_report_rules


THOUSAND = Decimal("1000")
RETAIL_LOAN_PREFIX = "5-6-4-20251231-02-零售贷款-0"
AMOUNT_EXPRESSION = "coalesce(try_cast(CURRENT_BALNC as decimal(38, 6)), 0) + coalesce(try_cast(ACCR_INTEREST as decimal(38, 6)), 0)"
REPORT_KEYWORD = "五、6-2 发放贷款和垫款-按担保方式分布情况"

GUARANTEE_ROWS = [
    ("B0308", "信用贷款", ("03",)),
    ("B0309", "保证贷款", ("04",)),
    ("B0310", "抵押贷款", ("01",)),
    ("B0311", "质押贷款", ("02",)),
    ("B0312", "发放贷款和垫款", ("01", "02", "03", "04")),
]

GUARANTEE_LABELS = {
    "01": "抵押贷款",
    "02": "质押贷款",
    "03": "信用贷款",
    "04": "保证贷款",
}


def build_loans_advances_guarantee_6_2_report(
    upload_dir: str | Path,
    report_config: dict[str, Any],
    institution_context: dict[str, Any] | None = None,
    rule_file_path: str | Path | None = None,
) -> pd.DataFrame:
    source_files = _find_retail_loan_csv_files(upload_dir)
    if not source_files:
        raise FileNotFoundError("未找到 5-6-4-20251231-02-零售贷款-01/02/03/04 CSV 文件。")

    rule_rows = _load_rule_rows(rule_file_path, report_config)
    base_report = _build_base_report(rule_rows, upload_dir, report_config, institution_context)
    base_rows_by_code = _rows_by_item_code(base_report)
    bill_discount_thousand = _load_company_personal_6_1_amount_thousand(
        rule_file_path,
        upload_dir,
        institution_context,
        "B0306",
    )
    credit_card_subject_thousand = _credit_card_subject_debit_total_thousand(upload_dir)
    trade_finance_credit_thousand = _subject_amount_thousand(upload_dir, "14011201", "贷方")
    summary = _query_guarantee_amounts_with_duckdb(source_files)
    corporate_summary = _query_corporate_guarantee_amounts(upload_dir)
    corporate_adjustment_total = _corporate_adjustment_total(corporate_summary)
    bill_extra_amount = _query_bill_extra_amount(upload_dir)
    period = str((institution_context or {}).get("period") or "")
    scope = str((institution_context or {}).get("scope") or "集团/本行")
    institution_code = str((institution_context or {}).get("code") or "")
    institution_name = str((institution_context or {}).get("name") or "")
    source_names = "; ".join(path.name for path in source_files)

    rows: list[dict[str, Any]] = []
    guarantee_codes_by_item = {item_code: codes for item_code, _item_name, codes in GUARANTEE_ROWS}
    for index, rule_row in enumerate(rule_rows, start=1):
        item_code = str(rule_row.get("指标编码") or "").strip()
        item_name = str(rule_row.get("指标名称") or "").strip()
        guarantee_codes = guarantee_codes_by_item.get(item_code, ())
        if not guarantee_codes:
            continue
        if item_code == "B0312" and item_code in base_rows_by_code:
            base_row = base_rows_by_code[item_code]
            if _base_row_has_amount(base_row):
                rows.append(_report_row_from_base(index, rule_row, base_row, period, scope, institution_code, institution_name))
            else:
                total_amount_thousand = _generated_amount_sum_thousand(rows)
                rows.append(
                    _report_row_from_amount(
                        index,
                        rule_row,
                        total_amount_thousand,
                        period,
                        scope,
                        institution_code,
                        institution_name,
                        "按B0308-B0311四类担保方式金额合计计算。",
                    )
                )
            continue

        balance_yuan = sum(
            (
                summary.get(code, {}).get("balance", Decimal("0"))
                + corporate_summary.get(code, {}).get("balance", Decimal("0"))
                for code in guarantee_codes
            ),
            Decimal("0"),
        )
        interest_yuan = sum(
            (
                summary.get(code, {}).get("interest", Decimal("0"))
                + corporate_summary.get(code, {}).get("interest", Decimal("0"))
                for code in guarantee_codes
            ),
            Decimal("0"),
        )
        adjustment_yuan = sum((corporate_summary.get(code, {}).get("adjustment", Decimal("0")) for code in guarantee_codes), Decimal("0"))
        if item_code == "B0308":
            balance_yuan += bill_extra_amount.get("balance", Decimal("0"))
            interest_yuan += bill_extra_amount.get("interest", Decimal("0"))
            adjustment_yuan -= corporate_adjustment_total
            balance_yuan += (credit_card_subject_thousand - trade_finance_credit_thousand) * THOUSAND
        if item_code == "B0311":
            balance_yuan += bill_discount_thousand * THOUSAND
        amount_thousand = float((balance_yuan + interest_yuan + adjustment_yuan) / THOUSAND)
        row_count = sum(
            (
                int(summary.get(code, {}).get("row_count", 0))
                + int(corporate_summary.get(code, {}).get("row_count", 0))
                for code in guarantee_codes
            ),
            0,
        )
        if item_code == "B0308":
            row_count += int(bill_extra_amount.get("row_count", 0))
        extra_note = ""
        if item_code == "B0308":
            extra_note = (
                f"追加信用卡相关科目借方轧差值：{float(credit_card_subject_thousand):,.2f}千元；"
                f"扣减14011201科目贷方轧差值：{float(trade_finance_credit_thousand):,.2f}千元；"
                f"扣减5-7-3-2 BIZ_TYPE=01利息调整：{float(corporate_adjustment_total / THOUSAND):,.2f}千元；"
            )
        if item_code == "B0311":
            extra_note = f"追加票据贴现B0306：{float(bill_discount_thousand):,.2f}千元；"
        rows.append(
            {
                "期间": period,
                "机构口径": scope,
                "机构编码": institution_code,
                "机构": institution_name,
                "序号": index,
                "报表名称": str(rule_row.get("报表名称") or report_config.get("display_name") or REPORT_KEYWORD),
                "指标编码": item_code,
                "指标名称": item_name,
                "数据来源": str(rule_row.get("数据来源") or source_names),
                "指标类型": str(rule_row.get("指标类型") or "明细汇总"),
                "加工规则": str(rule_row.get("加工规则") or ""),
                "生成金额-集团": amount_thousand,
                "生成金额-本行": amount_thousand,
                "计算状态": "已计算",
                "计算说明": (
                    f"DuckDB扫描文件：{source_names}；"
                    f"担保方式：{', '.join(GUARANTEE_LABELS.get(code, code) for code in guarantee_codes)}；"
                    f"明细行数：{row_count:,}；"
                    f"{extra_note}"
                    f"本金合计：{float(balance_yuan / THOUSAND):,.2f}千元；"
                    f"应计利息合计：{float(interest_yuan / THOUSAND):,.2f}千元；"
                    f"利息调整合计：{float(adjustment_yuan / THOUSAND):,.2f}千元。"
                ),
            }
        )
    return pd.DataFrame(rows)


def _credit_card_subject_debit_total_thousand(upload_dir: str | Path) -> Decimal:
    subjects = (
        "13015602",
        "13015604",
        "13015605",
        "13015609",
        "14041102",
        "14041103",
        "14041104",
        "14041107",
        "14041108",
        "14041109",
        "14041110",
        "14041112",
        "14041113",
        "14041114",
    )
    return sum((_subject_amount_thousand(upload_dir, subject, "借方") for subject in subjects), Decimal("0"))


def _subject_amount_thousand(upload_dir: str | Path, subject_code: str, direction: str) -> Decimal:
    from engine.rule_calculator import RuleCalculator

    calculator = RuleCalculator(
        upload_dir,
        subject_balance_prefix="1-12-",
        report_period="20251231",
    )
    amount = calculator._calculate_subject_balance_rule(f"{subject_code}科目余额{direction}轧差值")
    return _to_decimal(amount)


def _corporate_adjustment_total(corporate_summary: dict[str, dict[str, Decimal | int]]) -> Decimal:
    return sum(
        (_to_decimal(summary.get("adjustment")) for summary in corporate_summary.values()),
        Decimal("0"),
    )


def _base_row_has_amount(base_row: dict[str, Any]) -> bool:
    for column in ("生成金额-本行", "生成金额-集团"):
        value = base_row.get(column)
        if value is None or pd.isna(value):
            continue
        try:
            Decimal(str(value))
            return True
        except Exception:
            continue
    return False


def _generated_amount_sum_thousand(rows: list[dict[str, Any]]) -> Decimal:
    return sum((_to_decimal(row.get("生成金额-本行")) for row in rows), Decimal("0"))


def _report_row_from_amount(
    index: int,
    rule_row: dict[str, Any],
    amount_thousand: Decimal,
    period: str,
    scope: str,
    institution_code: str,
    institution_name: str,
    note: str,
) -> dict[str, Any]:
    amount = float(amount_thousand)
    return {
        "期间": period,
        "机构口径": scope,
        "机构编码": institution_code,
        "机构": institution_name,
        "序号": index,
        "报表名称": str(rule_row.get("报表名称") or REPORT_KEYWORD),
        "指标编码": str(rule_row.get("指标编码") or ""),
        "指标名称": str(rule_row.get("指标名称") or ""),
        "数据来源": str(rule_row.get("数据来源") or ""),
        "指标类型": str(rule_row.get("指标类型") or "复合指标"),
        "加工规则": str(rule_row.get("加工规则") or ""),
        "生成金额-集团": amount,
        "生成金额-本行": amount,
        "计算状态": "已计算",
        "计算说明": note,
    }


def _load_company_personal_6_1_amount_thousand(
    rule_file_path: str | Path | None,
    upload_dir: str | Path,
    institution_context: dict[str, Any] | None,
    item_code: str,
) -> Decimal:
    if not rule_file_path:
        return Decimal("0")
    rules = find_report_rules(rule_file_path, "五、6-1 发放贷款和垫款-按公司和个人分布情况")
    if rules.empty:
        return Decimal("0")
    context = institution_context or {}
    report_df = build_report_from_rules(
        rules,
        {
            "display_name": "五、6-1 发放贷款和垫款-按公司和个人分布情况",
            "output_prefix": "发放贷款和垫款-按公司和个人分布情况",
            "period_type": "时点",
        },
        upload_dir=str(upload_dir),
        institution_name=context.get("name"),
        institution_code=context.get("code"),
        institution_scope=context.get("scope"),
        report_period=str(context.get("period") or ""),
    )
    if report_df is None or report_df.empty or "指标编码" not in report_df.columns:
        return Decimal("0")
    target = report_df[report_df["指标编码"].astype(str).str.strip() == item_code]
    if target.empty:
        return Decimal("0")
    amount_column = "生成金额-本行" if "生成金额-本行" in target.columns else "生成金额-集团"
    return _to_decimal(target.iloc[0].get(amount_column))


def _build_base_report(
    rule_rows: list[dict[str, Any]],
    upload_dir: str | Path,
    report_config: dict[str, Any],
    institution_context: dict[str, Any] | None,
) -> pd.DataFrame:
    if not rule_rows:
        return pd.DataFrame()
    try:
        context = institution_context or {}
        return build_report_from_rules(
            pd.DataFrame(rule_rows),
            report_config,
            upload_dir=str(upload_dir),
            institution_name=context.get("name"),
            institution_code=context.get("code"),
            institution_scope=context.get("scope"),
            report_period=str(context.get("period") or ""),
        )
    except Exception:
        return pd.DataFrame()


def _rows_by_item_code(report_df: pd.DataFrame) -> dict[str, dict[str, Any]]:
    if report_df is None or report_df.empty or "指标编码" not in report_df.columns:
        return {}
    return {
        str(row.get("指标编码") or ""): {str(column): row.get(column) for column in report_df.columns}
        for _, row in report_df.iterrows()
    }


def _report_row_from_base(
    index: int,
    rule_row: dict[str, Any],
    base_row: dict[str, Any],
    period: str,
    scope: str,
    institution_code: str,
    institution_name: str,
) -> dict[str, Any]:
    return {
        "期间": period,
        "机构口径": scope,
        "机构编码": institution_code,
        "机构": institution_name,
        "序号": index,
        "报表名称": str(rule_row.get("报表名称") or base_row.get("报表名称") or REPORT_KEYWORD),
        "指标编码": str(rule_row.get("指标编码") or base_row.get("指标编码") or ""),
        "指标名称": str(rule_row.get("指标名称") or base_row.get("指标名称") or ""),
        "数据来源": str(rule_row.get("数据来源") or base_row.get("数据来源") or ""),
        "指标类型": str(rule_row.get("指标类型") or base_row.get("指标类型") or ""),
        "加工规则": str(rule_row.get("加工规则") or base_row.get("加工规则") or ""),
        "生成金额-集团": base_row.get("生成金额-集团"),
        "生成金额-本行": base_row.get("生成金额-本行"),
        "计算状态": str(base_row.get("计算状态") or "已计算"),
        "计算说明": str(base_row.get("计算说明") or "按规则文件科目余额规则计算。"),
    }


def _load_rule_rows(rule_file_path: str | Path | None, report_config: dict[str, Any]) -> list[dict[str, Any]]:
    keyword = str(report_config.get("rule_keyword") or report_config.get("display_name") or REPORT_KEYWORD)
    if rule_file_path:
        rules = find_report_rules(rule_file_path, keyword)
        if not rules.empty:
            parsed_rows = [
                _normalize_rule_record({str(column): row.get(column) for column in rules.columns})
                for _, row in rules.iterrows()
                if str(row.get("指标编码") or "").strip() in {item_code for item_code, _name, _codes in GUARANTEE_ROWS}
            ]
            if parsed_rows:
                return parsed_rows

        direct_rows = _load_rule_rows_direct(rule_file_path, keyword)
        if direct_rows:
            return direct_rows

    return [
        {
            "报表名称": str(report_config.get("display_name") or REPORT_KEYWORD),
            "指标编码": item_code,
            "指标名称": item_name,
            "数据来源": "",
            "指标类型": "明细汇总",
            "加工规则": "",
        }
        for item_code, item_name, _codes in GUARANTEE_ROWS
    ]


def _normalize_rule_record(record: dict[str, Any]) -> dict[str, Any]:
    result = dict(record)
    if not _clean_text(result.get("数据来源")):
        result["数据来源"] = _clean_text(result.get("指标数据来源"))
    if not _clean_text(result.get("加工规则")):
        result["加工规则"] = _clean_text(result.get("指标加工规则"))
    return result


def _load_rule_rows_direct(rule_file_path: str | Path, keyword: str) -> list[dict[str, Any]]:
    path = Path(rule_file_path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        target_codes = {item_code for item_code, _name, _codes in GUARANTEE_ROWS}
        for worksheet in workbook.worksheets:
            header_row_index, header = _detect_rule_header(worksheet)
            if header_row_index is None:
                continue
            rows: list[dict[str, Any]] = []
            for values in worksheet.iter_rows(min_row=header_row_index + 1, values_only=True):
                record = {
                    column: values[index] if index < len(values) else None
                    for index, column in enumerate(header)
                    if column
                }
                report_name = _clean_text(record.get("报表名称"))
                item_code = _clean_text(record.get("指标编码"))
                if report_name != keyword or item_code not in target_codes:
                    continue
                rows.append(
                    {
                        "报表名称": report_name,
                        "指标编码": item_code,
                        "指标名称": _clean_text(record.get("指标名称")),
                        "数据来源": _clean_text(record.get("指标数据来源") or record.get("数据来源")),
                        "指标类型": _clean_text(record.get("指标类型")),
                        "加工规则": _clean_text(record.get("指标加工规则") or record.get("加工规则")),
                    }
                )
            if rows:
                order_by_code = {item_code: index for index, (item_code, _name, _codes) in enumerate(GUARANTEE_ROWS)}
                return sorted(rows, key=lambda row: order_by_code.get(str(row.get("指标编码")), 999))
        return []
    finally:
        workbook.close()


def _detect_rule_header(worksheet: Any) -> tuple[int | None, list[str]]:
    for row_index, values in enumerate(worksheet.iter_rows(min_row=1, max_row=20, values_only=True), start=1):
        header = [_clean_text(value) for value in values]
        if "报表名称" in header and "指标编码" in header and "指标名称" in header:
            return row_index, header
    return None, []


def _clean_text(value: Any) -> str:
    if value is None or pd.isna(value):
        return ""
    return " ".join(str(value).replace("\r", " ").replace("\n", " ").split())


def _find_retail_loan_csv_files(upload_dir: str | Path) -> list[Path]:
    root = Path(upload_dir)
    files = [
        path
        for path in root.iterdir()
        if path.is_file()
        and path.name.startswith(RETAIL_LOAN_PREFIX)
        and path.suffix.lower() == ".csv"
        and not path.name.startswith("~$")
    ]
    return sorted(files, key=lambda path: path.name)


def _query_corporate_guarantee_amounts(upload_dir: str | Path) -> dict[str, dict[str, Decimal | int]]:
    source_file = _find_latest_file(upload_dir, "5-7-3-2-20251231")
    if source_file is None:
        return {}
    df = pd.read_excel(source_file, sheet_name="20251231", dtype=object, engine="openpyxl")
    required = {"BIZ_TYPE", "GUARANTEE_TYPE_CD", "CURRENT_BALNC", "ACCR_INTEREST"}
    missing = required - set(df.columns)
    if missing:
        raise ValueError(f"{source_file.name} 缺少列：{', '.join(sorted(missing))}")
    if "利息调整" not in df.columns:
        df["利息调整"] = 0

    data = df[df["BIZ_TYPE"].map(_normalize_code) == "01"].copy()
    data["_guarantee"] = data["GUARANTEE_TYPE_CD"].map(_normalize_code)
    data["_balance"] = data["CURRENT_BALNC"].map(_to_decimal)
    data["_interest"] = data["ACCR_INTEREST"].map(_to_decimal)
    data["_adjustment"] = data["利息调整"].map(_to_decimal)
    result: dict[str, dict[str, Decimal | int]] = {}
    for guarantee_code, group in data.groupby("_guarantee"):
        result[str(guarantee_code)] = {
            "row_count": int(len(group)),
            "balance": sum(group["_balance"], Decimal("0")),
            "interest": sum(group["_interest"], Decimal("0")),
            "adjustment": sum(group["_adjustment"], Decimal("0")),
        }
    return result


def _query_bill_extra_amount(upload_dir: str | Path) -> dict[str, Decimal | int]:
    source_file = _find_latest_file(upload_dir, "5-7-3-1-20251231")
    if source_file is None:
        return {"row_count": 0, "balance": Decimal("0"), "interest": Decimal("0")}
    df = pd.read_excel(source_file, sheet_name="20251231", dtype=object, engine="openpyxl")
    required = {"BIZ_TYPE", "BUSI_TYPE", "CURRENT_BALNC", "ACCR_INTEREST"}
    missing = required - set(df.columns)
    if missing:
        raise ValueError(f"{source_file.name} 缺少列：{', '.join(sorted(missing))}")
    data = df[
        (df["BIZ_TYPE"].map(_normalize_code) == "01")
        & (df["BUSI_TYPE"].map(_normalize_busi_type) == "01004")
    ].copy()
    return {
        "row_count": int(len(data)),
        "balance": sum(data["CURRENT_BALNC"].map(_to_decimal), Decimal("0")),
        "interest": sum(data["ACCR_INTEREST"].map(_to_decimal), Decimal("0")),
    }


def _find_latest_file(upload_dir: str | Path, prefix: str) -> Path | None:
    root = Path(upload_dir)
    candidates = sorted(
        (
            path
            for path in root.iterdir()
            if path.is_file()
            and path.name.startswith(prefix)
            and path.suffix.lower() in {".xlsx", ".xlsm", ".xls"}
            and not path.name.startswith("~$")
        ),
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    return candidates[0] if candidates else None


def _query_guarantee_amounts_with_duckdb(source_files: list[Path]) -> dict[str, dict[str, Decimal | int]]:
    file_list_sql = ", ".join(_sql_string(str(path.resolve()).replace("\\", "/")) for path in source_files)
    sql = f"""
        select
            trim(GUARANTEE_TYPE_CD) as guarantee_type_cd,
            count(*) as row_count,
            sum(coalesce(try_cast(CURRENT_BALNC as decimal(38, 6)), 0)) as balance,
            sum(coalesce(try_cast(ACCR_INTEREST as decimal(38, 6)), 0)) as interest,
            sum({AMOUNT_EXPRESSION}) as amount
        from read_csv_auto([{file_list_sql}], all_varchar=true, union_by_name=true, sample_size=20480)
        where trim(coalesce(GUARANTEE_TYPE_CD, '')) <> ''
          and trim(coalesce(BIZ_TYPE, '')) = '02'
        group by 1
    """
    with duckdb.connect() as connection:
        rows = connection.execute(sql).fetchall()

    result: dict[str, dict[str, Decimal | int]] = {}
    for guarantee_type_cd, row_count, balance, interest, _amount in rows:
        code = str(guarantee_type_cd or "").strip()
        result[code] = {
            "row_count": int(row_count or 0),
            "balance": _to_decimal(balance),
            "interest": _to_decimal(interest),
        }
    return result


def _to_decimal(value: Any) -> Decimal:
    if value is None or pd.isna(value):
        return Decimal("0")
    if isinstance(value, Decimal):
        return value
    text = str(value).replace(",", "").strip()
    if not text or text == "-" or text.lower() == "nan":
        return Decimal("0")
    return Decimal(text)


def _normalize_code(value: Any) -> str:
    text = _clean_text(value)
    if text.endswith(".0"):
        text = text[:-2]
    return text.zfill(2) if text.isdigit() else text


def _normalize_busi_type(value: Any) -> str:
    text = _clean_text(value)
    if text.endswith(".0"):
        text = text[:-2]
    return text.zfill(5) if text.isdigit() else text


def _sql_string(value: str) -> str:
    return "'" + value.replace("'", "''") + "'"
