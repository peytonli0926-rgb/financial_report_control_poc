from __future__ import annotations

import re
from dataclasses import dataclass
from decimal import Decimal, ROUND_HALF_UP, InvalidOperation
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


CENT = Decimal("0.01")
ZERO = Decimal("0.00")


@dataclass
class MapperConfig:
    mapping_mode: str = "strict"
    amount_tolerance: Decimal = CENT
    header_row_1: int = 1
    header_row_2: int = 2
    data_start_row: int = 3
    left_org_col: str = "E"
    left_org_detail_col: str = "F"
    left_interest_col: str = "K"
    left_balance_col: str = "M"
    left_manual_interest_col: str = "AN"
    risk_org_col: str = "F"
    risk_balance_col: str = "AA"
    risk_interest_col: str = "AB"
    zero_anchor_enabled: bool = True
    zero_anchor_target_category: str = "境内非银行"
    zero_anchor_keywords: tuple[str, ...] = ("三方存管", "黄金交易所", "交易所")


def run_interbank_interest_mapping(
    principal_file: str | Path,
    risk_file: str | Path,
    output_file: str | Path,
    detail_files: list[str | Path] | None = None,
    manual_file: str | Path | None = None,
    config: MapperConfig | None = None,
) -> Path:
    cfg = config or MapperConfig()
    principal_path = Path(principal_file)
    risk_path = Path(risk_file)
    output_path = Path(output_file)

    left_df = read_excel_table(principal_path, cfg)
    risk_df = read_excel_table(risk_path, cfg)
    manual_df = read_excel_table(manual_file, cfg) if manual_file else None

    prepared_left = prepare_left_df(left_df, cfg, manual_df)
    prepared_risk = prepare_risk_df(risk_df, cfg)
    mapped_left, mapped_risk, mapping_df = run_mapping(prepared_left, prepared_risk, cfg)
    merged_detail_df = build_merged_detail_df(mapped_left, detail_files or [], cfg)
    unmapped_df = build_unmapped_risk_df(mapped_risk)
    attention_df = build_attention_df(mapped_left, mapped_risk, unmapped_df)
    category_summary_df = build_category_summary_df(mapped_left)
    check_summary_df = build_check_summary_df(mapped_left, mapped_risk, unmapped_df, attention_df)
    rules_df = build_rules_df(cfg)

    write_mapping_workbook(
        output_path,
        check_summary_df,
        mapped_left,
        mapped_risk,
        mapping_df,
        unmapped_df,
        attention_df,
        category_summary_df,
        rules_df,
        merged_detail_df,
    )
    return output_path


def read_excel_table(file_path: str | Path, config: MapperConfig) -> pd.DataFrame:
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"文件不存在：{path}")

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        max_col = ws.max_column
        header_1: list[Any] | None = None
        header_2: list[Any] | None = None
        rows: list[list[Any]] = []
        row_numbers: list[int] = []
        blank_streak = 0
        for row_no, row_values in enumerate(ws.iter_rows(min_row=1, max_col=max_col, values_only=True), start=1):
            values = list(row_values)
            if row_no == config.header_row_1:
                header_1 = values
                continue
            if row_no == config.header_row_2:
                header_2 = values
                continue
            if row_no < config.data_start_row:
                continue
            if _is_blank_row(values):
                blank_streak += 1
                if rows and blank_streak >= 500:
                    break
                continue
            blank_streak = 0
            rows.append(values)
            row_numbers.append(row_no)
        columns = _build_headers(header_1 or [], header_2 or [], max_col)
    finally:
        wb.close()

    df = pd.DataFrame(rows, columns=columns)
    df.insert(0, "EXCEL_ROW_NO", row_numbers)
    df.attrs["source_file"] = str(path)
    return df


def prepare_left_df(left_df: pd.DataFrame, config: MapperConfig, manual_df: pd.DataFrame | None = None) -> pd.DataFrame:
    df = left_df.copy()
    org_col = _find_column(df, ["机构代码", "ORG_ID"], config.left_org_col)
    detail_col = _find_column(df, ["机构明细", "分理处", "机构详情"], config.left_org_detail_col)
    balance_col = _find_column(df, ["折合人民币余额", "本金余额", "余额"], config.left_balance_col)
    interest_col = _find_column(df, ["应计利息"], config.left_interest_col)
    manual_col = _find_column_optional(df, ["人工核对", "KP：风险取值", "风险取值"], config.left_manual_interest_col)

    if manual_df is not None and len(manual_df) == len(df):
        manual_source_col = _find_column_optional(manual_df, ["人工核对", "KP：风险取值", "风险取值"], config.left_manual_interest_col)
        manual_values = manual_df[manual_source_col].tolist() if manual_source_col else [None] * len(df)
    else:
        manual_values = df[manual_col].tolist() if manual_col else [None] * len(df)

    df["LEFT_ROW_NO"] = df["EXCEL_ROW_NO"]
    df["LEFT_ORG_ID_RAW"] = df[org_col].map(normalize_org_id)
    df["LEFT_ORG_DETAIL"] = df[detail_col].fillna("").astype(str)
    df["LEFT_ORG_ID_FINAL"] = [derive_left_org_id(raw, detail) for raw, detail in zip(df["LEFT_ORG_ID_RAW"], df["LEFT_ORG_DETAIL"])]
    df["LEFT_BALANCE"] = df[balance_col].map(to_decimal)
    df["LEFT_BALANCE_KEY"] = df["LEFT_BALANCE"]
    df["LEFT_ORIGINAL_ACCR_INTEREST"] = df[interest_col].map(to_decimal)
    df["MANUAL_ACCR_INTEREST"] = [to_decimal(value) for value in manual_values]
    df["LEFT_CURRENCY"] = df.apply(lambda row: _detect_currency(_row_text(row)), axis=1)
    df["LEFT_TERM_TYPE"] = df.apply(lambda row: _detect_term_type(_row_text(row)), axis=1)
    df["LEFT_PRODUCT_TYPE"] = df.apply(lambda row: _detect_product_type(_row_text(row)), axis=1)
    df["ACCR_INTEREST"] = ZERO
    df["RISK_MAPPED_ACCR_INTEREST"] = ZERO
    df["RETAINED_ORIGINAL_ACCR_INTEREST"] = ZERO
    df["MATCH_STATUS"] = "未匹配"
    df["MATCH_RULE"] = ""
    df["RISK_ROW_NOS"] = ""
    df["ATTENTION_FLAG"] = "否"
    df["ATTENTION_REASON"] = ""
    df["COUNTERPARTY_REGION"] = df.apply(lambda row: classify_counterparty(_row_text(row))[0], axis=1)
    df["COUNTERPARTY_TYPE"] = df.apply(lambda row: classify_counterparty(_row_text(row))[1], axis=1)
    df["COUNTERPARTY_CATEGORY"] = df.apply(lambda row: classify_counterparty(_row_text(row))[2], axis=1)
    df["NORMAL_MATCH_USED"] = False
    return df


def prepare_risk_df(risk_df: pd.DataFrame, config: MapperConfig) -> pd.DataFrame:
    df = risk_df.copy()
    org_col = _find_column(df, ["ORG_ID", "机构代码"], config.risk_org_col)
    balance_col = _find_column(df, ["CURRENT_BALNC", "当前余额", "本金余额"], config.risk_balance_col)
    interest_col = _find_column(df, ["ACCR_INTEREST", "应计利息"], config.risk_interest_col)

    df["RISK_ROW_NO"] = df["EXCEL_ROW_NO"]
    df["RISK_ORG_ID"] = df[org_col].map(normalize_org_id)
    df["RISK_BALANCE"] = df[balance_col].map(to_decimal)
    df["RISK_BALANCE_KEY"] = df["RISK_BALANCE"]
    df["RISK_ACCR_INTEREST"] = df[interest_col].map(to_decimal)
    df["RISK_CURRENCY"] = df.apply(lambda row: _detect_currency(_row_text(row)), axis=1)
    df["RISK_TERM_TYPE"] = df.apply(lambda row: _detect_term_type(_row_text(row)), axis=1)
    df["RISK_PRODUCT_TYPE"] = df.apply(lambda row: _detect_product_type(_row_text(row)), axis=1)
    df["RISK_MAPPED_AMOUNT"] = ZERO
    df["RISK_UNMAPPED_AMOUNT"] = df["RISK_ACCR_INTEREST"]
    df["RISK_MATCH_STATUS"] = "未匹配"
    df["RISK_MATCH_RULE"] = ""
    df["TARGET_LEFT_ROW_NOS"] = ""
    df["ATTENTION_FLAG"] = "否"
    df["ATTENTION_REASON"] = ""
    df["RISK_USED"] = False
    return df


def run_mapping(left_df: pd.DataFrame, risk_df: pd.DataFrame, config: MapperConfig) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    mappings: list[dict[str, Any]] = []
    apply_exact_unique_match(left_df, risk_df, "LEFT_ORG_ID_FINAL", "R1_修正机构码_金额精确匹配", mappings)
    apply_exact_unique_match(left_df, risk_df, "LEFT_ORG_ID_RAW", "R2_原始机构码_金额精确匹配", mappings)
    apply_manual_interest_duplicate_match(left_df, risk_df, mappings)
    apply_duplicate_sequential_match(left_df, risk_df, mappings)
    apply_group_sum_allocation(left_df, risk_df, mappings)
    handle_zero_balance_interest(left_df, risk_df, config, mappings)
    mark_unmatched_risk(risk_df)
    mapping_df = pd.DataFrame(mappings)
    update_left_and_risk_status(left_df, risk_df, mapping_df)
    apply_manual_validation(left_df)
    flag_unknown_category(left_df)
    return left_df, risk_df, mapping_df


def apply_exact_unique_match(
    left_df: pd.DataFrame,
    risk_df: pd.DataFrame,
    left_org_col: str,
    rule_name: str,
    mappings: list[dict[str, Any]],
) -> None:
    risk_candidates = _unmatched_nonzero_risks(risk_df)
    left_available = left_df[~left_df["NORMAL_MATCH_USED"]]
    risk_groups = _group_indexes(risk_candidates, ["RISK_ORG_ID", "RISK_BALANCE_KEY"])
    left_groups = _group_indexes(left_available, [left_org_col, "LEFT_BALANCE_KEY"])
    for key, risk_indexes in risk_groups.items():
        left_indexes = left_groups.get(key, [])
        if len(risk_indexes) == 1 and len(left_indexes) > 1:
            left_indexes = _filter_left_indexes_by_risk_attributes(left_df, risk_df.loc[risk_indexes[0]], left_indexes)
        if len(risk_indexes) == 1 and len(left_indexes) == 1:
            _add_mapping(left_df, risk_df, left_indexes[0], risk_indexes[0], risk_df.at[risk_indexes[0], "RISK_ACCR_INTEREST"], Decimal("1.00"), rule_name, "已匹配", "否", "", mappings)


def apply_manual_interest_duplicate_match(left_df: pd.DataFrame, risk_df: pd.DataFrame, mappings: list[dict[str, Any]]) -> None:
    risk_candidates = _unmatched_nonzero_risks(risk_df)
    left_available = left_df[~left_df["NORMAL_MATCH_USED"]]
    risk_groups = _group_indexes(risk_candidates, ["RISK_ORG_ID", "RISK_BALANCE_KEY"])
    left_groups = _group_indexes(left_available, ["LEFT_ORG_ID_FINAL", "LEFT_BALANCE_KEY"])
    for key, risk_indexes in risk_groups.items():
        left_indexes = left_groups.get(key, [])
        if len(risk_indexes) <= 1 or len(left_indexes) <= 1:
            continue
        remaining_left = set(left_indexes)
        for risk_idx in sorted(risk_indexes, key=lambda idx: risk_df.at[idx, "RISK_ROW_NO"]):
            risk_interest = risk_df.at[risk_idx, "RISK_ACCR_INTEREST"]
            matched_left = [
                left_idx
                for left_idx in sorted(remaining_left, key=lambda idx: left_df.at[idx, "LEFT_ROW_NO"])
                if left_df.at[left_idx, "MANUAL_ACCR_INTEREST"] != ZERO
                and abs(left_df.at[left_idx, "MANUAL_ACCR_INTEREST"] - risk_interest) <= CENT
            ]
            matched_left = _filter_left_indexes_by_risk_attributes(left_df, risk_df.loc[risk_idx], matched_left)
            if len(matched_left) == 1:
                left_idx = matched_left[0]
                _add_mapping(
                    left_df,
                    risk_df,
                    left_idx,
                    risk_idx,
                    risk_interest,
                    Decimal("1.00"),
                    "R5_人工AN辅助重复行匹配",
                    "已匹配",
                    "否",
                    "",
                    mappings,
                )
                remaining_left.remove(left_idx)


def apply_duplicate_sequential_match(left_df: pd.DataFrame, risk_df: pd.DataFrame, mappings: list[dict[str, Any]]) -> None:
    risk_candidates = _unmatched_nonzero_risks(risk_df)
    left_available = left_df[~left_df["NORMAL_MATCH_USED"]]
    risk_groups = _group_indexes(risk_candidates, ["RISK_ORG_ID", "RISK_BALANCE_KEY"])
    left_groups = _group_indexes(left_available, ["LEFT_ORG_ID_FINAL", "LEFT_BALANCE_KEY"])
    for key, risk_indexes in risk_groups.items():
        left_indexes = left_groups.get(key, [])
        if len(risk_indexes) <= 1 or len(left_indexes) <= 1:
            continue
        risk_indexes = sorted(risk_indexes, key=lambda idx: risk_df.at[idx, "RISK_ROW_NO"])
        left_indexes = sorted(left_indexes, key=lambda idx: left_df.at[idx, "LEFT_ROW_NO"])
        pair_count = min(len(risk_indexes), len(left_indexes))
        attention = "是" if len(risk_indexes) != len(left_indexes) else "否"
        reason = "同机构同金额重复行数量不一致，已按原始行号顺序匹配可确定部分" if attention == "是" else ""
        for risk_idx, left_idx in zip(risk_indexes[:pair_count], left_indexes[:pair_count]):
            _add_mapping(left_df, risk_df, left_idx, risk_idx, risk_df.at[risk_idx, "RISK_ACCR_INTEREST"], Decimal("1.00"), "R3_重复金额_顺序一对一分配", "已匹配", attention, reason, mappings)
        for risk_idx in risk_indexes[pair_count:]:
            _flag_risk(risk_df, risk_idx, reason or "同机构同金额重复行无法唯一分配")


def apply_group_sum_allocation(left_df: pd.DataFrame, risk_df: pd.DataFrame, mappings: list[dict[str, Any]]) -> None:
    for risk_idx, risk_row in _unmatched_nonzero_risks(risk_df).sort_values("RISK_ROW_NO").iterrows():
        if risk_row["RISK_BALANCE"] == ZERO:
            continue
        candidates = left_df[
            (~left_df["NORMAL_MATCH_USED"])
            & (left_df["LEFT_ORG_ID_FINAL"] == risk_row["RISK_ORG_ID"])
            & (left_df["LEFT_BALANCE"] != ZERO)
        ].sort_values("LEFT_ROW_NO")
        combo_indexes = _find_unique_sum_combo(candidates, risk_row["RISK_BALANCE"])
        if combo_indexes is None:
            continue
        allocations = allocate_interest_by_balance(risk_row["RISK_ACCR_INTEREST"], [left_df.loc[idx] for idx in combo_indexes])
        for left_idx, amount, ratio in allocations:
            _add_mapping(left_df, risk_df, left_idx, risk_idx, amount, ratio, "R4_同机构本金汇总分摊", "已匹配", "否", "", mappings)


def handle_zero_balance_interest(
    left_df: pd.DataFrame,
    risk_df: pd.DataFrame,
    config: MapperConfig,
    mappings: list[dict[str, Any]],
) -> None:
    zero_risks = _unmatched_nonzero_risks(risk_df)
    zero_risks = zero_risks[zero_risks["RISK_BALANCE"] == ZERO].sort_values("RISK_ROW_NO")
    for risk_idx, risk_row in zero_risks.iterrows():
        reason = "风险表AA为0但AB不为0，无法通过本金匹配"
        if config.mapping_mode != "force_full":
            _flag_risk(risk_df, risk_idx, reason)
            continue
        left_idx = _find_zero_anchor_left(left_df, risk_row, config)
        if left_idx is None:
            _flag_risk(risk_df, risk_idx, reason + "，未找到强制挂载候选行")
            continue
        _add_mapping(
            left_df,
            risk_df,
            left_idx,
            risk_idx,
            risk_row["RISK_ACCR_INTEREST"],
            Decimal("0.00"),
            "R6_零本金利息_强制挂载",
            "已映射_需复核",
            "是",
            reason + "，已按配置强制挂载",
            mappings,
            normal_used=False,
        )


def build_merged_detail_df(mapped_left: pd.DataFrame, detail_files: list[str | Path], config: MapperConfig) -> pd.DataFrame:
    frames = [_business_left_output(mapped_left, source_label="5-2-1-1")]
    for path_like in detail_files:
        path = Path(path_like)
        if not path.exists():
            continue
        df = read_excel_table(path, config)
        df["ACCR_INTEREST"] = df[_find_column(df, ["应计利息"], config.left_interest_col)].map(to_decimal)
        df["SOURCE_FILE"] = path.name
        df["COUNTERPARTY_CATEGORY"] = df.apply(lambda row: classify_counterparty(_row_text(row))[2], axis=1)
        frames.append(df)
    return pd.concat(frames, ignore_index=True, sort=False) if frames else pd.DataFrame()


def write_mapping_workbook(
    output_path: Path,
    check_summary_df: pd.DataFrame,
    mapped_left: pd.DataFrame,
    mapped_risk: pd.DataFrame,
    mapping_df: pd.DataFrame,
    unmapped_df: pd.DataFrame,
    attention_df: pd.DataFrame,
    category_summary_df: pd.DataFrame,
    rules_df: pd.DataFrame,
    merged_detail_df: pd.DataFrame,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        _to_excel(check_summary_df, writer, "核对汇总")
        _to_excel(_left_result_output(mapped_left), writer, "5-2-1-1_映射结果")
        _to_excel(_risk_result_output(mapped_risk), writer, "5-2-2-1_风险映射明细")
        _to_excel(mapping_df, writer, "风险映射关系")
        _to_excel(unmapped_df, writer, "未映射风险明细")
        _to_excel(attention_df, writer, "需人工关注")
        _to_excel(category_summary_df, writer, "分类汇总")
        _to_excel(rules_df, writer, "规则说明")
        _to_excel(merged_detail_df, writer, "5-2-1_合并明细")
        for ws in writer.book.worksheets:
            _format_sheet(ws)


def normalize_org_id(value: Any) -> str | None:
    if value is None or pd.isna(value):
        return None
    text = str(value).strip()
    if not text or text.lower() == "nan":
        return None
    if re.fullmatch(r"\d+\.0+", text):
        text = text.split(".", 1)[0]
    text = re.sub(r"\s+", "", text)
    if text.isdigit() and len(text) < 6:
        text = text.zfill(6)
    return text


def derive_left_org_id(f_value: Any, g_value: Any) -> str | None:
    text = "" if g_value is None or pd.isna(g_value) else str(g_value).strip()
    match = re.match(r"^(\d{6})", text)
    if match:
        return match.group(1)
    return normalize_org_id(f_value)


def to_decimal(value: Any) -> Decimal:
    if value is None or pd.isna(value):
        return ZERO
    if isinstance(value, Decimal):
        return value.quantize(CENT, rounding=ROUND_HALF_UP)
    text = str(value).strip().replace(",", "")
    if not text or text.lower() == "nan":
        return ZERO
    if text.startswith("(") and text.endswith(")"):
        text = "-" + text[1:-1]
    try:
        return Decimal(text).quantize(CENT, rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError):
        return ZERO


def classify_counterparty(row_text: str) -> tuple[str, str, str]:
    compact_text = row_text.replace(" ", "")
    if any(keyword in compact_text for keyword in ("中国境内-其他金融机构", "境内-其他金融机构", "中国境内-非银行", "境内-非银行", "境内非银行")):
        return "境内", "非银行", "境内非银行"
    if any(keyword in compact_text for keyword in ("中国境外-其他金融机构", "境外-其他金融机构", "中国境外-非银行", "境外-非银行", "境外非银行")):
        return "境外", "非银行", "境外非银行"
    if any(keyword in compact_text for keyword in ("中国境内-银行", "境内-银行", "境内银行")):
        return "境内", "银行", "境内银行"
    if any(keyword in compact_text for keyword in ("中国境外-银行", "境外-银行", "境外银行")):
        return "境外", "银行", "境外银行"

    domestic_keywords = ("中国境内", "境内", "中华人民共和国", "CN", "CHN")
    overseas_keywords = ("境外", "香港", "澳门", "台湾", "海外", "国外", "境外机构")
    bank_keywords = ("银行", "农村信用社", "农信", "农商行", "城商行", "村镇银行", "政策性银行", "商业银行", "人民银行")
    non_bank_keywords = ("证券", "基金", "保险", "期货", "信托", "财务公司", "金融租赁", "消费金融", "资产管理", "资管", "交易所", "清算所", "登记结算", "黄金交易所", "三方存管", "券商", "其他金融机构")
    region = "未知"
    if any(keyword in row_text for keyword in overseas_keywords):
        region = "境外"
    elif any(keyword in row_text for keyword in domestic_keywords):
        region = "境内"
    org_type = "未知"
    if any(keyword in row_text for keyword in non_bank_keywords):
        org_type = "非银行"
    elif any(keyword in row_text for keyword in bank_keywords):
        org_type = "银行"
    if region == "境内" and org_type == "银行":
        return region, org_type, "境内银行"
    if region == "境内" and org_type == "非银行":
        return region, org_type, "境内非银行"
    if region == "境外" and org_type == "银行":
        return region, org_type, "境外银行"
    if region == "境外" and org_type == "非银行":
        return region, org_type, "境外非银行"
    return region, org_type, "未知"


def _filter_left_indexes_by_risk_attributes(left_df: pd.DataFrame, risk_row: pd.Series, left_indexes: list[int]) -> list[int]:
    filtered = list(left_indexes)
    for left_col, risk_col in [
        ("LEFT_CURRENCY", "RISK_CURRENCY"),
        ("LEFT_TERM_TYPE", "RISK_TERM_TYPE"),
        ("LEFT_PRODUCT_TYPE", "RISK_PRODUCT_TYPE"),
    ]:
        risk_value = str(risk_row.get(risk_col) or "")
        if not risk_value:
            continue
        candidates = [idx for idx in filtered if str(left_df.at[idx, left_col] or "") == risk_value]
        if candidates:
            filtered = candidates
    return filtered


def _detect_currency(row_text: str) -> str:
    text = row_text.upper()
    if "美元" in row_text or "USD" in text:
        return "USD"
    if "港币" in row_text or "港元" in row_text or "HKD" in text:
        return "HKD"
    if "欧元" in row_text or "EUR" in text:
        return "EUR"
    if "日元" in row_text or "JPY" in text:
        return "JPY"
    if "人民币" in row_text or "CNY" in text or "RMB" in text:
        return "CNY"
    return ""


def _detect_term_type(row_text: str) -> str:
    if "定期" in row_text:
        return "定期"
    if "活期" in row_text:
        return "活期"
    if "保证金" in row_text:
        return "保证金"
    return ""


def _detect_product_type(row_text: str) -> str:
    if "保证金" in row_text:
        return "保证金"
    if "备付金" in row_text:
        return "备付金"
    if "三方存管" in row_text:
        return "三方存管"
    if "交易所" in row_text:
        return "交易所"
    if "境外" in row_text:
        return "境外存放"
    if "存放境内商业银行" in row_text:
        return "境内商业银行"
    if "存放境内其他银行" in row_text or "其他银行业" in row_text:
        return "境内其他银行业"
    if "其他金融机构" in row_text:
        return "其他金融机构"
    return ""


def allocate_interest_by_balance(risk_interest: Decimal, candidate_left_rows: list[pd.Series]) -> list[tuple[int, Decimal, Decimal]]:
    total_balance = sum((row["LEFT_BALANCE"] for row in candidate_left_rows), ZERO)
    if total_balance == ZERO:
        return []
    allocated: list[tuple[int, Decimal, Decimal]] = []
    running_total = ZERO
    for index, row in enumerate(candidate_left_rows):
        ratio = (row["LEFT_BALANCE"] / total_balance).quantize(Decimal("0.00000001"), rounding=ROUND_HALF_UP)
        if index < len(candidate_left_rows) - 1:
            amount = (risk_interest * row["LEFT_BALANCE"] / total_balance).quantize(CENT, rounding=ROUND_HALF_UP)
            running_total += amount
        else:
            amount = risk_interest - running_total
        allocated.append((row.name, amount, ratio))
    return allocated


def _add_mapping(
    left_df: pd.DataFrame,
    risk_df: pd.DataFrame,
    left_idx: int,
    risk_idx: int,
    allocated_interest: Decimal,
    ratio: Decimal,
    rule: str,
    status: str,
    attention_flag: str,
    attention_reason: str,
    mappings: list[dict[str, Any]],
    normal_used: bool = True,
) -> None:
    mappings.append(
        {
            "RISK_ROW_NO": risk_df.at[risk_idx, "RISK_ROW_NO"],
            "LEFT_ROW_NO": left_df.at[left_idx, "LEFT_ROW_NO"],
            "RISK_ORG_ID": risk_df.at[risk_idx, "RISK_ORG_ID"],
            "LEFT_ORG_ID_FINAL": left_df.at[left_idx, "LEFT_ORG_ID_FINAL"],
            "RISK_BALANCE": risk_df.at[risk_idx, "RISK_BALANCE"],
            "LEFT_BALANCE": left_df.at[left_idx, "LEFT_BALANCE"],
            "RISK_ACCR_INTEREST": risk_df.at[risk_idx, "RISK_ACCR_INTEREST"],
            "ALLOCATED_INTEREST": allocated_interest,
            "ALLOCATION_RATIO": ratio,
            "MATCH_RULE": rule,
            "MATCH_STATUS": status,
            "ATTENTION_FLAG": attention_flag,
            "ATTENTION_REASON": attention_reason,
        }
    )
    risk_df.at[risk_idx, "RISK_USED"] = True
    if normal_used:
        left_df.at[left_idx, "NORMAL_MATCH_USED"] = True


def update_left_and_risk_status(left_df: pd.DataFrame, risk_df: pd.DataFrame, mapping_df: pd.DataFrame) -> None:
    if not mapping_df.empty:
        for left_row_no, group in mapping_df.groupby("LEFT_ROW_NO", sort=False):
            idx = left_df.index[left_df["LEFT_ROW_NO"] == left_row_no][0]
            mapped_amount = sum(group["ALLOCATED_INTEREST"], ZERO)
            left_df.at[idx, "RISK_MAPPED_ACCR_INTEREST"] = mapped_amount
            left_df.at[idx, "ACCR_INTEREST"] = mapped_amount
            left_df.at[idx, "MATCH_STATUS"] = "已映射_需复核" if (group["ATTENTION_FLAG"] == "是").any() else "已匹配"
            left_df.at[idx, "MATCH_RULE"] = "；".join(dict.fromkeys(group["MATCH_RULE"].astype(str)))
            left_df.at[idx, "RISK_ROW_NOS"] = ",".join(group["RISK_ROW_NO"].astype(str))
            if (group["ATTENTION_FLAG"] == "是").any():
                left_df.at[idx, "ATTENTION_FLAG"] = "是"
                left_df.at[idx, "ATTENTION_REASON"] = "；".join(reason for reason in group["ATTENTION_REASON"].astype(str) if reason)
        for risk_row_no, group in mapping_df.groupby("RISK_ROW_NO", sort=False):
            idx = risk_df.index[risk_df["RISK_ROW_NO"] == risk_row_no][0]
            mapped_amount = sum(group["ALLOCATED_INTEREST"], ZERO)
            risk_df.at[idx, "RISK_MAPPED_AMOUNT"] = mapped_amount
            risk_df.at[idx, "RISK_UNMAPPED_AMOUNT"] = risk_df.at[idx, "RISK_ACCR_INTEREST"] - mapped_amount
            risk_df.at[idx, "RISK_MATCH_STATUS"] = "已映射_需复核" if (group["ATTENTION_FLAG"] == "是").any() else "已匹配"
            risk_df.at[idx, "RISK_MATCH_RULE"] = "；".join(dict.fromkeys(group["MATCH_RULE"].astype(str)))
            risk_df.at[idx, "TARGET_LEFT_ROW_NOS"] = ",".join(group["LEFT_ROW_NO"].astype(str))
            if (group["ATTENTION_FLAG"] == "是").any():
                risk_df.at[idx, "ATTENTION_FLAG"] = "是"
                risk_df.at[idx, "ATTENTION_REASON"] = "；".join(reason for reason in group["ATTENTION_REASON"].astype(str) if reason)
    retain_unmapped_left_original_interest(left_df)


def retain_unmapped_left_original_interest(left_df: pd.DataFrame) -> None:
    mask = (
        (left_df["RISK_MAPPED_ACCR_INTEREST"] == ZERO)
        & (left_df["LEFT_ORIGINAL_ACCR_INTEREST"] != ZERO)
    )
    left_df.loc[mask, "RETAINED_ORIGINAL_ACCR_INTEREST"] = left_df.loc[mask, "LEFT_ORIGINAL_ACCR_INTEREST"]
    left_df.loc[mask, "ACCR_INTEREST"] = left_df.loc[mask, "LEFT_ORIGINAL_ACCR_INTEREST"]
    left_df.loc[mask, "MATCH_STATUS"] = "主表原值保留_需复核"
    left_df.loc[mask, "MATCH_RULE"] = "R8_风险表无对应记录_保留主表K列原值"
    left_df.loc[mask, "ATTENTION_FLAG"] = "是"
    left_df.loc[mask, "ATTENTION_REASON"] = left_df.loc[mask, "ATTENTION_REASON"].apply(
        lambda value: _append_reason(value, "5-2-2-1未找到对应风险明细，保留5-2-1-1原K列应计利息")
    )


def apply_manual_validation(left_df: pd.DataFrame) -> None:
    diffs: list[Decimal] = []
    statuses: list[str] = []
    for _, row in left_df.iterrows():
        auto = row["ACCR_INTEREST"]
        manual = row["MANUAL_ACCR_INTEREST"]
        diff = auto - manual
        diffs.append(diff)
        if manual == ZERO and auto == ZERO:
            statuses.append("人工为空")
        elif manual == ZERO and auto != ZERO:
            statuses.append("人工为空但自动有值")
        elif manual != ZERO and auto == ZERO:
            statuses.append("人工有值但自动未匹配")
        elif abs(diff) <= CENT:
            statuses.append("与人工一致")
        else:
            statuses.append("与人工不一致")
    left_df["AUTO_MANUAL_DIFF"] = diffs
    left_df["MANUAL_CHECK_STATUS"] = statuses
    mask = left_df["MANUAL_CHECK_STATUS"].isin(["与人工不一致", "人工有值但自动未匹配"])
    left_df.loc[mask, "ATTENTION_FLAG"] = "是"
    left_df.loc[mask, "ATTENTION_REASON"] = left_df.loc[mask, "ATTENTION_REASON"].apply(
        lambda value: _append_reason(value, "自动映射结果与人工AN列差异超过0.01")
    )


def flag_unknown_category(left_df: pd.DataFrame) -> None:
    mask = (left_df["COUNTERPARTY_CATEGORY"] == "未知") & (left_df["ACCR_INTEREST"] != ZERO)
    left_df.loc[mask, "ATTENTION_FLAG"] = "是"
    left_df.loc[mask, "ATTENTION_REASON"] = left_df.loc[mask, "ATTENTION_REASON"].apply(
        lambda value: _append_reason(value, "无法识别交易对手分类")
    )


def mark_unmatched_risk(risk_df: pd.DataFrame) -> None:
    for idx, row in risk_df.iterrows():
        if row["RISK_ACCR_INTEREST"] != ZERO and not row["RISK_USED"]:
            _flag_risk(risk_df, idx, "找不到同机构同金额本金明细")


def build_unmapped_risk_df(risk_df: pd.DataFrame) -> pd.DataFrame:
    rows = risk_df[(risk_df["RISK_ACCR_INTEREST"] != ZERO) & (risk_df["RISK_UNMAPPED_AMOUNT"] != ZERO)].copy()
    if rows.empty:
        return pd.DataFrame(columns=["RISK_ROW_NO", "RISK_ORG_ID", "RISK_BALANCE", "RISK_ACCR_INTEREST", "UNMATCH_REASON", "SUGGESTED_ACTION"])
    return pd.DataFrame(
        {
            "RISK_ROW_NO": rows["RISK_ROW_NO"],
            "RISK_ORG_ID": rows["RISK_ORG_ID"],
            "RISK_BALANCE": rows["RISK_BALANCE"],
            "RISK_ACCR_INTEREST": rows["RISK_UNMAPPED_AMOUNT"],
            "UNMATCH_REASON": rows["ATTENTION_REASON"].replace("", "未映射"),
            "SUGGESTED_ACTION": "请确认应挂载至哪一条5-2-1-1明细，或补充兜底挂载规则",
        }
    )


def build_attention_df(left_df: pd.DataFrame, risk_df: pd.DataFrame, unmapped_df: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for _, row in left_df[left_df["ATTENTION_FLAG"] == "是"].iterrows():
        rows.append(
            {
                "ISSUE_TYPE": "主表复核",
                "LEFT_ROW_NO": row["LEFT_ROW_NO"],
                "RISK_ROW_NO": row["RISK_ROW_NOS"],
                "ORG_ID": row["LEFT_ORG_ID_FINAL"],
                "BALANCE": row["LEFT_BALANCE"],
                "ACCR_INTEREST": row["ACCR_INTEREST"],
                "ATTENTION_REASON": row["ATTENTION_REASON"],
                "SUGGESTED_ACTION": "请复核分类、人工差异或强制挂载规则",
            }
        )
    for _, row in risk_df[risk_df["ATTENTION_FLAG"] == "是"].iterrows():
        rows.append(
            {
                "ISSUE_TYPE": "风险表复核",
                "LEFT_ROW_NO": row["TARGET_LEFT_ROW_NOS"],
                "RISK_ROW_NO": row["RISK_ROW_NO"],
                "ORG_ID": row["RISK_ORG_ID"],
                "BALANCE": row["RISK_BALANCE"],
                "ACCR_INTEREST": row["RISK_ACCR_INTEREST"],
                "ATTENTION_REASON": row["ATTENTION_REASON"],
                "SUGGESTED_ACTION": "请确认风险表利息应落到哪条主表明细",
            }
        )
    for _, row in unmapped_df.iterrows():
        rows.append(
            {
                "ISSUE_TYPE": "未映射风险",
                "LEFT_ROW_NO": "",
                "RISK_ROW_NO": row["RISK_ROW_NO"],
                "ORG_ID": row["RISK_ORG_ID"],
                "BALANCE": row["RISK_BALANCE"],
                "ACCR_INTEREST": row["RISK_ACCR_INTEREST"],
                "ATTENTION_REASON": row["UNMATCH_REASON"],
                "SUGGESTED_ACTION": row["SUGGESTED_ACTION"],
            }
        )
    return pd.DataFrame(rows)


def build_category_summary_df(left_df: pd.DataFrame) -> pd.DataFrame:
    if left_df.empty:
        return pd.DataFrame()
    grouped = left_df.groupby(["COUNTERPARTY_CATEGORY", "LEFT_ORG_ID_FINAL"], dropna=False)
    rows = []
    for (category, org_id), group in grouped:
        rows.append(
            {
                "COUNTERPARTY_CATEGORY": category,
                "LEFT_ORG_ID_FINAL": org_id,
                "本金余额合计": sum(group["LEFT_BALANCE"], ZERO),
                "ACCR_INTEREST合计": sum(group["ACCR_INTEREST"], ZERO),
                "明细行数": len(group),
                "风险表来源行数": len({item for text in group["RISK_ROW_NOS"].astype(str) for item in text.split(",") if item}),
                "需人工关注金额": sum(group.loc[group["ATTENTION_FLAG"] == "是", "ACCR_INTEREST"], ZERO),
            }
        )
    return pd.DataFrame(rows)


def build_check_summary_df(
    left_df: pd.DataFrame,
    risk_df: pd.DataFrame,
    unmapped_df: pd.DataFrame,
    attention_df: pd.DataFrame,
) -> pd.DataFrame:
    risk_total = sum(risk_df["RISK_ACCR_INTEREST"], ZERO)
    mapped_total = sum(left_df["ACCR_INTEREST"], ZERO)
    risk_mapped_total = sum(left_df["RISK_MAPPED_ACCR_INTEREST"], ZERO)
    retained_total = sum(left_df["RETAINED_ORIGINAL_ACCR_INTEREST"], ZERO)
    unmapped_total = sum(unmapped_df["RISK_ACCR_INTEREST"], ZERO) if not unmapped_df.empty else ZERO
    manual_total = sum(left_df["MANUAL_ACCR_INTEREST"], ZERO)
    rows = [
        ("5-2-2-1 AB列ACCR_INTEREST合计", risk_total),
        ("风险表AB自动映射到5-2-1-1金额合计", risk_mapped_total),
        ("风险表无对应记录但保留主表K列原值合计", retained_total),
        ("最终5-2-1-1 ACCR_INTEREST/K列合计", mapped_total),
        ("未映射ACCR_INTEREST合计", unmapped_total),
        ("风险表总额校验差异", risk_total - risk_mapped_total - unmapped_total),
        ("已映射风险表行数", int((risk_df["RISK_MAPPED_AMOUNT"] != ZERO).sum())),
        ("未映射风险表行数", int(((risk_df["RISK_ACCR_INTEREST"] != ZERO) & (risk_df["RISK_UNMAPPED_AMOUNT"] != ZERO)).sum())),
        ("需人工关注行数", len(attention_df)),
        ("人工AN列合计", manual_total),
        ("自动结果与人工AN列差异", mapped_total - manual_total),
    ]
    for category in ["境内银行", "境内非银行", "境外银行", "境外非银行", "未知"]:
        rows.append((f"{category} ACCR_INTEREST合计", sum(left_df.loc[left_df["COUNTERPARTY_CATEGORY"] == category, "ACCR_INTEREST"], ZERO)))
    return pd.DataFrame(rows, columns=["项目", "内容"])


def build_rules_df(config: MapperConfig) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {"规则": "R1_修正机构码_金额精确匹配", "说明": "主表机构明细前6位机构码加本金余额，匹配风险表机构码加当前余额。"},
            {"规则": "R2_原始机构码_金额精确匹配", "说明": "主表原机构码加本金余额，匹配风险表机构码加当前余额。"},
            {"规则": "R3_重复金额_顺序一对一分配", "说明": "同机构同金额重复时，按Excel原始行号顺序一对一分配。"},
            {"规则": "R4_同机构本金汇总分摊", "说明": "同机构多行本金合计等于风险余额时，按本金比例分摊AB列利息。"},
            {"规则": "R6_零本金利息_强制挂载", "说明": f"当前模式为 {config.mapping_mode}；force_full 模式下按关键字或分类强制挂载并标记需复核。"},
            {"规则": "R8_风险表无对应记录_保留主表K列原值", "说明": "主表原K列有利息但风险表没有对应记录时，保留主表K列金额以保证加工后明细能与科目账核对，并标记需复核。"},
            {"规则": "数据源约束", "说明": "优先使用5-2-2-1风险表AB列；仅当风险表无对应记录且主表原K列非零时，保留主表K列原值并单独披露。"},
        ]
    )


def _build_headers(header_1: list[Any], header_2: list[Any], max_col: int) -> list[str]:
    headers: list[str] = []
    used: dict[str, int] = {}
    first_row = list(header_1) + [None] * max(0, max_col - len(header_1))
    second_row = list(header_2) + [None] * max(0, max_col - len(header_2))
    for idx in range(max_col):
        first = _clean_header(first_row[idx])
        second = _clean_header(second_row[idx])
        name = second or first or f"COL_{get_column_letter(idx + 1)}"
        if first and second and first != second and not isinstance(first_row[idx], pd.Timestamp):
            name = second
        count = used.get(name, 0)
        used[name] = count + 1
        if count:
            name = f"{name}_{count + 1}"
        headers.append(name)
    return headers


def _clean_header(value: Any) -> str:
    if value is None or pd.isna(value):
        return ""
    return str(value).strip()


def _find_column(df: pd.DataFrame, keywords: list[str], fallback_letter: str) -> str:
    for keyword in keywords:
        for col in df.columns:
            if keyword and keyword.upper() in str(col).upper():
                return str(col)
    fallback_col_name = f"COL_{fallback_letter.upper()}"
    if fallback_col_name in df.columns:
        return fallback_col_name
    idx = _col_letter_to_index(fallback_letter)
    if "EXCEL_ROW_NO" in df.columns:
        idx += 1
    if idx < len(df.columns):
        return str(df.columns[idx])
    raise KeyError(f"找不到列：keywords={keywords}, fallback={fallback_letter}")


def _find_column_optional(df: pd.DataFrame, keywords: list[str], fallback_letter: str) -> str | None:
    try:
        return _find_column(df, keywords, fallback_letter)
    except KeyError:
        return None


def _col_letter_to_index(letter: str) -> int:
    result = 0
    for char in letter.upper():
        if "A" <= char <= "Z":
            result = result * 26 + ord(char) - ord("A") + 1
    return result - 1


def _is_blank_row(values: list[Any]) -> bool:
    return all(value is None or value == "" for value in values)


def _row_text(row: pd.Series) -> str:
    return " ".join(str(value) for value in row.tolist() if value is not None and not pd.isna(value))


def _unmatched_nonzero_risks(risk_df: pd.DataFrame) -> pd.DataFrame:
    return risk_df[(risk_df["RISK_ACCR_INTEREST"] != ZERO) & (~risk_df["RISK_USED"])]


def _group_indexes(df: pd.DataFrame, columns: list[str]) -> dict[tuple[Any, ...], list[int]]:
    groups: dict[tuple[Any, ...], list[int]] = {}
    for idx, row in df.iterrows():
        key = tuple(row[col] for col in columns)
        if any(value in (None, "") for value in key):
            continue
        groups.setdefault(key, []).append(idx)
    return groups


def _find_unique_sum_combo(candidates: pd.DataFrame, target: Decimal) -> list[int] | None:
    items = [(idx, int(row["LEFT_BALANCE"] * 100)) for idx, row in candidates.iterrows() if row["LEFT_BALANCE"] > ZERO]
    target_cents = int(target * 100)
    if target_cents <= 0 or len(items) > 24:
        return None
    found: list[list[int]] = []

    def dfs(start: int, remaining: int, path: list[int]) -> None:
        if len(found) > 1:
            return
        if remaining == 0:
            found.append(path.copy())
            return
        if remaining < 0:
            return
        for pos in range(start, len(items)):
            idx, amount = items[pos]
            if amount <= remaining:
                path.append(idx)
                dfs(pos + 1, remaining - amount, path)
                path.pop()

    dfs(0, target_cents, [])
    return found[0] if len(found) == 1 and len(found[0]) > 1 else None


def _find_zero_anchor_left(left_df: pd.DataFrame, risk_row: pd.Series, config: MapperConfig) -> int | None:
    same_org = left_df[left_df["LEFT_ORG_ID_FINAL"] == risk_row["RISK_ORG_ID"]]
    risk_text = _row_text(risk_row)
    risk_term = str(risk_row.get("RISK_TERM_TYPE") or "")

    if _has_any(risk_text, ("三方存管", "黄金交易所", "交易所", "券商")):
        for frame in [same_org, left_df]:
            for keyword in config.zero_anchor_keywords:
                if keyword not in risk_text:
                    continue
                matches = frame[frame.apply(lambda row: keyword in _row_text(row), axis=1)]
                if not matches.empty:
                    return int(matches.sort_values("LEFT_ROW_NO").index[0])
        category_matches = _zero_anchor_candidates(left_df, config.zero_anchor_target_category, risk_term="")
        if not category_matches.empty:
            return int(category_matches.sort_values("LEFT_ROW_NO").index[0])

    if "存放同业" in risk_text:
        bank_matches = _zero_anchor_candidates(same_org, "境内银行", risk_term=risk_term)
        if not bank_matches.empty:
            return int(bank_matches.sort_values("LEFT_ROW_NO").index[0])
        bank_matches = _zero_anchor_candidates(left_df, "境内银行", risk_term=risk_term)
        if not bank_matches.empty:
            return int(bank_matches.sort_values("LEFT_ROW_NO").index[0])

    category_matches = left_df[left_df["COUNTERPARTY_CATEGORY"] == config.zero_anchor_target_category]
    if not category_matches.empty:
        return int(category_matches.sort_values("LEFT_ROW_NO").index[0])
    return int(left_df.sort_values("LEFT_ROW_NO").index[0]) if not left_df.empty else None


def _zero_anchor_candidates(frame: pd.DataFrame, category: str, risk_term: str = "") -> pd.DataFrame:
    if frame.empty:
        return frame
    matches = frame[frame["COUNTERPARTY_CATEGORY"] == category]
    if risk_term:
        term_matches = matches[matches["LEFT_TERM_TYPE"] == risk_term]
        if not term_matches.empty:
            matches = term_matches
    positive_matches = matches[matches["LEFT_BALANCE"] > ZERO]
    return positive_matches if not positive_matches.empty else matches


def _has_any(text: str, keywords: tuple[str, ...]) -> bool:
    return any(keyword in text for keyword in keywords)


def _flag_risk(risk_df: pd.DataFrame, idx: int, reason: str) -> None:
    risk_df.at[idx, "ATTENTION_FLAG"] = "是"
    risk_df.at[idx, "ATTENTION_REASON"] = _append_reason(risk_df.at[idx, "ATTENTION_REASON"], reason)
    risk_df.at[idx, "RISK_MATCH_RULE"] = risk_df.at[idx, "RISK_MATCH_RULE"] or "UNMATCHED"


def _append_reason(existing: Any, reason: str) -> str:
    existing_text = str(existing or "")
    if not existing_text:
        return reason
    if reason in existing_text:
        return existing_text
    return f"{existing_text}；{reason}"


def _business_left_output(mapped_left: pd.DataFrame, source_label: str) -> pd.DataFrame:
    output = mapped_left.copy()
    output["SOURCE_FILE"] = source_label
    return output


def _left_result_output(mapped_left: pd.DataFrame) -> pd.DataFrame:
    output = mapped_left.drop(columns=["NORMAL_MATCH_USED"], errors="ignore").copy()
    interest_col = _find_existing_interest_col(output)
    if interest_col:
        output[interest_col] = output["ACCR_INTEREST"]
    return output


def _find_existing_interest_col(df: pd.DataFrame) -> str | None:
    for col in df.columns:
        if "应计利息" in str(col):
            return str(col)
    return None


def _risk_result_output(mapped_risk: pd.DataFrame) -> pd.DataFrame:
    return mapped_risk.drop(columns=["RISK_USED"], errors="ignore").copy()


def _to_excel(df: pd.DataFrame, writer: pd.ExcelWriter, sheet_name: str) -> None:
    safe_df = df.copy()
    for col in safe_df.columns:
        if safe_df[col].map(lambda value: isinstance(value, Decimal)).any():
            safe_df[col] = safe_df[col].map(lambda value: float(value) if isinstance(value, Decimal) else value)
    safe_df.to_excel(writer, sheet_name=sheet_name[:31], index=False)


def _format_sheet(ws) -> None:
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for column_cells in ws.columns:
        column_letter = get_column_letter(column_cells[0].column)
        values = [str(cell.value) for cell in column_cells[:200] if cell.value is not None]
        width = min(max([len(value) for value in values] + [10]) + 2, 45)
        ws.column_dimensions[column_letter].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0.00'
