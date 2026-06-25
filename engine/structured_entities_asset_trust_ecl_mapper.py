from __future__ import annotations

from decimal import Decimal, InvalidOperation
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


OUTPUT_FILE_NAME = "6-3-4_资管计划信托股权明细表_ECL_FINAL匹配结果.xlsx"


def run_asset_trust_equity_ecl_mapping(
    upload_dir: str | Path,
    output_file: str | Path | None = None,
) -> Path:
    upload_path = Path(upload_dir)
    position_file = _latest_file_by_prefix(upload_path, "6-3-4-20251231")
    impairment_file = _latest_impairment_file(upload_path)
    if position_file is None:
        raise FileNotFoundError("未找到 6-3-4-20251231 资管计划&信托&股权明细表，请先上传。")
    if impairment_file is None:
        raise FileNotFoundError("未找到包含 ECL_FINAL 的 5-7-3-2-20251231 减值明细表，请先上传。")

    position_df, header_row, sheet_name = _load_position_df(position_file)
    impairment_df = _load_impairment_df(impairment_file)
    result_df, unmatched_df, summary_df = _merge_ecl(position_df, impairment_df, position_file, impairment_file)
    _write_ecl_back_to_workbook(position_file, result_df, header_row, sheet_name)

    if output_file is not None:
        output_path = Path(output_file)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            result_df.to_excel(writer, sheet_name="匹配结果", index=False)
            unmatched_df.to_excel(writer, sheet_name="未匹配明细", index=False)
            summary_df.to_excel(writer, sheet_name="匹配汇总", index=False)
        return output_path
    return position_file


def _latest_file_by_prefix(upload_dir: Path, prefix: str) -> Path | None:
    files = sorted(
        [path for path in upload_dir.iterdir() if path.is_file() and path.name.startswith(prefix) and path.suffix.lower() in {".xlsx", ".xls"}],
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    return files[0] if files else None


def _latest_impairment_file(upload_dir: Path) -> Path | None:
    candidates = sorted(
        [path for path in upload_dir.iterdir() if path.is_file() and path.name.startswith("5-7-3-2-20251231") and path.suffix.lower() in {".xlsx", ".xls"}],
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    for path in candidates:
        try:
            columns = pd.read_excel(path, sheet_name="20251231", nrows=0).columns
        except Exception:
            try:
                columns = pd.read_excel(path, sheet_name=0, nrows=0).columns
            except Exception:
                continue
        if {"BIZ_TYPE", "CURRENT_BALNC", "ACCR_INTEREST", "ECL_FINAL"}.issubset(set(columns)):
            return path
    return None


def _load_position_df(path: Path) -> tuple[pd.DataFrame, int, str]:
    sheet_name = "持仓明细"
    try:
        preview = pd.read_excel(path, sheet_name=sheet_name, header=None, nrows=20)
    except Exception:
        sheet_name = 0
        preview = pd.read_excel(path, sheet_name=sheet_name, header=None, nrows=20)

    required = {"本金/初始本金(元)", "应收利息"}
    header_index = _detect_header_index(preview, required)
    df = pd.read_excel(path, sheet_name=sheet_name, header=header_index)
    df.columns = [_normalize_column(column) for column in df.columns]
    missing = {_normalize_column(column) for column in required} - set(df.columns)
    if missing:
        raise ValueError(f"6-3-4 资管计划&信托&股权明细表缺少列：{', '.join(sorted(missing))}")
    df = df.dropna(how="all").copy()
    df["_excel_row"] = df.index + header_index + 2
    return df, header_index + 1, str(sheet_name)


def _detect_header_index(preview: pd.DataFrame, required_columns: set[str]) -> int:
    required = {_normalize_column(column) for column in required_columns}
    for index, row in preview.iterrows():
        values = {_normalize_column(value) for value in row.tolist() if not pd.isna(value)}
        if required.issubset(values):
            return int(index)
    raise ValueError("未能识别6-3-4资管计划&信托&股权明细表表头行。")


def _load_impairment_df(path: Path) -> pd.DataFrame:
    try:
        df = pd.read_excel(path, sheet_name="20251231")
    except Exception:
        df = pd.read_excel(path, sheet_name=0)
    required = {"BIZ_TYPE", "CURRENT_BALNC", "ACCR_INTEREST", "ECL_FINAL"}
    missing = required - set(df.columns)
    if missing:
        raise ValueError(f"5-7-3-2 减值明细表缺少列：{', '.join(sorted(missing))}")
    biz_type = df["BIZ_TYPE"].map(_normalize_code)
    return df[biz_type == "06"].copy()


def _merge_ecl(
    position_df: pd.DataFrame,
    impairment_df: pd.DataFrame,
    position_file: Path,
    impairment_file: Path,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    position = position_df.copy()
    impairment = impairment_df.copy()

    position["匹配_CURRENT_BALNC"] = position["本金/初始本金(元)"].map(_normalize_amount)
    position["匹配_ACCR_INTEREST"] = position["应收利息"].map(_normalize_amount)
    position["匹配键"] = _join_key(position["匹配_CURRENT_BALNC"], position["匹配_ACCR_INTEREST"])

    impairment["匹配_CURRENT_BALNC"] = impairment["CURRENT_BALNC"].map(_normalize_amount)
    impairment["匹配_ACCR_INTEREST"] = impairment["ACCR_INTEREST"].map(_normalize_amount)
    impairment["匹配键"] = _join_key(impairment["匹配_CURRENT_BALNC"], impairment["匹配_ACCR_INTEREST"])

    grouped = (
        impairment.groupby("匹配键", dropna=False)
        .agg(
            ECL_FINAL=("ECL_FINAL", lambda values: pd.to_numeric(values, errors="coerce").fillna(0.0).sum()),
            减值明细匹配行数=("ECL_FINAL", "size"),
            BUSI_PK_ID=("BUSI_PK_ID", _join_unique_values if "BUSI_PK_ID" in impairment.columns else "size"),
            减值明细_CURRENT_BALNC=("CURRENT_BALNC", _join_unique_values),
            减值明细_ACCR_INTEREST=("ACCR_INTEREST", _join_unique_values),
        )
        .reset_index()
    )

    result = position.merge(grouped, on="匹配键", how="left")
    result["ECL_FINAL"] = pd.to_numeric(result["ECL_FINAL"], errors="coerce")
    result["匹配状态"] = result["ECL_FINAL"].notna().map({True: "已匹配", False: "未匹配"})
    result["未匹配原因"] = ""
    result.loc[result["匹配状态"] == "未匹配", "未匹配原因"] = (
        "未在5-7-3-2减值明细BIZ_TYPE=06中按CURRENT_BALNC+ACCR_INTEREST匹配到记录"
    )
    result["ECL_FINAL"] = result["ECL_FINAL"].fillna(0.0)
    unmatched = result[result["匹配状态"] == "未匹配"].copy()

    summary = pd.DataFrame(
        [
            ("6-3-4明细表", position_file.name),
            ("5-7-3-2减值明细表", impairment_file.name),
            ("匹配规则", "6-3-4：本金/初始本金(元)+应收利息；5-7-3-2：CURRENT_BALNC+ACCR_INTEREST；减值明细过滤BIZ_TYPE=06"),
            ("6-3-4明细行数", len(position)),
            ("5-7-3-2 BIZ_TYPE=06行数", len(impairment)),
            ("已匹配行数", int((result["匹配状态"] == "已匹配").sum())),
            ("未匹配行数", int((result["匹配状态"] == "未匹配").sum())),
            ("匹配ECL_FINAL合计", float(result.loc[result["匹配状态"] == "已匹配", "ECL_FINAL"].sum())),
            ("重复匹配键6-3-4行数", int(position.duplicated("匹配键", keep=False).sum())),
            ("重复匹配键减值明细行数", int(impairment.duplicated("匹配键", keep=False).sum())),
        ],
        columns=["项目", "内容"],
    )
    return result, unmatched, summary


def _write_ecl_back_to_workbook(path: Path, result_df: pd.DataFrame, header_row: int, sheet_name: str) -> None:
    workbook = load_workbook(path)
    worksheet = workbook[sheet_name] if sheet_name in workbook.sheetnames else workbook.worksheets[0]
    header_map = {
        _normalize_column(cell.value): cell.column
        for cell in worksheet[header_row]
        if cell.value is not None and str(cell.value).strip()
    }
    ecl_column = _ensure_column(worksheet, header_map, header_row, "ECL_FINAL")
    status_column = _ensure_column(worksheet, header_map, header_row, "ECL匹配状态")
    key_column = _ensure_column(worksheet, header_map, header_row, "ECL匹配键")
    reason_column = _ensure_column(worksheet, header_map, header_row, "ECL未匹配原因")

    for _, row in result_df.iterrows():
        excel_row = int(row["_excel_row"])
        worksheet.cell(row=excel_row, column=ecl_column, value=float(row["ECL_FINAL"]))
        worksheet.cell(row=excel_row, column=status_column, value=str(row["匹配状态"]))
        worksheet.cell(row=excel_row, column=key_column, value=str(row["匹配键"]))
        worksheet.cell(row=excel_row, column=reason_column, value=str(row["未匹配原因"]))
    workbook.save(path)


def _ensure_column(worksheet, header_map: dict[str, int], header_row: int, column_name: str) -> int:
    normalized = _normalize_column(column_name)
    if normalized in header_map:
        return header_map[normalized]
    next_column = worksheet.max_column + 1
    worksheet.cell(row=header_row, column=next_column, value=column_name)
    header_map[normalized] = next_column
    return next_column


def _normalize_code(value) -> str:
    text = str(value or "").strip()
    if text.endswith(".0"):
        text = text[:-2]
    return text.zfill(2)


def _normalize_amount(value) -> str:
    if pd.isna(value):
        return "0.00"
    text = str(value).replace(",", "").strip()
    if not text or text.lower() == "nan" or text == "-":
        return "0.00"
    try:
        return str(Decimal(text).quantize(Decimal("0.01")))
    except (InvalidOperation, ValueError):
        return text


def _normalize_column(value) -> str:
    return "".join(str(value).split())


def _join_key(current_balnc: pd.Series, accr_interest: pd.Series) -> pd.Series:
    return current_balnc.astype(str) + "|" + accr_interest.astype(str)


def _join_unique_values(values: pd.Series) -> str:
    unique_values = []
    for value in values:
        if pd.isna(value):
            continue
        text = str(value)
        if text not in unique_values:
            unique_values.append(text)
    return ";".join(unique_values[:20])
