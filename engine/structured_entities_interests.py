from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re
from typing import Any

import pandas as pd
from openpyxl import load_workbook


DISCLOSURE_COLUMNS = [
    "交易性金融资产",
    "债权投资",
    "其他债权投资",
    "其他权益工具投资",
]


@dataclass(frozen=True)
class SourceSpec:
    name: str
    prefixes: tuple[str, ...]
    keywords: tuple[str, ...]


@dataclass(frozen=True)
class IndicatorSpec:
    code: str
    name: str
    source_name: str
    disclosure_column: str
    subject_rule: str = ""


SOURCE_SPECS = [
    SourceSpec("资产支持证券", ("6-3-1-20251231",), ("资产证券化", "资产支持")),
    SourceSpec("基金投资", ("6-3-2-20251231",), ()),
    SourceSpec("资管计划&信托", ("6-3-4-20251231",), ("资管", "资产管理", "信托")),
    SourceSpec("理财产品", ("6-3-4-20251231",), ("理财",)),
    SourceSpec("其他投资", ("6-3-4-20251231",), ("股权", "股票", "其他")),
]


INDICATOR_SPECS = [
    IndicatorSpec(
        "B1003",
        "资产支持证券—交易性金融资产余额-结构化主体",
        "资产支持证券",
        "交易性金融资产",
        "11030113科目余额借方轧差值+11030213科目余额借方轧差值+11030313科目余额借方轧差值",
    ),
    IndicatorSpec("B1004", "资产支持证券—债权投资余额-结构化主体", "资产支持证券", "债权投资"),
    IndicatorSpec("B1005", "资产支持证券—其他债权投资余额-结构化主体", "资产支持证券", "其他债权投资"),
    IndicatorSpec("B1006", "基金投资 - 交易性金融资产余额-结构化主体", "基金投资", "交易性金融资产"),
    IndicatorSpec("B1007", "资管计划&信托—交易性金融资产余额-结构化主体", "资管计划&信托", "交易性金融资产"),
    IndicatorSpec("B1008", "资管计划&信托—债权投资余额-结构化主体", "资管计划&信托", "债权投资"),
    IndicatorSpec(
        "B0414",
        "理财产品—交易性金融资产余额-结构化主体",
        "理财产品",
        "交易性金融资产",
        "11010102科目余额借方轧差值+11010202科目余额借方轧差值+11010302科目余额借方轧差值+"
        "11020102科目余额借方轧差值+11020202科目余额借方轧差值+11110102科目余额借方轧差值+"
        "11110202科目余额借方轧差值+11110302科目余额借方轧差值+11120102科目余额借方轧差值+"
        "11120202科目余额借方轧差值+13013102科目余额借方轧差值+13013302科目余额借方轧差值",
    ),
    IndicatorSpec("B1010", "其他投资—交易性金融资产余额-结构化主体", "其他投资", "交易性金融资产"),
    IndicatorSpec("B1011", "其他投资—其他权益工具投资余额-结构化主体", "其他投资", "其他权益工具投资"),
]


def build_unconsolidated_structured_entities_6_3_report(
    upload_dir: str | Path,
    report_config: dict[str, Any],
    institution_context: dict[str, Any] | None = None,
) -> pd.DataFrame:
    context = institution_context or {}
    period = str(context.get("period") or "")
    report_name = str(report_config.get("display_name") or "六、3 在未纳入合并财务报表范围的结构化主体中的权益")

    summaries = {summary["name"]: summary for summary in [_summarize_spec(Path(upload_dir), spec) for spec in SOURCE_SPECS]}
    upload_path = Path(upload_dir)
    rows = [_indicator_row(upload_path, period, report_name, indicator, summaries.get(indicator.source_name)) for indicator in INDICATOR_SPECS]

    for index, row in enumerate(rows, start=1):
        row["序号"] = index
    return pd.DataFrame(rows)


def _indicator_row(upload_dir: Path, period: str, report_name: str, indicator: IndicatorSpec, summary: dict[str, Any] | None) -> dict[str, Any]:
    if indicator.code == "B1007":
        return _b1007_row(upload_dir, period, report_name, indicator)
    if indicator.code == "B1008":
        return _b1008_row(upload_dir, period, report_name, indicator)
    if indicator.subject_rule:
        return _subject_rule_row(upload_dir, period, report_name, indicator)

    values = summary["values"] if summary else {column: 0.0 for column in DISCLOSURE_COLUMNS}
    amount = values[indicator.disclosure_column]
    rule_text = f"按上传明细表筛选{indicator.source_name}，按会计分类归集至{indicator.disclosure_column}，金额单位由元换算为千元。"
    if indicator.source_name == "资产支持证券" and summary and "ECL_FINAL" in str(summary.get("source", "")):
        rule_text = (
            "按6-3-1债券投资持仓表ECL_FINAL整合结果，筛选产品分类=金融债-资产证券化产品，"
            f"会计分类归集至{indicator.disclosure_column}，金额=本金+应收利息+利息调整-ECL_FINAL，单位由元换算为千元。"
        )
    elif indicator.source_name == "基金投资":
        rule_text = "按6-3-2基金投资持仓表，筛选是否合并<>'是'，金额=本金+应收利息+公允价值变动，单位由元换算为千元。"
    return {
        "期间": period,
        "机构口径": "集团",
        "报表名称": report_name,
        "指标编码": indicator.code,
        "指标名称": indicator.name,
        "数据来源": summary["source"] if summary else "未找到上传明细表",
        "指标类型": "基础指标",
        "加工规则": rule_text,
        "交易性金融资产": amount if indicator.disclosure_column == "交易性金融资产" else pd.NA,
        "债权投资": amount if indicator.disclosure_column == "债权投资" else pd.NA,
        "其他债权投资": amount if indicator.disclosure_column == "其他债权投资" else pd.NA,
        "其他权益工具投资": amount if indicator.disclosure_column == "其他权益工具投资" else pd.NA,
        "生成金额-集团": amount,
        "生成金额-本行": pd.NA,
        "计算状态": summary["status"] if summary else "未计算",
        "计算说明": summary["note"] if summary else "未找到对应来源明细。",
    }


def _b1007_row(upload_dir: Path, period: str, report_name: str, indicator: IndicatorSpec) -> dict[str, Any]:
    asset_amount, asset_note, asset_source = _b1007_asset_management_amount(upload_dir)
    trust_amount, trust_note, trust_source = _b1007_trust_amount(upload_dir)
    amount = (asset_amount + trust_amount) / 1000.0
    sources = "、".join(source for source in [asset_source, trust_source] if source)
    return {
        "期间": period,
        "机构口径": "集团",
        "报表名称": report_name,
        "指标编码": indicator.code,
        "指标名称": indicator.name,
        "数据来源": sources or "未找到上传明细表",
        "指标类型": "基础指标",
        "加工规则": (
            "6-3-3资管计划持仓表：产品类型='资管计划' AND 是否合并='否'，金额=本金+公允价值变动；"
            "6-3-4资管计划&信托&股权明细表：产品代码/业务编号='信托计划' AND 会计分类='FVTPL' "
            "AND 是否第三方结构化主体<>'是'，金额=本金/初始本金+公允价值变动；单位由元换算为千元。"
        ),
        "交易性金融资产": amount,
        "债权投资": pd.NA,
        "其他债权投资": pd.NA,
        "其他权益工具投资": pd.NA,
        "生成金额-集团": amount,
        "生成金额-本行": pd.NA,
        "计算状态": "已计算" if sources else "未计算",
        "计算说明": f"{asset_note} {trust_note}".strip(),
    }


def _b1008_row(upload_dir: Path, period: str, report_name: str, indicator: IndicatorSpec) -> dict[str, Any]:
    amount, note, source = _b1008_trust_ac_amount(upload_dir)
    value = amount / 1000.0
    return {
        "期间": period,
        "机构口径": "集团",
        "报表名称": report_name,
        "指标编码": indicator.code,
        "指标名称": indicator.name,
        "数据来源": source or "未找到上传明细表",
        "指标类型": "基础指标",
        "加工规则": (
            "6-3-4资管计划&信托&股权明细表：产品类型='信托' AND 是否第三方结构化主体<>'是' "
            "AND 会计分类='AC'，金额=本金/初始本金+应收利息-ECL_FINAL，单位由元换算为千元。"
        ),
        "交易性金融资产": pd.NA,
        "债权投资": value,
        "其他债权投资": pd.NA,
        "其他权益工具投资": pd.NA,
        "生成金额-集团": value,
        "生成金额-本行": pd.NA,
        "计算状态": "已计算" if source else "未计算",
        "计算说明": note,
    }


def _b1008_trust_ac_amount(upload_dir: Path) -> tuple[float, str, str]:
    path = _latest_file_by_prefix(upload_dir, "6-3-4-20251231")
    if path is None:
        return 0.0, "未找到6-3-4资管计划&信托&股权明细表，B1008按0处理。", ""
    records = _worksheet_records(path, {"产品类型", "会计分类", "是否第三方结构化主体", "本金/初始本金(元)", "应收利息", "ECL_FINAL"})
    selected = [
        record
        for record in records
        if str(record.get("产品类型") or "").strip() == "信托"
        and str(record.get("会计分类") or "").strip().upper() == "AC"
        and str(record.get("是否第三方结构化主体") or "").strip() != "是"
    ]
    amount = sum(
        _to_float(record.get("本金/初始本金(元)"))
        + _to_float(record.get("应收利息"))
        - _to_float(record.get("ECL_FINAL"))
        for record in selected
    )
    return float(amount), f"{path.name}筛选信托、AC且非第三方结构化主体{len(selected)}行。", path.name


def _b1007_asset_management_amount(upload_dir: Path) -> tuple[float, str, str]:
    path = _latest_file_by_prefix(upload_dir, "6-3-3-20251231")
    if path is None:
        return 0.0, "未找到6-3-3资管计划持仓表，资管计划部分按0处理。", ""
    frame = _load_position_frame(path)
    required = {"产品类型", "本金(元)", "公允价值变动(元)"}
    missing = required - set(frame.columns)
    if missing:
        return 0.0, f"{path.name}缺少列：{', '.join(sorted(missing))}，资管计划部分按0处理。", path.name
    selected = frame.loc[frame["产品类型"].astype(str).str.strip().eq("资管计划")].copy()
    if "是否合并" in selected.columns:
        selected = selected.loc[selected["是否合并"].astype(str).str.strip().eq("否")].copy()
    else:
        selected = selected.iloc[0:0].copy()
    amount = _numeric_series(selected["本金(元)"]).add(_numeric_series(selected["公允价值变动(元)"]), fill_value=0.0).sum()
    return float(amount), f"{path.name}筛选资管计划{len(selected)}行。", path.name


def _b1007_trust_amount(upload_dir: Path) -> tuple[float, str, str]:
    path = _latest_file_by_prefix(upload_dir, "6-3-4-20251231")
    if path is None:
        return 0.0, "未找到6-3-4资管计划&信托&股权明细表，信托计划部分按0处理。", ""
    records = _worksheet_records(path, {"产品代码/业务编号", "会计分类", "是否第三方结构化主体", "本金/初始本金(元)", "公允价值变动(元)"})
    selected = [
        record
        for record in records
        if str(record.get("产品代码/业务编号") or "").strip() == "信托计划"
        and str(record.get("会计分类") or "").strip().upper() == "FVTPL"
        and str(record.get("是否第三方结构化主体") or "").strip() != "是"
    ]
    amount = sum(_to_float(record.get("本金/初始本金(元)")) + _to_float(record.get("公允价值变动(元)")) for record in selected)
    return float(amount), f"{path.name}筛选信托计划FVTPL且非第三方结构化主体{len(selected)}行。", path.name


def _worksheet_records(path: Path, required_columns: set[str]) -> list[dict[str, Any]]:
    workbook = load_workbook(path, data_only=False)
    worksheet = workbook["持仓明细"] if "持仓明细" in workbook.sheetnames else workbook.worksheets[0]
    required = {_normalize_column(column) for column in required_columns}
    header_row = None
    header_map: dict[int, str] = {}
    for row_index in range(1, min(worksheet.max_row, 30) + 1):
        values = {_normalize_column(worksheet.cell(row_index, column).value) for column in range(1, worksheet.max_column + 1)}
        if required.issubset(values):
            header_row = row_index
            for column in range(1, worksheet.max_column + 1):
                value = worksheet.cell(row_index, column).value
                if value is not None and str(value).strip():
                    header_map[column] = _normalize_column(value)
            break
    if header_row is None:
        raise ValueError(f"未能识别{path.name}表头行。")
    records: list[dict[str, Any]] = []
    for row_index in range(header_row + 1, worksheet.max_row + 1):
        record = {name: _cell_value(worksheet.cell(row_index, column).value) for column, name in header_map.items()}
        if any(value not in {None, ""} for value in record.values()):
            records.append(record)
    return records


def _cell_value(value: Any) -> Any:
    if isinstance(value, str) and value.startswith("="):
        formula = value[1:].strip()
        if re.fullmatch(r"[+-]?\d+(?:\.\d+)?", formula):
            return float(formula)
    return value


def _subject_rule_row(upload_dir: Path, period: str, report_name: str, indicator: IndicatorSpec) -> dict[str, Any]:
    amount, status, note = _calculate_subject_rule(upload_dir, indicator.subject_rule)
    value = amount if amount is not None else pd.NA
    return {
        "期间": period,
        "机构口径": "集团",
        "报表名称": report_name,
        "指标编码": indicator.code,
        "指标名称": indicator.name,
        "数据来源": "科目余额表",
        "指标类型": "基础指标",
        "加工规则": indicator.subject_rule,
        "交易性金融资产": value if indicator.disclosure_column == "交易性金融资产" else pd.NA,
        "债权投资": value if indicator.disclosure_column == "债权投资" else pd.NA,
        "其他债权投资": value if indicator.disclosure_column == "其他债权投资" else pd.NA,
        "其他权益工具投资": value if indicator.disclosure_column == "其他权益工具投资" else pd.NA,
        "生成金额-集团": value,
        "生成金额-本行": pd.NA,
        "计算状态": status,
        "计算说明": note,
    }


def _calculate_subject_rule(upload_dir: Path, rule_text: str) -> tuple[float | None, str, str]:
    from engine.rule_calculator import RuleCalculator

    calculator = RuleCalculator(upload_dir, subject_balance_prefix="1-1-", report_period="20251231")
    amount = calculator._calculate_subject_balance_rule(rule_text)
    if amount is not None:
        return amount, "已计算", "按科目余额借方轧差规则计算，含系统通用审计调整/合并抵销处理，单位转换为千元。"

    term_matches = list(re.finditer(r"([+\-])?\s*(\d{4,8})\s*科目余额\s*(借方|贷方)\s*轧差值", rule_text))
    if term_matches:
        subject_df = calculator._load_subject_balance_df()
        adjustment_df = calculator._load_adjustment_df()
        total = 0.0
        missing_codes: list[str] = []
        for match in term_matches:
            sign = -1.0 if match.group(1) == "-" else 1.0
            code = match.group(2)
            side = match.group(3)
            total += sign * calculator._account_amount(code, side, True, "余额")
            if not _has_account_code(subject_df, code) and not _has_account_code(adjustment_df, code):
                missing_codes.append(code)
        note = "按科目余额借方轧差规则逐项计算，含系统通用审计调整/合并抵销处理，单位转换为千元。"
        if missing_codes:
            note += f" 未找到科目按0处理：{', '.join(missing_codes)}。"
        return total / 1000.0, "已计算", note

    subject_df = calculator._load_subject_balance_df()
    codes = re.findall(r"(?<!\d)(\d{4,8})(?!\d)\s*科目余额借方轧差值", rule_text)
    missing_codes = [code for code in codes if subject_df.loc[subject_df["account_code"].astype(str) == code].empty]
    if codes and len(missing_codes) == len(codes):
        return 0.0, "已计算", f"规则已解析；科目余额表未找到科目：{', '.join(missing_codes)}，按0列示。"
    return None, "未计算", "科目余额规则未匹配到账户或基础文件。"


def _has_account_code(frame: pd.DataFrame, code: str) -> bool:
    if frame.empty:
        return False
    for column in ["account_code", "level1_code", "level2_code", "level3_code"]:
        if column in frame.columns and frame[column].astype(str).eq(code).any():
            return True
    return False


def _summarize_spec(upload_dir: Path, spec: SourceSpec) -> dict[str, Any]:
    values = {column: 0.0 for column in DISCLOSURE_COLUMNS}
    notes: list[str] = []
    sources: list[str] = []

    for prefix in spec.prefixes:
        path = _source_file_for_spec(upload_dir, spec, prefix)
        if path is None:
            notes.append(f"未找到{prefix}开头的上传文件。")
            continue
        frame = _load_position_frame(path, sheet_name="整合结果" if _is_asset_support_ecl_file(spec, path) else 0)
        selected = _select_rows(frame, spec)
        if selected.empty:
            notes.append(f"{path.name}未筛选到{spec.name}明细。")
            sources.append(path.name)
            continue
        amount = _amount_series(selected, include_ecl_final=_is_asset_support_ecl_file(spec, path))
        categories = selected.apply(_disclosure_column_for_row, axis=1)
        for column in DISCLOSURE_COLUMNS:
            values[column] += float(amount[categories == column].sum() / 1000)
        notes.append(f"{path.name}筛选{len(selected)}行。")
        sources.append(path.name)

    status = "已计算" if sources else "未计算"
    return {
        "name": spec.name,
        "values": values,
        "source": "、".join(dict.fromkeys(sources)) if sources else "未找到上传明细表",
        "status": status,
        "note": " ".join(notes) if notes else "按上传明细表自动汇总。",
    }


def _source_file_for_spec(upload_dir: Path, spec: SourceSpec, prefix: str) -> Path | None:
    if spec.name == "资产支持证券":
        output_dir = upload_dir.parent / "output"
        ecl_file = _latest_file_by_pattern(output_dir, "6-3-1_*ECL_FINAL*.xlsx")
        if ecl_file is not None:
            return ecl_file
    return _latest_file_by_prefix(upload_dir, prefix)


def _is_asset_support_ecl_file(spec: SourceSpec, path: Path) -> bool:
    return spec.name == "资产支持证券" and "ECL_FINAL" in path.name


def _latest_file_by_pattern(directory: Path, pattern: str) -> Path | None:
    if not directory.exists():
        return None
    files = sorted(
        [path for path in directory.glob(pattern) if path.is_file() and path.suffix.lower() in {".xlsx", ".xls"}],
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    return files[0] if files else None


def _latest_file_by_prefix(upload_dir: Path, prefix: str, suffixes: set[str] | None = None) -> Path | None:
    allowed_suffixes = suffixes or {".xlsx", ".xls"}
    files = sorted(
        [path for path in upload_dir.iterdir() if path.is_file() and path.name.startswith(prefix) and path.suffix.lower() in allowed_suffixes],
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    return files[0] if files else None


def _load_position_frame(path: Path, sheet_name: str | int = 0) -> pd.DataFrame:
    preview = pd.read_excel(path, sheet_name=sheet_name, header=None, nrows=12)
    header_index = int(preview.notna().sum(axis=1).idxmax())
    frame = pd.read_excel(path, sheet_name=sheet_name, header=header_index).dropna(how="all").copy()
    frame.columns = [_normalize_column(column) for column in frame.columns]
    return frame


def _select_rows(frame: pd.DataFrame, spec: SourceSpec) -> pd.DataFrame:
    if spec.name == "基金投资":
        selected = frame.copy()
        if "是否合并" in selected.columns:
            selected = selected.loc[~selected["是否合并"].astype(str).str.strip().eq("是")].copy()
        return selected

    text_columns = [column for column in ["产品类型", "产品分类", "产品代码/业务编号", "产品名称", "分类2"] if column in frame.columns]
    if not text_columns:
        return frame.iloc[0:0].copy()
    combined = frame[text_columns].fillna("").astype(str).agg(" ".join, axis=1)
    mask = combined.str.contains("|".join(spec.keywords), regex=True, na=False)

    if spec.name == "其他投资":
        excluded = combined.str.contains("资管|资产管理|信托|理财", regex=True, na=False)
        mask = mask & ~excluded
    return frame[mask].copy()


def _amount_series(frame: pd.DataFrame, include_ecl_final: bool = False) -> pd.Series:
    if include_ecl_final:
        amount_columns = ["本金(元)", "应收利息(元)", "应收利息（元）", "利息调整(元)"]
        total = pd.Series(0.0, index=frame.index)
        for column in dict.fromkeys(_normalize_column(item) for item in amount_columns):
            if column in frame.columns:
                total = total + pd.to_numeric(frame[column], errors="coerce").fillna(0.0)
        if "ECL_FINAL" in frame.columns:
            ecl = pd.to_numeric(frame["ECL_FINAL"], errors="coerce").fillna(0.0)
            total = total - ecl
        return total
    if "总额" in frame.columns:
        return pd.to_numeric(frame["总额"], errors="coerce").fillna(0.0)
    market_col = _first_existing_column(frame, ["持仓市值/账面价值(元)", "持仓市值(元)", "账面价值(元)", "本金/初始本金(元)", "本金(元)"])
    interest_col = _first_existing_column(frame, ["应收利息", "应收利息(元)", "应收利息（元）"])
    market = pd.to_numeric(frame[market_col], errors="coerce").fillna(0.0) if market_col else pd.Series(0.0, index=frame.index)
    interest = pd.to_numeric(frame[interest_col], errors="coerce").fillna(0.0) if interest_col else pd.Series(0.0, index=frame.index)
    return market + interest


def _disclosure_column_for_row(row: pd.Series) -> str:
    category = _text(row.get("分类1"))
    if category in DISCLOSURE_COLUMNS:
        return category

    accounting = _text(row.get("会计分类")).upper()
    product_text = " ".join(_text(row.get(column)) for column in ["产品类型", "产品分类", "产品代码/业务编号", "产品名称"])
    if "FVTPL" in accounting or "交易" in accounting:
        return "交易性金融资产"
    if accounting == "AC" or "摊余成本" in accounting:
        return "债权投资"
    if "FVOCI" in accounting:
        if any(keyword in product_text for keyword in ["股权", "股票", "权益"]):
            return "其他权益工具投资"
        return "其他债权投资"
    return "交易性金融资产"


def _first_existing_column(frame: pd.DataFrame, candidates: list[str]) -> str | None:
    for candidate in candidates:
        normalized = _normalize_column(candidate)
        if normalized in frame.columns:
            return normalized
    return None


def _numeric_series(series: pd.Series) -> pd.Series:
    return pd.to_numeric(series, errors="coerce").fillna(0.0)


def _to_float(value: Any) -> float:
    if value is None or pd.isna(value):
        return 0.0
    text = str(value).replace(",", "").strip()
    if not text or text.lower() == "nan" or text == "-":
        return 0.0
    try:
        return float(Decimal(text))
    except (InvalidOperation, ValueError):
        return 0.0


def _normalize_column(value: Any) -> str:
    return "".join(str(value).split())


def _text(value: Any) -> str:
    if pd.isna(value):
        return ""
    return str(value).strip()
