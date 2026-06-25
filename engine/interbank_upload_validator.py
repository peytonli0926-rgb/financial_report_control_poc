from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, ROUND_HALF_UP, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


CENT = Decimal("0.01")
ZERO = Decimal("0.00")
PRINCIPAL_ACCOUNT_CODES = (
    "10110101",
    "10110102",
    "10110103",
    "10110104",
    "10110501",
    "10110502",
    "10110503",
    "10110701",
)
INTEREST_ACCOUNT_CODES = (
    "13010401",
    "13010402",
    "13010403",
    "13010404",
    "13010602",
)


@dataclass
class ValidationResult:
    name: str
    ledger_total: Decimal
    subject_total: Decimal
    difference: Decimal
    passed: bool
    message: str


def validate_parent_interbank_principal(upload_dir: str | Path) -> ValidationResult | None:
    upload_path = Path(upload_dir)
    detail_file = _find_file(upload_path, "5-2-1-1-1231")
    subject_file = _find_file(upload_path, "1-12-")
    if detail_file is None or subject_file is None:
        return None
    interest_total = sum_interbank_detail_amounts([detail_file], amount_col_index=11)
    ledger_total = sum_interbank_detail_amounts([detail_file], amount_col_index=13)
    book_total = (ledger_total + interest_total).quantize(CENT, rounding=ROUND_HALF_UP)
    subject_total = sum_subject_net_debit(subject_file, PRINCIPAL_ACCOUNT_CODES)
    result = _result("母行存放同业本金校验", ledger_total, subject_total)
    result.message = (
        f"母行存放同业明细与科目余额校验{'通过' if result.passed else '未通过'}："
        f"文件 {detail_file.name}；"
        f"K列应计利息合计 {interest_total}，"
        f"M列折合人民币余额合计 {ledger_total}，"
        f"K+M合计 {book_total}；"
        f"本行科目余额本金口径 {subject_total}，"
        f"M列差异 {result.difference}。"
    )
    return result


def validate_parent_interbank_interest(upload_dir: str | Path, mapped_file: str | Path | None = None) -> ValidationResult | None:
    upload_path = Path(upload_dir)
    detail_file = Path(mapped_file) if mapped_file else _find_file(upload_path, "5-2-1-1-1231")
    subject_file = _find_file(upload_path, "1-12-")
    if detail_file is None or subject_file is None or not detail_file.exists():
        return None
    if mapped_file:
        ledger_total = sum_mapped_result_interest(detail_file)
    else:
        ledger_total = sum_interbank_detail_amounts([detail_file], amount_col_index=11)
    subject_total = sum_subject_net_debit(subject_file, INTEREST_ACCOUNT_CODES)
    return _result("母行存放同业应计利息校验", ledger_total, subject_total)


def validate_subsidiary_interbank_principal(upload_dir: str | Path) -> ValidationResult | None:
    upload_path = Path(upload_dir)
    detail_files = _subsidiary_detail_files(upload_path)
    subject_file = _find_file(upload_path, "1-5-")
    if not detail_files or subject_file is None:
        return None
    ledger_total = sum_interbank_detail_amounts(detail_files, amount_col_index=13)
    subject_total = sum_subject_net_debit(subject_file, PRINCIPAL_ACCOUNT_CODES)
    return _result("子公司存放同业本金校验", ledger_total, subject_total)


def validate_subsidiary_interbank_interest(upload_dir: str | Path) -> ValidationResult | None:
    upload_path = Path(upload_dir)
    detail_files = _subsidiary_detail_files(upload_path)
    subject_file = _find_file(upload_path, "1-5-")
    if not detail_files or subject_file is None:
        return None
    ledger_total = sum_interbank_detail_amounts(detail_files, amount_col_index=11)
    subject_total = sum_subject_net_debit(subject_file, INTEREST_ACCOUNT_CODES)
    return _result("子公司存放同业应计利息校验", ledger_total, subject_total)


def validate_subsidiary_interbank_file_principal(upload_dir: str | Path, detail_file: str | Path) -> ValidationResult | None:
    upload_path = Path(upload_dir)
    subject_file = _find_file(upload_path, "1-5-")
    detail_path = Path(detail_file)
    if subject_file is None or not detail_path.exists():
        return None
    institution_codes = extract_detail_institution_codes(detail_path)
    ledger_total = sum_interbank_detail_amounts([detail_path], amount_col_index=13)
    subject_total = sum_subject_net_debit(subject_file, PRINCIPAL_ACCOUNT_CODES, institution_codes=institution_codes)
    return _result(f"{detail_path.name} 本金校验", ledger_total, subject_total)


def validate_subsidiary_interbank_file_interest(upload_dir: str | Path, detail_file: str | Path) -> ValidationResult | None:
    upload_path = Path(upload_dir)
    subject_file = _find_file(upload_path, "1-5-")
    detail_path = Path(detail_file)
    if subject_file is None or not detail_path.exists():
        return None
    institution_codes = extract_detail_institution_codes(detail_path)
    ledger_total = sum_interbank_detail_amounts([detail_path], amount_col_index=11)
    subject_total = sum_subject_net_debit(subject_file, INTEREST_ACCOUNT_CODES, institution_codes=institution_codes)
    return _result(f"{detail_path.name} 应计利息校验", ledger_total, subject_total)


def sum_interbank_detail_amounts(files: list[Path], amount_col_index: int) -> Decimal:
    total = ZERO
    for file_path in files:
        wb = load_workbook(file_path, read_only=True, data_only=True)
        try:
            ws = wb[wb.sheetnames[0]]
            blank_streak = 0
            for row_no, row in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
                values = list(row)
                if all(value is None or value == "" for value in values):
                    blank_streak += 1
                    if blank_streak >= 500:
                        break
                    continue
                blank_streak = 0
                if len(values) >= amount_col_index:
                    total += to_decimal(values[amount_col_index - 1])
        finally:
            wb.close()
    return total.quantize(CENT, rounding=ROUND_HALF_UP)


def sum_subject_net_debit(
    file_path: str | Path,
    account_codes: tuple[str, ...],
    institution_codes: set[str] | None = None,
) -> Decimal:
    target_codes = set(account_codes)
    total = ZERO
    wb = load_workbook(file_path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        for row in ws.iter_rows(min_row=11, values_only=True):
            values = list(row)
            if len(values) < 10:
                continue
            institution_code = normalize_institution_code(values[0] if values else None)
            if institution_codes and institution_code not in institution_codes:
                continue
            account_code = normalize_account_code(values[2] if len(values) > 2 else None)
            if account_code not in target_codes:
                continue
            ending_debit = to_decimal(values[8])
            ending_credit = to_decimal(values[9])
            total += ending_debit - ending_credit
    finally:
        wb.close()
    return total.quantize(CENT, rounding=ROUND_HALF_UP)


def extract_detail_institution_codes(file_path: str | Path) -> set[str]:
    codes: set[str] = set()
    wb = load_workbook(file_path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        blank_streak = 0
        for row in ws.iter_rows(min_row=3, values_only=True):
            values = list(row)
            if all(value is None or value == "" for value in values):
                blank_streak += 1
                if blank_streak >= 500:
                    break
                continue
            blank_streak = 0
            if len(values) >= 5:
                code = normalize_institution_code(values[4])
                if code:
                    codes.add(code)
    finally:
        wb.close()
    return codes


def sum_mapped_result_interest(file_path: str | Path) -> Decimal:
    total = ZERO
    wb = load_workbook(file_path, read_only=True, data_only=True)
    try:
        ws = wb["5-2-1-1_映射结果"] if "5-2-1-1_映射结果" in wb.sheetnames else wb[wb.sheetnames[0]]
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        interest_col = None
        for index, value in enumerate(headers, start=1):
            if str(value or "") == "ACCR_INTEREST":
                interest_col = index
                break
        if interest_col is None:
            interest_col = 11
        for row in ws.iter_rows(min_row=2, values_only=True):
            if len(row) >= interest_col:
                total += to_decimal(row[interest_col - 1])
    finally:
        wb.close()
    return total.quantize(CENT, rounding=ROUND_HALF_UP)


def normalize_account_code(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip().replace("\xa0", "").replace(" ", "")
    if text.endswith(".0"):
        text = text[:-2]
    return "".join(ch for ch in text if ch.isdigit())


def normalize_institution_code(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip().replace("\xa0", "").replace(" ", "")


def to_decimal(value: Any) -> Decimal:
    if value is None:
        return ZERO
    if isinstance(value, Decimal):
        return value.quantize(CENT, rounding=ROUND_HALF_UP)
    text = str(value).strip().replace(",", "")
    if not text or text.lower() == "nan":
        return ZERO
    if text.startswith("(") and text.endswith(")"):
        text = "-" + text[1:-1]
    try:
        return Decimal(text).quantize(CENT, rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError):
        return ZERO


def _result(name: str, ledger_total: Decimal, subject_total: Decimal) -> ValidationResult:
    difference = (ledger_total - subject_total).quantize(CENT, rounding=ROUND_HALF_UP)
    passed = abs(difference) <= CENT
    status = "对上" if passed else "未对上"
    message = f"{name}{status}：上传/映射金额合计 {ledger_total}，科目余额借方轧差 {subject_total}，差异 {difference}"
    return ValidationResult(name, ledger_total, subject_total, difference, passed, message)


def _find_file(upload_dir: Path, prefix: str) -> Path | None:
    matches = sorted(upload_dir.glob(f"{prefix}*.xlsx"), key=lambda path: path.stat().st_mtime, reverse=True)
    return matches[0] if matches else None


def _subsidiary_detail_files(upload_dir: Path) -> list[Path]:
    files = []
    for path in sorted(upload_dir.glob("5-2-1-*-1231*.xlsx")):
        if path.name.startswith("5-2-1-1-"):
            continue
        files.append(path)
    return files
