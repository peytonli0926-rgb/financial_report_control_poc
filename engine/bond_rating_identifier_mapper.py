from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook


RATING_PREFIX = "10-1-3-1-1231"
POSITION_PREFIX = "6-3-1-20251231"
SUMMARY_FILE_NAME = "6-3-1_债券投资持仓表_评级标识同步结果.xlsx"

MATCH_COLUMNS = ("债券代码", "会计分类", "总额")
RATING_COLUMN = "评级标识"
RATING_TOTAL_ALIASES = ("总额", "金额(元)", "金额", "总额(元)")


@dataclass(frozen=True)
class RatingSyncResult:
    position_file: Path
    rating_file: Path
    summary_file: Path
    position_rows: int
    rating_rows: int
    updated_rows: int
    unmatched_position_rows: int
    duplicate_rating_key_rows: int
    conflicting_rating_key_rows: int


def sync_bond_rating_identifier(
    upload_dir: str | Path,
    output_dir: str | Path,
) -> RatingSyncResult:
    upload_path = Path(upload_dir)
    output_path = Path(output_dir)
    position_file = _latest_file_by_prefix(upload_path, POSITION_PREFIX)
    rating_file = _latest_file_by_prefix(upload_path, RATING_PREFIX)
    if position_file is None:
        raise FileNotFoundError("未找到 6-3-1-20251231 债券投资持仓表，请先上传。")
    if rating_file is None:
        raise FileNotFoundError("未找到 10-1-3-1-1231 信用风险-风险集中度-债券评级表，请先上传。")
    if position_file.suffix.lower() != ".xlsx":
        raise ValueError("评级标识同步需要写回6-3-1持仓表，当前仅支持 .xlsx 目标文件。")

    rating_df = _load_rating_df(rating_file)
    rating_lookup, duplicate_key_rows, conflicting_key_rows, conflict_df = _build_rating_lookup(rating_df)
    updated_rows, position_rows, unmatched_df = _update_position_workbook(position_file, rating_lookup)

    summary_file = output_path / SUMMARY_FILE_NAME
    output_path.mkdir(parents=True, exist_ok=True)
    _write_summary(
        summary_file=summary_file,
        position_file=position_file,
        rating_file=rating_file,
        position_rows=position_rows,
        rating_rows=len(rating_df),
        updated_rows=updated_rows,
        unmatched_df=unmatched_df,
        conflict_df=conflict_df,
        duplicate_rating_key_rows=duplicate_key_rows,
        conflicting_rating_key_rows=conflicting_key_rows,
    )

    return RatingSyncResult(
        position_file=position_file,
        rating_file=rating_file,
        summary_file=summary_file,
        position_rows=position_rows,
        rating_rows=len(rating_df),
        updated_rows=updated_rows,
        unmatched_position_rows=len(unmatched_df),
        duplicate_rating_key_rows=duplicate_key_rows,
        conflicting_rating_key_rows=conflicting_key_rows,
    )


def _latest_file_by_prefix(upload_dir: Path, prefix: str) -> Path | None:
    if not upload_dir.exists():
        return None
    files = sorted(
        [
            path
            for path in upload_dir.iterdir()
            if path.is_file()
            and path.name.startswith(prefix)
            and path.suffix.lower() in {".xlsx", ".xls"}
        ],
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    return files[0] if files else None


def _load_rating_df(path: Path) -> pd.DataFrame:
    required = {"债券代码", "会计分类", RATING_COLUMN}
    header_index = _detect_excel_header_index(path, required, total_aliases=RATING_TOTAL_ALIASES)
    df = pd.read_excel(path, sheet_name=0, header=header_index)
    df = _normalize_rating_columns(df)
    required = set(MATCH_COLUMNS) | {RATING_COLUMN}
    missing = required - set(df.columns)
    if missing:
        raise ValueError(f"10-1-3-1-1231 债券评级表缺少列：{', '.join(sorted(missing))}")
    df = df[df["债券代码"].notna()].copy()
    df["匹配键"] = _build_match_key(df)
    df["评级标识_规范"] = df[RATING_COLUMN].map(_normalize_text)
    return df[df["评级标识_规范"] != ""].copy()


def _detect_excel_header_index(path: Path, required_columns: set[str], total_aliases: tuple[str, ...] = ("总额",)) -> int:
    preview = pd.read_excel(path, sheet_name=0, header=None, nrows=30)
    for index, row in preview.iterrows():
        values = {_normalize_header(value) for value in row.tolist() if not pd.isna(value)}
        if required_columns.issubset(values) and any(alias in values for alias in total_aliases):
            return int(index)
    raise ValueError(
        f"未能识别 {path.name} 的表头行，需包含：{', '.join(sorted(required_columns))}，"
        f"且金额列需为：{', '.join(total_aliases)} 之一"
    )


def _normalize_rating_columns(df: pd.DataFrame) -> pd.DataFrame:
    result = df.copy()
    for alias in RATING_TOTAL_ALIASES:
        if alias in result.columns:
            result = result.rename(columns={alias: "总额"})
            break
    return result


def _build_rating_lookup(rating_df: pd.DataFrame) -> tuple[dict[str, Any], int, int, pd.DataFrame]:
    duplicate_key_mask = rating_df.duplicated("匹配键", keep=False)
    duplicate_key_rows = int(duplicate_key_mask.sum())

    grouped = (
        rating_df.groupby("匹配键", dropna=False)
        .agg(
            债券代码=("债券代码", _join_unique_values),
            会计分类=("会计分类", _join_unique_values),
            总额=("总额", _join_unique_values),
            评级标识=(RATING_COLUMN, _join_unique_values),
            评级标识种类数=("评级标识_规范", lambda values: len(set(values))),
            评级表行数=(RATING_COLUMN, "size"),
        )
        .reset_index()
    )
    conflict_df = grouped[grouped["评级标识种类数"] > 1].copy()
    conflict_keys = set(conflict_df["匹配键"].astype(str))
    valid = grouped[~grouped["匹配键"].astype(str).isin(conflict_keys)].copy()
    lookup = dict(zip(valid["匹配键"].astype(str), valid["评级标识"]))
    conflicting_key_rows = int(rating_df["匹配键"].astype(str).isin(conflict_keys).sum())
    return lookup, duplicate_key_rows, conflicting_key_rows, conflict_df


def _update_position_workbook(position_file: Path, rating_lookup: dict[str, Any]) -> tuple[int, int, pd.DataFrame]:
    workbook = load_workbook(position_file)
    worksheet = workbook.worksheets[0]
    header_row, column_map = _detect_openpyxl_header(worksheet, set(MATCH_COLUMNS))
    rating_column_index = column_map.get(RATING_COLUMN)
    if rating_column_index is None:
        rating_column_index = worksheet.max_column + 1
        worksheet.cell(row=header_row, column=rating_column_index, value=RATING_COLUMN)
        column_map[RATING_COLUMN] = rating_column_index

    updated_rows = 0
    position_rows = 0
    unmatched_rows: list[dict[str, Any]] = []
    for row_index in range(header_row + 1, worksheet.max_row + 1):
        worksheet.cell(row=row_index, column=rating_column_index).value = None
        bond_code = worksheet.cell(row=row_index, column=column_map["债券代码"]).value
        if _normalize_text(bond_code) == "":
            continue
        position_rows += 1
        accounting_class = worksheet.cell(row=row_index, column=column_map["会计分类"]).value
        total_amount = _position_total_amount(worksheet, column_map, row_index)
        match_key = _join_key(bond_code, accounting_class, total_amount)
        rating = rating_lookup.get(match_key)
        if rating is None:
            unmatched_rows.append(
                {
                    "Excel行号": row_index,
                    "债券代码": bond_code,
                    "会计分类": accounting_class,
                    "总额": total_amount,
                    "匹配键": match_key,
                    "未匹配原因": "评级表中未找到相同债券代码、会计分类、总额且评级标识唯一的记录",
                }
            )
            continue
        worksheet.cell(row=row_index, column=rating_column_index, value=rating)
        updated_rows += 1

    workbook.save(position_file)
    return updated_rows, position_rows, pd.DataFrame(unmatched_rows)


def _position_total_amount(worksheet, column_map: dict[str, int], row_index: int) -> Any:
    total_value = worksheet.cell(row=row_index, column=column_map["总额"]).value
    if _normalize_amount(total_value) != "0.00":
        return total_value

    accounting_class = worksheet.cell(row=row_index, column=column_map["会计分类"]).value
    accounting_text = _normalize_accounting_class(accounting_class)
    principal = _decimal_cell(worksheet, column_map, row_index, "本金(元)")
    accrued_interest = _decimal_cell(worksheet, column_map, row_index, "应计利息(元)")
    receivable_interest = _decimal_cell(worksheet, column_map, row_index, "应收利息（元）")
    if accounting_text == "AC":
        return principal + accrued_interest + receivable_interest + _decimal_cell(worksheet, column_map, row_index, "利息调整(元)")
    if accounting_text in {"FVTPL", "FVOCI"}:
        return principal + accrued_interest + receivable_interest + _decimal_cell(worksheet, column_map, row_index, "公允价值变动(元)")
    return total_value


def _decimal_cell(worksheet, column_map: dict[str, int], row_index: int, column_name: str) -> Decimal:
    column_index = column_map.get(column_name)
    if column_index is None:
        return Decimal("0")
    return _to_decimal(worksheet.cell(row=row_index, column=column_index).value)


def _detect_openpyxl_header(worksheet, required_columns: set[str]) -> tuple[int, dict[str, int]]:
    max_scan_rows = min(30, worksheet.max_row)
    for row_index in range(1, max_scan_rows + 1):
        column_map: dict[str, int] = {}
        values = set()
        for column_index in range(1, worksheet.max_column + 1):
            header = _normalize_header(worksheet.cell(row=row_index, column=column_index).value)
            if not header:
                continue
            values.add(header)
            column_map.setdefault(header, column_index)
        if required_columns.issubset(values):
            return row_index, column_map
    raise ValueError("未能识别6-3-1债券投资持仓表表头行。")


def _write_summary(
    summary_file: Path,
    position_file: Path,
    rating_file: Path,
    position_rows: int,
    rating_rows: int,
    updated_rows: int,
    unmatched_df: pd.DataFrame,
    conflict_df: pd.DataFrame,
    duplicate_rating_key_rows: int,
    conflicting_rating_key_rows: int,
) -> None:
    summary_df = pd.DataFrame(
        [
            ("6-3-1持仓表", position_file.name),
            ("10-1-3-1评级表", rating_file.name),
            ("匹配规则", "债券代码+会计分类+总额一致，同步评级表的评级标识至6-3-1持仓表"),
            ("6-3-1持仓行数", position_rows),
            ("评级表有效评级行数", rating_rows),
            ("已更新评级标识行数", updated_rows),
            ("未匹配6-3-1行数", len(unmatched_df)),
            ("评级表重复匹配键行数", duplicate_rating_key_rows),
            ("评级表评级冲突行数", conflicting_rating_key_rows),
        ],
        columns=["项目", "内容"],
    )
    with pd.ExcelWriter(summary_file, engine="openpyxl") as writer:
        summary_df.to_excel(writer, sheet_name="同步汇总", index=False)
        unmatched_df.to_excel(writer, sheet_name="未匹配6-3-1明细", index=False)
        conflict_df.to_excel(writer, sheet_name="评级表冲突键", index=False)


def _build_match_key(df: pd.DataFrame) -> pd.Series:
    return (
        df["债券代码"].map(_normalize_bond_code)
        + "|"
        + df["会计分类"].map(_normalize_accounting_class)
        + "|"
        + df["总额"].map(_normalize_amount)
    )


def _join_key(bond_code: Any, accounting_class: Any, total_amount: Any) -> str:
    return "|".join(
        (
            _normalize_bond_code(bond_code),
            _normalize_accounting_class(accounting_class),
            _normalize_amount(total_amount),
        )
    )


def _normalize_header(value: Any) -> str:
    return str(value or "").strip()


def _normalize_text(value: Any) -> str:
    if pd.isna(value):
        return ""
    text = str(value).strip()
    return "" if text.lower() == "nan" else text


def _normalize_bond_code(value: Any) -> str:
    text = _normalize_text(value).replace(" ", "")
    if text.endswith(".0"):
        text = text[:-2]
    return text


def _normalize_accounting_class(value: Any) -> str:
    return _normalize_text(value).upper()


def _normalize_amount(value: Any) -> str:
    try:
        return str(_to_decimal(value).quantize(Decimal("0.01")))
    except (InvalidOperation, ValueError):
        return str(value).replace(",", "").strip()


def _to_decimal(value: Any) -> Decimal:
    if pd.isna(value):
        return Decimal("0")
    text = str(value).replace(",", "").strip()
    if not text or text.lower() == "nan" or text == "-":
        return Decimal("0")
    if text.startswith("="):
        return Decimal("0")
    return Decimal(text)


def _join_unique_values(values: pd.Series) -> str:
    unique_values: list[str] = []
    for value in values:
        text = _normalize_text(value)
        if text and text not in unique_values:
            unique_values.append(text)
    return ";".join(unique_values[:20])
