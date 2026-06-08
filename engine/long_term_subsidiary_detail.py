from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


CURRENT_DATA_SHEET_KEYWORD = "子公司当年数据上传"
SUMMARY_SHEET_NAME = "生成说明"
CURRENT_AMOUNT_COLUMN = "当年投资金额"
CURRENT_PERIOD_COLUMN = "本期"
PREVIOUS_PERIOD_COLUMN = "上期"


def build_long_term_subsidiary_detail_report(
    rule_df: pd.DataFrame,
    report_config: dict[str, Any],
    upload_dir: str | Path,
    report_period: str | None = None,
    institution_name: str | None = None,
    **_: Any,
) -> pd.DataFrame:
    detail_type = str(report_config.get("detail_report_type", ""))
    source_path = find_long_term_subsidiary_data_file(upload_dir, report_period)
    source_df = read_long_term_subsidiary_current_data(source_path)
    period = _normalize_period(report_period) or _period_from_file_name(source_path.name)
    period_label = _period_label(period)
    previous_period_label = _previous_period_label(period)

    if detail_type == "subsidiary_investment":
        return _build_subsidiary_investment_df(
            source_df,
            rule_df,
            report_config,
            source_path,
            period,
            period_label,
            previous_period_label,
            institution_name,
        )
    if detail_type == "subsidiary_control":
        return _build_subsidiary_control_df(
            source_df,
            rule_df,
            report_config,
            source_path,
            period,
            period_label,
            previous_period_label,
            institution_name,
        )
    raise ValueError(f"不支持的长期股权投资明细报表类型：{detail_type}")


def find_long_term_subsidiary_data_file(upload_dir: str | Path, report_period: str | None = None) -> Path:
    upload_path = Path(upload_dir)
    if not upload_path.exists():
        raise FileNotFoundError(f"上传目录不存在：{upload_path}")

    candidates = [
        path
        for path in upload_path.glob("1-11-*.xls*")
        if path.is_file() and not path.name.startswith("~$")
    ]
    if not candidates:
        raise FileNotFoundError("未找到 1-11- 长投子公司数据文件。")

    period = _normalize_period(report_period)
    if period:
        period_candidates = [path for path in candidates if period in path.name]
        if period_candidates:
            return max(period_candidates, key=lambda path: path.stat().st_mtime_ns)
    return max(candidates, key=lambda path: path.stat().st_mtime_ns)


def read_long_term_subsidiary_current_data(file_path: str | Path) -> pd.DataFrame:
    path = Path(file_path)
    workbook = pd.ExcelFile(path)
    sheet_names = sorted(
        workbook.sheet_names,
        key=lambda name: 0 if CURRENT_DATA_SHEET_KEYWORD in str(name) else 1,
    )
    for sheet_name in sheet_names:
        for header_row in range(0, 10):
            try:
                df = pd.read_excel(path, sheet_name=sheet_name, header=header_row, dtype=object)
            except Exception:
                continue
            normalized = _normalize_columns(df)
            if _has_required_current_data_columns(normalized):
                return _clean_current_data(normalized)
    raise ValueError("1-11 长投子公司数据未找到包含“子公司名称、当年投资金额、本行持股比例”等字段的明细表。")


def write_long_term_subsidiary_detail_excel(
    report_df: pd.DataFrame,
    report_config: dict[str, Any],
    output_path: str | Path,
) -> Path:
    if report_df is None or report_df.empty:
        raise ValueError("report_df 不能为空。")

    path = Path(output_path)
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("output_path 必须是 .xlsx 或 .xlsm 文件路径。")
    path.parent.mkdir(parents=True, exist_ok=True)

    display_name = str(report_config.get("display_name", "长期股权投资明细"))
    business_df = _business_columns(report_df)
    summary_df = _summary_df(report_df, display_name)

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        business_df.to_excel(writer, sheet_name=display_name[:31], index=False)
        report_df.to_excel(writer, sheet_name="生成数据", index=False)
        summary_df.to_excel(writer, sheet_name=SUMMARY_SHEET_NAME, index=False)

        _format_business_sheet(writer.book[display_name[:31]], business_df)
        _format_table_sheet(writer.book["生成数据"], report_df)
        _format_table_sheet(writer.book[SUMMARY_SHEET_NAME], summary_df)

    return path


def _build_subsidiary_investment_df(
    source_df: pd.DataFrame,
    rule_df: pd.DataFrame,
    report_config: dict[str, Any],
    source_path: Path,
    period: str,
    period_label: str,
    previous_period_label: str,
    institution_name: str | None,
) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    total = pd.to_numeric(source_df[CURRENT_AMOUNT_COLUMN], errors="coerce").sum()
    for index, row in source_df.iterrows():
        company_name = _normalize_text(row.get("子公司名称"))
        amount = _parse_number(row.get(CURRENT_AMOUNT_COLUMN))
        rows.append(
            {
                "序号": len(rows) + 1,
                "期间": period,
                "机构": institution_name or "",
                "报表名称": report_config.get("display_name", "五、8-2 长期股权投资-对子公司的投资"),
                "指标编码": _company_metric_code(rule_df, company_name) or "B0489/B0490",
                "指标名称": company_name,
                "指标类型": "明细",
                "数据来源": source_path.name,
                "生成金额-集团": amount,
                "生成金额-本行": amount,
                "子公司名称": company_name,
                period_label: amount,
                previous_period_label: pd.NA,
                "备注": _normalize_text(row.get("备注")),
                "取数字段": "子公司名称、当年投资金额(千元)",
                "计算状态": "已生成",
                "计算说明": "按 1-11 长投子公司数据逐行展开；上期明细需由报表平台或历史 1-11 数据补充。",
            }
        )

    rows.append(
        {
            "序号": len(rows) + 1,
            "期间": period,
            "机构": institution_name or "",
            "报表名称": report_config.get("display_name", "五、8-2 长期股权投资-对子公司的投资"),
            "指标编码": _total_metric_code(rule_df, "B0487"),
            "指标名称": "合计",
            "指标类型": "明细合计",
            "数据来源": source_path.name,
            "生成金额-集团": float(total),
            "生成金额-本行": float(total),
            "子公司名称": "合计",
            period_label: float(total),
            previous_period_label: pd.NA,
            "备注": "",
            "取数字段": "当年投资金额(千元)合计",
            "计算状态": "已生成",
            "计算说明": "按明细行当年投资金额汇总。",
        }
    )
    return pd.DataFrame(rows)


def _build_subsidiary_control_df(
    source_df: pd.DataFrame,
    rule_df: pd.DataFrame,
    report_config: dict[str, Any],
    source_path: Path,
    period: str,
    period_label: str,
    previous_period_label: str,
    institution_name: str | None,
) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    metric_codes = _control_metric_code_map(rule_df)
    for _, row in source_df.iterrows():
        company_name = _normalize_text(row.get("子公司名称"))
        shareholding_ratio = _parse_number(row.get("本行持股比例"))
        voting_ratio = _parse_number(row.get("本行表决权比例"))
        registered_capital = _parse_number(row.get("注册资本"))
        registered_place = _normalize_text(row.get("主要经营地/注册地区"))
        established_time = _normalize_text(row.get("成立时间"))
        main_business = _normalize_text(row.get("主营业务"))
        rows.append(
            {
                "序号": len(rows) + 1,
                "期间": period,
                "机构": institution_name or "",
                "报表名称": report_config.get("display_name", "五、8-4 长期股权投资-主要子公司持股比例及控制权情况表"),
                "指标编码": "/".join(metric_codes.values()),
                "指标名称": company_name,
                "指标类型": "明细",
                "数据来源": source_path.name,
                "子公司名称": company_name,
                metric_codes["子公司名称"]: company_name,
                f"{metric_codes['本行持股比例']}集团本期": shareholding_ratio,
                f"{metric_codes['本行持股比例']}集团上期": pd.NA,
                f"{metric_codes['本行表决权比例']}集团本期": voting_ratio,
                f"{metric_codes['本行表决权比例']}集团上期": pd.NA,
                f"{metric_codes['注册资本']}集团本期": registered_capital,
                f"{metric_codes['注册资本']}集团上期": pd.NA,
                metric_codes["注册地"]: registered_place,
                metric_codes["注册时间"]: established_time,
                metric_codes["主营业务"]: main_business,
                f"本行持股比例-{period_label}": shareholding_ratio,
                f"本行持股比例-{previous_period_label}": pd.NA,
                f"本行表决权比例-{period_label}": voting_ratio,
                f"本行表决权比例-{previous_period_label}": pd.NA,
                f"注册资本-{period_label}": registered_capital,
                f"注册资本-{previous_period_label}": pd.NA,
                "注册地": registered_place,
                "成立时间": established_time,
                "主营业务": main_business,
                "取数字段": "子公司名称、本行持股比例、本行表决权比例、注册资本、注册地、成立时间、主营业务",
                "计算状态": "已生成",
                "计算说明": (
                    "按 B0489 展开子公司行；B0492-B0497 在每个子公司维度上取 1-11 长投子公司数据。"
                    "上期比例和注册资本需由报表平台或历史 1-11 数据补充。"
                ),
            }
        )
    return pd.DataFrame(rows)


def _normalize_columns(df: pd.DataFrame) -> pd.DataFrame:
    normalized = df.copy()
    normalized.columns = [_normalize_header(column) for column in normalized.columns]
    return normalized


def _normalize_header(value: Any) -> str:
    text = _normalize_text(value)
    text = text.replace("(千元)", "").replace("（千元）", "").replace("\n", "")
    for prefix in ("当年", "本期"):
        if prefix in text and "投资金额" in text:
            return CURRENT_AMOUNT_COLUMN
    if "注册资本" in text:
        return "注册资本"
    return text


def _has_required_current_data_columns(df: pd.DataFrame) -> bool:
    required_columns = {"子公司名称", CURRENT_AMOUNT_COLUMN, "本行持股比例", "本行表决权比例", "注册资本"}
    return required_columns.issubset(set(df.columns))


def _clean_current_data(df: pd.DataFrame) -> pd.DataFrame:
    cleaned = df.copy()
    cleaned = cleaned.loc[cleaned["子公司名称"].map(_normalize_text).ne("")]
    cleaned = cleaned.loc[~cleaned["子公司名称"].map(_normalize_text).eq("合计")]
    cleaned = cleaned.dropna(how="all").reset_index(drop=True)
    return cleaned


def _company_metric_code(rule_df: pd.DataFrame, company_name: str) -> str:
    if rule_df is None or rule_df.empty or not company_name:
        return ""
    for _, row in rule_df.iterrows():
        item_name = _normalize_text(row.get("指标名称"))
        if company_name in item_name:
            return _normalize_text(row.get("指标编码"))
    return ""


def _total_metric_code(rule_df: pd.DataFrame, fallback: str) -> str:
    if rule_df is None or rule_df.empty:
        return fallback
    for _, row in rule_df.iterrows():
        item_name = _normalize_text(row.get("指标名称"))
        code = _normalize_text(row.get("指标编码"))
        if code == fallback or "合计" in item_name:
            return code or fallback
    return fallback


def _control_metric_code_map(rule_df: pd.DataFrame) -> dict[str, str]:
    default_codes = {
        "子公司名称": "B0489",
        "本行持股比例": "B0492",
        "本行表决权比例": "B0493",
        "注册资本": "B0494",
        "注册地": "B0495",
        "注册时间": "B0496",
        "主营业务": "B0497",
    }
    if rule_df is None or rule_df.empty:
        return default_codes

    result = dict(default_codes)
    for _, row in rule_df.iterrows():
        item_name = _normalize_text(row.get("指标名称"))
        item_code = _normalize_text(row.get("指标编码"))
        if not item_name or not item_code:
            continue
        if "子公司名称" in item_name:
            result["子公司名称"] = item_code
        elif "持股比例" in item_name:
            result["本行持股比例"] = item_code
        elif "表决权比例" in item_name:
            result["本行表决权比例"] = item_code
        elif "注册资本" in item_name:
            result["注册资本"] = item_code
        elif "注册地" in item_name:
            result["注册地"] = item_code
        elif "注册时间" in item_name or "成立时间" in item_name:
            result["注册时间"] = item_code
        elif "主营业务" in item_name:
            result["主营业务"] = item_code
    return result


def _business_columns(report_df: pd.DataFrame) -> pd.DataFrame:
    subsidiary_control_columns = _subsidiary_control_business_columns(report_df)
    if subsidiary_control_columns:
        return report_df[subsidiary_control_columns].copy()

    technical_columns = {
        "期间",
        "机构",
        "报表名称",
        "指标编码",
        "指标类型",
        "数据来源",
        "取数字段",
        "计算状态",
        "计算说明",
    }
    return report_df[[column for column in report_df.columns if column not in technical_columns]].copy()


def _subsidiary_control_business_columns(report_df: pd.DataFrame) -> list[str]:
    if report_df.empty or "子公司名称" not in report_df.columns:
        return []
    if not any(str(column).startswith("本行持股比例-") for column in report_df.columns):
        return []

    ordered_columns = ["子公司名称"]
    for prefix in ("本行持股比例-", "本行表决权比例-", "注册资本-"):
        ordered_columns.extend([column for column in report_df.columns if str(column).startswith(prefix)])
    ordered_columns.extend(["注册地", "成立时间", "主营业务"])
    return [column for column in ordered_columns if column in report_df.columns]


def _summary_df(report_df: pd.DataFrame, display_name: str) -> pd.DataFrame:
    source_file = ""
    if "数据来源" in report_df.columns and not report_df.empty:
        source_file = _normalize_text(report_df["数据来源"].iloc[0])
    return pd.DataFrame(
        [
            {"项目": "报表名称", "内容": display_name},
            {"项目": "生成时间", "内容": datetime.now().strftime("%Y-%m-%d %H:%M:%S")},
            {"项目": "数据来源", "内容": source_file},
            {"项目": "生成明细行数", "内容": len(report_df)},
            {
                "项目": "生成说明",
                "内容": "B0489 等明细指标按 1-11 长投子公司数据逐行展开；上期列不从本期数据推断。",
            },
        ]
    )


def _format_business_sheet(worksheet: Worksheet, df: pd.DataFrame) -> None:
    _format_table_sheet(worksheet, df)
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)


def _format_table_sheet(worksheet: Worksheet, df: pd.DataFrame) -> None:
    worksheet.freeze_panes = "A2"
    if worksheet.max_row >= 1 and worksheet.max_column >= 1:
        worksheet.auto_filter.ref = worksheet.dimensions
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            header = str(df.columns[cell.column - 1])
            if any(keyword in header for keyword in ["金额", "注册资本"]) or _looks_like_period_header(header):
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right")
            if "比例" in header:
                cell.number_format = "0.0%"
                cell.alignment = Alignment(horizontal="right")
    for column_cells in worksheet.columns:
        width = 10
        for cell in column_cells:
            if cell.value is not None:
                width = max(width, _display_width(str(cell.value)) + 2)
        worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(width, 48)


def _looks_like_period_header(value: str) -> bool:
    return "年" in value and "月" in value and "日" in value


def _parse_number(value: Any) -> float | None:
    if value is None or pd.isna(value):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = _normalize_text(value).replace(",", "")
    if not text or text in {"-", "－", "—", "报表平台取数"}:
        return None
    if text.endswith("%"):
        try:
            return float(text[:-1]) / 100
        except ValueError:
            return None
    try:
        return float(text)
    except ValueError:
        return None


def _normalize_text(value: Any) -> str:
    if value is None or pd.isna(value):
        return ""
    return " ".join(str(value).replace("\r", " ").replace("\n", " ").split())


def _normalize_period(value: Any) -> str:
    text = _normalize_text(value)
    if text.endswith(".0"):
        text = text[:-2]
    digits = "".join(ch for ch in text if ch.isdigit())
    return digits[:8] if len(digits) >= 8 else ""


def _period_from_file_name(file_name: str) -> str:
    import re

    match = re.search(r"(?<!\d)(20\d{6})(?!\d)", file_name)
    return match.group(1) if match else ""


def _period_label(period: str) -> str:
    if len(period) == 8:
        return f"{period[:4]}年{int(period[4:6])}月{int(period[6:8])}日"
    return CURRENT_PERIOD_COLUMN


def _previous_period_label(period: str) -> str:
    if len(period) == 8:
        return f"{int(period[:4]) - 1}年{int(period[4:6])}月{int(period[6:8])}日"
    return PREVIOUS_PERIOD_COLUMN


def _display_width(value: str) -> int:
    width = 0
    for char in value:
        width += 2 if ord(char) > 127 else 1
    return width
