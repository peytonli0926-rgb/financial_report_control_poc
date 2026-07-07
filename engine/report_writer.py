from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


RESULT_SHEET_NAME = "资产负债表生成结果"
INCOME_RESULT_SHEET_NAME = "利润表生成结果"
SUMMARY_SHEET_NAME = "生成说明"
PDF_VALIDATION_SHEET_NAME = "PDF校验结果"
PDF_VALIDATION_SUMMARY_SHEET_NAME = "PDF校验说明"
PDF_DIFFERENCE_SHEET_NAME = "差异项目清单"
AMOUNT_COLUMNS = ["生成金额-集团", "生成金额-本行"]
AMOUNT_NUMBER_FORMAT = '#,##0.00'


def write_balance_sheet_report(
    balance_df: pd.DataFrame,
    balance_check: dict[str, Any],
    output_path: str | Path,
) -> Path:
    """Write the generated balance sheet and generation summary to Excel."""
    if balance_df is None:
        raise ValueError("balance_df 不能为空。")
    if not isinstance(balance_df, pd.DataFrame):
        raise TypeError("balance_df 必须是 pandas DataFrame。")
    if balance_check is None:
        raise ValueError("balance_check 不能为空。")
    if not isinstance(balance_check, dict):
        raise TypeError("balance_check 必须是 dict。")

    path = Path(output_path)
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("output_path 必须是 .xlsx 或 .xlsm 文件路径。")

    try:
        path.parent.mkdir(parents=True, exist_ok=True)

        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            balance_df.to_excel(writer, sheet_name=RESULT_SHEET_NAME, index=False)

            summary_df = _build_summary_df(balance_df, balance_check)
            summary_df.to_excel(writer, sheet_name=SUMMARY_SHEET_NAME, index=False)

            result_ws = writer.book[RESULT_SHEET_NAME]
            summary_ws = writer.book[SUMMARY_SHEET_NAME]

            _format_result_sheet(result_ws, balance_df)
            _format_summary_sheet(summary_ws, summary_df)

        return path
    except OSError as exc:
        raise OSError(f"写入 Excel 报告失败：{exc}") from exc
    except Exception as exc:
        raise ValueError(f"生成 Excel 报告失败：{exc}") from exc


def write_pdf_validation_report(
    validation_df: pd.DataFrame,
    output_path: str | Path,
    report_name: str = "资产负债表",
    tolerance: float = 1,
) -> Path:
    """Write PDF validation result, summary, and difference list to Excel."""
    if validation_df is None:
        raise ValueError("validation_df 不能为空。")
    if not isinstance(validation_df, pd.DataFrame):
        raise TypeError("validation_df 必须是 pandas DataFrame。")

    path = Path(output_path)
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("output_path 必须是 .xlsx 或 .xlsm 文件路径。")

    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        difference_df = _build_difference_df(validation_df)
        summary_df = _build_pdf_validation_summary_df(report_name, tolerance)

        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            validation_df.to_excel(
                writer,
                sheet_name=PDF_VALIDATION_SHEET_NAME,
                index=False,
            )
            summary_df.to_excel(
                writer,
                sheet_name=PDF_VALIDATION_SUMMARY_SHEET_NAME,
                index=False,
            )
            difference_df.to_excel(
                writer,
                sheet_name=PDF_DIFFERENCE_SHEET_NAME,
                index=False,
            )

            _format_table_sheet(
                writer.book[PDF_VALIDATION_SHEET_NAME],
                validation_df,
            )
            _format_table_sheet(
                writer.book[PDF_VALIDATION_SUMMARY_SHEET_NAME],
                summary_df,
            )
            _format_table_sheet(
                writer.book[PDF_DIFFERENCE_SHEET_NAME],
                difference_df,
            )

        return path
    except OSError as exc:
        raise OSError(f"写入 PDF 校验报告失败：{exc}") from exc
    except Exception as exc:
        raise ValueError(f"生成 PDF 校验报告失败：{exc}") from exc


def write_report_generation_excel(
    report_df: pd.DataFrame,
    report_config: dict[str, Any],
    output_path: str | Path,
    balance_check: dict[str, Any] | None = None,
    trace_df: pd.DataFrame | None = None,
) -> Path:
    """Write a generated report using report type configuration."""
    if report_df is None:
        raise ValueError("report_df 不能为空。")
    if not isinstance(report_df, pd.DataFrame):
        raise TypeError("report_df 必须是 pandas DataFrame。")

    path = Path(output_path)
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("output_path 必须是 .xlsx 或 .xlsm 文件路径。")

    display_name = str(report_config.get("display_name", "报表"))
    result_sheet_name = _safe_sheet_name(f"{display_name}生成结果")

    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        summary_df = _build_generic_summary_df(report_df, report_config, balance_check)

        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            report_df.to_excel(writer, sheet_name=result_sheet_name, index=False)
            summary_df.to_excel(writer, sheet_name=SUMMARY_SHEET_NAME, index=False)
            if trace_df is not None and not trace_df.empty:
                trace_sheet_name = _safe_sheet_name("指标追溯明细")
                trace_df.to_excel(writer, sheet_name=trace_sheet_name, index=False)

            _format_table_sheet(writer.book[result_sheet_name], report_df)
            _format_table_sheet(writer.book[SUMMARY_SHEET_NAME], summary_df)
            if trace_df is not None and not trace_df.empty:
                _format_table_sheet(writer.book[trace_sheet_name], trace_df)

        return path
    except OSError as exc:
        raise OSError(f"写入{display_name} Excel 报告失败：{exc}") from exc
    except Exception as exc:
        raise ValueError(f"生成{display_name} Excel 报告失败：{exc}") from exc


def write_pdf_validation_excel(
    validation_df: pd.DataFrame,
    report_config: dict[str, Any],
    output_path: str | Path,
    tolerance: float = 1,
) -> Path:
    """Write PDF validation result using report type configuration."""
    return write_pdf_validation_report(
        validation_df,
        output_path,
        report_name=str(report_config.get("display_name", "报表")),
        tolerance=tolerance,
    )


def _safe_sheet_name(name: str) -> str:
    """Return an Excel-compatible worksheet title."""
    clean_name = str(name or "Sheet").translate(str.maketrans({char: " " for char in r'[]:*?/\\'})).strip()
    if not clean_name:
        clean_name = "Sheet"
    return clean_name[:31]


def write_income_statement_report(
    income_df: pd.DataFrame,
    output_path: str | Path,
) -> Path:
    """Write the generated income statement and generation summary to Excel."""
    if income_df is None:
        raise ValueError("income_df 不能为空。")
    if not isinstance(income_df, pd.DataFrame):
        raise TypeError("income_df 必须是 pandas DataFrame。")

    path = Path(output_path)
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("output_path 必须是 .xlsx 或 .xlsm 文件路径。")

    try:
        path.parent.mkdir(parents=True, exist_ok=True)

        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            income_df.to_excel(writer, sheet_name=INCOME_RESULT_SHEET_NAME, index=False)

            summary_df = _build_income_summary_df(income_df)
            summary_df.to_excel(writer, sheet_name=SUMMARY_SHEET_NAME, index=False)

            _format_table_sheet(writer.book[INCOME_RESULT_SHEET_NAME], income_df)
            _format_table_sheet(writer.book[SUMMARY_SHEET_NAME], summary_df)

        return path
    except OSError as exc:
        raise OSError(f"写入利润表 Excel 报告失败：{exc}") from exc
    except Exception as exc:
        raise ValueError(f"生成利润表 Excel 报告失败：{exc}") from exc


def _build_summary_df(
    balance_df: pd.DataFrame,
    balance_check: dict[str, Any],
) -> pd.DataFrame:
    report_name = _detect_report_name(balance_df)
    group_result = _format_balance_result(
        balance_check.get("group_balanced"),
        balance_check.get("group_diff"),
    )
    parent_result = _format_balance_result(
        balance_check.get("parent_balanced"),
        balance_check.get("parent_diff"),
    )

    return pd.DataFrame(
        [
            {"项目": "报表名称", "内容": report_name},
            {"项目": "生成时间", "内容": datetime.now().strftime("%Y-%m-%d %H:%M:%S")},
            {
                "项目": "数据来源说明",
                "内容": "由加工规则文件中已加工得到的集团金额和本行金额生成，金额单位为人民币千元。",
            },
            {"项目": "集团平衡检查结果", "内容": group_result},
            {"项目": "本行平衡检查结果", "内容": parent_result},
        ]
    )


def _build_income_summary_df(income_df: pd.DataFrame) -> pd.DataFrame:
    report_name = _detect_report_name(income_df, "利润表")
    return pd.DataFrame(
        [
            {"项目": "报表名称", "内容": report_name},
            {"项目": "生成时间", "内容": datetime.now().strftime("%Y-%m-%d %H:%M:%S")},
            {
                "项目": "数据来源说明",
                "内容": "由加工规则文件中已加工得到的集团金额和本行金额生成，金额单位为人民币千元。",
            },
        ]
    )


def _build_generic_summary_df(
    report_df: pd.DataFrame,
    report_config: dict[str, Any],
    balance_check: dict[str, Any] | None,
) -> pd.DataFrame:
    display_name = str(report_config.get("display_name", "报表"))
    report_name = _detect_report_name(report_df, display_name)
    rows = [
        {"项目": "报表名称", "内容": report_name},
        {"项目": "报表类型", "内容": display_name},
        {"项目": "期间类型", "内容": report_config.get("period_type", "")},
        {"项目": "生成时间", "内容": datetime.now().strftime("%Y-%m-%d %H:%M:%S")},
        {
            "项目": "数据来源说明",
            "内容": "由加工规则文件中已加工得到的集团金额和本行金额生成，金额单位为人民币千元。",
        },
    ]

    if balance_check is not None:
        rows.extend(
            [
                {
                    "项目": "集团平衡检查结果",
                    "内容": _format_balance_result(
                        balance_check.get("group_balanced"),
                        balance_check.get("group_diff"),
                    ),
                },
                {
                    "项目": "本行平衡检查结果",
                    "内容": _format_balance_result(
                        balance_check.get("parent_balanced"),
                        balance_check.get("parent_diff"),
                    ),
                },
            ]
        )

    return pd.DataFrame(rows)


def _build_pdf_validation_summary_df(report_name: str, tolerance: float) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {"项目": "校验对象", "内容": report_name},
            {"项目": "单位", "内容": "人民币千元"},
            {"项目": "容差", "内容": f"{tolerance:g} 千元"},
            {
                "项目": "一致判断规则",
                "内容": f"差异绝对值 <= {tolerance:g} 千元判断为一致；PDF金额为空判断为未提取到PDF披露值；其他情况判断为不一致。",
            },
            {"项目": "生成时间", "内容": datetime.now().strftime("%Y-%m-%d %H:%M:%S")},
        ]
    )


def _build_difference_df(validation_df: pd.DataFrame) -> pd.DataFrame:
    if validation_df.empty:
        return validation_df.copy()

    status_columns = [
        column
        for column in ["是否一致-集团", "是否一致-本行"]
        if column in validation_df.columns
    ]
    if not status_columns:
        return validation_df.iloc[0:0].copy()

    mask = pd.Series(False, index=validation_df.index)
    for column in status_columns:
        mask = mask | validation_df[column].isin(["不一致", "未提取到PDF披露值"])

    return validation_df.loc[mask].copy()


def _detect_report_name(balance_df: pd.DataFrame, default_name: str = "资产负债表") -> str:
    if "报表名称" not in balance_df.columns or balance_df.empty:
        return default_name

    report_names = balance_df["报表名称"].dropna().astype(str)
    return report_names.iloc[0] if not report_names.empty else default_name


def _format_balance_result(balanced: Any, diff: Any) -> str:
    if diff is None:
        return "未完成检查：未找到资产总计或负债和股东权益总计。"

    status = "平衡" if bool(balanced) else "不平衡"
    return f"{status}，差额：{float(diff):,.2f} 千元"


def _format_result_sheet(ws: Worksheet, balance_df: pd.DataFrame) -> None:
    _format_table_sheet(ws, balance_df)


def _format_summary_sheet(ws: Worksheet, summary_df: pd.DataFrame) -> None:
    _format_table_sheet(ws, summary_df)


def _format_table_sheet(ws: Worksheet, df: pd.DataFrame) -> None:
    _bold_header(ws)
    ws.freeze_panes = "A2"

    if ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = ws.dimensions

    _format_amount_columns(ws, df)
    _format_percentage_columns(ws, df)
    _auto_fit_columns(ws)


def _bold_header(ws: Worksheet) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True)


def _format_amount_columns(ws: Worksheet, balance_df: pd.DataFrame) -> None:
    amount_column_indexes = [
        column_index + 1
        for column_index, column in enumerate(balance_df.columns)
        if column in AMOUNT_COLUMNS or "金额" in str(column)
    ]

    for column_index in amount_column_indexes:
        for row_index in range(2, ws.max_row + 1):
            ws.cell(row=row_index, column=column_index).number_format = (
                AMOUNT_NUMBER_FORMAT
            )


def _format_percentage_columns(ws: Worksheet, df: pd.DataFrame) -> None:
    headers = {
        column_index + 1: str(column)
        for column_index, column in enumerate(df.columns)
    }
    code_column = next((index for index, header in headers.items() if header == "指标编码"), None)
    item_column = next((index for index, header in headers.items() if header == "指标名称"), None)
    field_column = next((index for index, header in headers.items() if header == "指标项"), None)

    for row_index in range(2, ws.max_row + 1):
        row_text = " ".join(
            str(ws.cell(row=row_index, column=column_index).value or "")
            for column_index in (code_column, item_column, field_column)
            if column_index
        )
        is_ratio_row = (
            "B0492" in row_text
            or "B0493" in row_text
            or any(f"B076{index}" in row_text for index in range(5))
            or "\u6bd4\u4f8b" in row_text
            or "\u6298\u73b0\u7387" in row_text
            or "\u589e\u957f\u7387" in row_text
        )
        for column_index, header in headers.items():
            if not (is_ratio_row or "比例" in header):
                continue
            if not _is_percentage_value_column(header):
                continue
            cell = ws.cell(row=row_index, column=column_index)
            if isinstance(cell.value, (int, float)) and -1 <= float(cell.value) <= 1:
                cell.number_format = "0.0%"


def _is_percentage_value_column(header: str) -> bool:
    if "\u9875\u7801" in header or "\u539f\u59cb\u884c\u6587\u672c" in header:
        return False
    if "\u6bd4\u4f8b" in header:
        return True
    if any(token in header for token in ("\u672c\u671f", "\u4e0a\u671f", "\u672c\u5e74", "\u4e0a\u5e74", "\u751f\u6210\u91d1\u989d", "\u91d1\u989d")):
        return True
    return header in {"\u751f\u6210\u503c", "PDF\u62ab\u9732\u503c", "\u5dee\u5f02\u503c", "PDF\u6307\u6807\u503c"} or header.startswith("PDF\u96c6\u56e2") or header.startswith("PDF\u672c\u884c")

def _auto_fit_columns(ws: Worksheet) -> None:
    for column_cells in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)

        for cell in column_cells:
            if cell.value is None:
                continue
            cell_length = _display_width(str(cell.value))
            max_length = max(max_length, cell_length)

        ws.column_dimensions[column_letter].width = min(max(max_length + 2, 10), 60)


def _display_width(value: str) -> int:
    """Estimate Excel column width, counting CJK characters as wider text."""
    width = 0
    for char in value:
        width += 2 if ord(char) > 127 else 1
    return width
