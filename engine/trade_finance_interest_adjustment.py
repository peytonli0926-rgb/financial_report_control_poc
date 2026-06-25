from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook


TRADE_PREFIX = "5-6-3-20251231"
IMPAIRMENT_PREFIX = "5-7-3-2-20251231"
TARGET_COLUMN = "利息调整"
TARGET_BIZ_TYPE = "01"
TARGET_BUSI_TYPES = {"01002", "01003", "01004", "01005"}
TRADE_BALANCE_COLUMNS = ("人民币余额", "人民币金额")
TRADE_INTEREST_ADJUSTMENT_COLUMNS = (
    TARGET_COLUMN,
    "利息调整（科目14011201）",
    "利息调整(科目14011201)",
)
IMPAIRMENT_REQUIRED_COLUMNS = {"BIZ_TYPE", "BUSI_TYPE", "CURRENT_BALNC", "START_DT", "DUBIL_MATR_DT"}


@dataclass(frozen=True)
class InterestAdjustmentSyncResult:
    trade_file: Path
    impairment_file: Path
    matched_rows: int
    unmatched_rows: int
    source_rows: int
    target_rows: int
    cleared_rows: int


def sync_trade_finance_interest_adjustment(
    trade_file: str | Path,
    impairment_file: str | Path,
) -> InterestAdjustmentSyncResult:
    trade_path = Path(trade_file)
    impairment_path = Path(impairment_file)

    source_rows = _load_trade_rows(trade_path)
    source_by_key: dict[tuple[str, str, str], list[Any]] = {}
    for row in source_rows:
        source_by_key.setdefault(row["key"], []).append(row["interest_adjustment"])

    sheets = pd.read_excel(impairment_path, sheet_name=None, dtype=object, engine="openpyxl")
    target_sheet_name, target_df = _find_dataframe_sheet_with_columns(
        sheets,
        IMPAIRMENT_REQUIRED_COLUMNS,
        preferred_sheet="20251231",
    )
    target_df = target_df.copy()
    if TARGET_COLUMN not in target_df.columns:
        target_df[TARGET_COLUMN] = 0

    non_blank_mask = ~target_df.drop(columns=[TARGET_COLUMN], errors="ignore").isna().all(axis=1)
    target_df.loc[non_blank_mask, TARGET_COLUMN] = 0
    cleared_rows = int(non_blank_mask.sum())
    matched_rows = 0
    target_rows = 0
    target_mask = non_blank_mask & target_df.apply(
        lambda row: _is_target_impairment_row(row.get("BIZ_TYPE"), row.get("BUSI_TYPE")),
        axis=1,
    )

    for index, row in target_df.loc[target_mask].iterrows():
        target_rows += 1
        key = (
            _normalize_amount(row.get("CURRENT_BALNC")),
            _normalize_date(row.get("START_DT")),
            _normalize_date(row.get("DUBIL_MATR_DT")),
        )
        candidates = source_by_key.get(key) or []
        if candidates:
            interest_adjustment = candidates.pop(0)
            if interest_adjustment not in (None, "", 0, 0.0, "0"):
                target_df.at[index, TARGET_COLUMN] = interest_adjustment
            matched_rows += 1

    sheets[target_sheet_name] = target_df
    _write_workbook(impairment_path, sheets)
    return InterestAdjustmentSyncResult(
        trade_file=trade_path,
        impairment_file=impairment_path,
        matched_rows=matched_rows,
        unmatched_rows=max(target_rows - matched_rows, 0),
        source_rows=len(source_rows),
        target_rows=target_rows,
        cleared_rows=cleared_rows,
    )


def find_latest_impairment_file(upload_dir: str | Path) -> Path | None:
    upload_path = Path(upload_dir)
    candidates = sorted(
        [
            path
            for path in upload_path.glob(f"{IMPAIRMENT_PREFIX}*.xlsx")
            if path.is_file()
        ],
        key=lambda path: ("减值明细" in path.name, path.stat().st_mtime),
        reverse=True,
    )
    for path in candidates:
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
            worksheet = workbook["20251231"] if "20251231" in workbook.sheetnames else workbook.worksheets[0]
            headers = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
            normalized_headers = {_normalize_column(header) for header in headers}
            if {_normalize_column(column) for column in IMPAIRMENT_REQUIRED_COLUMNS}.issubset(normalized_headers):
                return path
        except Exception:
            continue
    return None


def _load_trade_rows(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, data_only=True)
    worksheet, header_row, header_map = _find_sheet_with_header_aliases(
        workbook,
        [
            TRADE_BALANCE_COLUMNS,
            ("借据借款日期",),
            ("借据到期日",),
            TRADE_INTEREST_ADJUSTMENT_COLUMNS,
        ],
    )
    balance_col = _column_from_aliases(header_map, TRADE_BALANCE_COLUMNS)
    start_col = header_map[_normalize_column("借据借款日期")]
    maturity_col = header_map[_normalize_column("借据到期日")]
    interest_col = _column_from_aliases(header_map, TRADE_INTEREST_ADJUSTMENT_COLUMNS)

    rows: list[dict[str, Any]] = []
    for excel_row in range(header_row + 1, worksheet.max_row + 1):
        if _is_blank_row(worksheet, excel_row):
            continue
        key = (
            _normalize_amount(worksheet.cell(excel_row, balance_col).value),
            _normalize_date(worksheet.cell(excel_row, start_col).value),
            _normalize_date(worksheet.cell(excel_row, maturity_col).value),
        )
        rows.append(
            {
                "key": key,
                "interest_adjustment": worksheet.cell(excel_row, interest_col).value,
            }
        )
    return rows


def _find_sheet_with_headers(workbook, required_columns: set[str], preferred_sheet: str | None = None):
    return _find_sheet_with_header_aliases(
        workbook,
        [(column,) for column in required_columns],
        preferred_sheet=preferred_sheet,
    )


def _find_dataframe_sheet_with_columns(
    sheets: dict[str, pd.DataFrame],
    required_columns: set[str],
    preferred_sheet: str | None = None,
) -> tuple[str, pd.DataFrame]:
    sheet_names = []
    if preferred_sheet and preferred_sheet in sheets:
        sheet_names.append(preferred_sheet)
    sheet_names.extend(name for name in sheets if name not in sheet_names)

    required = {_normalize_column(column) for column in required_columns}
    for sheet_name in sheet_names:
        df = sheets[sheet_name]
        column_map = {_normalize_column(column): column for column in df.columns}
        if required.issubset(column_map):
            return sheet_name, df.rename(columns={column_map[name]: name for name in required})
    raise ValueError(f"未找到包含列 {', '.join(sorted(required_columns))} 的工作表")


def _write_workbook(path: Path, sheets: dict[str, pd.DataFrame]) -> None:
    temp_path = path.with_name(f"{path.stem}.tmp{path.suffix}")
    with pd.ExcelWriter(temp_path, engine="openpyxl") as writer:
        for sheet_name, df in sheets.items():
            df.to_excel(writer, sheet_name=sheet_name, index=False)
    temp_path.replace(path)


def _find_sheet_with_header_aliases(
    workbook,
    required_column_aliases: list[tuple[str, ...]],
    preferred_sheet: str | None = None,
):
    worksheets = []
    if preferred_sheet and preferred_sheet in workbook.sheetnames:
        worksheets.append(workbook[preferred_sheet])
    worksheets.extend(sheet for sheet in workbook.worksheets if sheet not in worksheets)

    for worksheet in worksheets:
        for row in worksheet.iter_rows(min_row=1, max_row=min(30, worksheet.max_row)):
            header_map = {
                _normalize_column(cell.value): cell.column
                for cell in row
                if cell.value is not None and str(cell.value).strip()
            }
            if all(_column_from_aliases(header_map, aliases, raise_missing=False) for aliases in required_column_aliases):
                return worksheet, row[0].row, header_map
    required_text = ", ".join("/".join(aliases) for aliases in required_column_aliases)
    raise ValueError(f"未找到包含列 {required_text} 的工作表")


def _column_from_aliases(header_map: dict[str, int], aliases: tuple[str, ...], raise_missing: bool = True) -> int | None:
    for alias in aliases:
        column = header_map.get(_normalize_column(alias))
        if column:
            return column
    if raise_missing:
        raise ValueError(f"未找到列：{'/'.join(aliases)}")
    return None


def _ensure_column(worksheet, header_map: dict[str, int], header_row: int, column_name: str) -> int:
    normalized = _normalize_column(column_name)
    if normalized in header_map:
        return header_map[normalized]
    next_column = worksheet.max_column + 1
    worksheet.cell(header_row, next_column, value=column_name)
    header_map[normalized] = next_column
    return next_column


def _normalize_column(value: Any) -> str:
    return "".join(str(value or "").split())


def _normalize_amount(value: Any) -> str:
    if value is None:
        return "0.00"
    text = str(value).replace(",", "").strip()
    if not text or text.lower() == "nan" or text == "-":
        return "0.00"
    try:
        return str(Decimal(text).quantize(Decimal("0.01")))
    except (InvalidOperation, ValueError):
        return text


def _normalize_date(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip()
    if not text or text.lower() == "nan":
        return ""
    if text.endswith(".0"):
        text = text[:-2]
    if len(text) == 8 and text.isdigit():
        return f"{text[:4]}-{text[4:6]}-{text[6:]}"
    return text.replace("/", "-")


def _is_target_impairment_row(biz_type: Any, busi_type: Any) -> bool:
    return (
        _normalize_code(biz_type, width=2) == TARGET_BIZ_TYPE
        and _normalize_code(busi_type, width=5) in TARGET_BUSI_TYPES
    )


def _normalize_code(value: Any, width: int) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if not text or text.lower() == "nan":
        return ""
    if text.endswith(".0"):
        text = text[:-2]
    return text.zfill(width) if text.isdigit() else text


def _is_blank_row(worksheet, row_number: int) -> bool:
    return all(
        cell.value is None or str(cell.value).strip() == ""
        for cell in worksheet[row_number]
    )
