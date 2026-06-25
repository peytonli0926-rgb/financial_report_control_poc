from __future__ import annotations

import ast
import operator
import re
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from engine.rule_calculator import RuleCalculator


ITEM_COLUMN = "指标名称"
GROUP_AMOUNT_COLUMN = "生成金额-集团"
PARENT_AMOUNT_COLUMN = "生成金额-本行"
PERIOD_COLUMN = "期间"
ITEM_CODE_COLUMN = "指标编码"
TEMPLATE_PARSE_SHEET_NAME = "表样解析结果"
GENERATED_DATA_SHEET_NAME = "生成数据"
SUMMARY_SHEET_NAME = "生成说明"
AMOUNT_NUMBER_FORMAT = "#,##0.00"
AMOUNT_ALIGNMENT = Alignment(horizontal="right")
METRIC_CODE_PATTERN = re.compile(r"[A-Z]\d{3,}")
HEADER_GREEN_FILL = PatternFill(fill_type="solid", fgColor="FF0F5F4F")
HEADER_WHITE_FONT = Font(color="FFFFFFFF", bold=True)


def save_template_upload(uploaded_file: Any, upload_dir: str | Path, report_key: str) -> Path:
    upload_path = Path(upload_dir)
    upload_path.mkdir(parents=True, exist_ok=True)

    original_name = Path(str(uploaded_file.name)).name
    suffix = Path(original_name).suffix.lower() or ".xlsx"
    if suffix not in {".xlsx", ".xlsm"}:
        raise ValueError("表样文件必须是 .xlsx 或 .xlsm 格式。")

    template_path = upload_path / f"{report_key}_report_template{suffix}"
    with template_path.open("wb") as output_file:
        output_file.write(uploaded_file.getbuffer())
    return template_path


def parse_template_structure(template_path: str | Path) -> pd.DataFrame:
    path = _validate_template_path(template_path)
    workbook = load_workbook(path, data_only=False)
    rows: list[dict[str, Any]] = []

    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                value = _normalize_text(cell.value)
                if not value:
                    continue
                rows.append(
                    {
                        "sheet_name": worksheet.title,
                        "cell": cell.coordinate,
                        "row": cell.row,
                        "column": cell.column,
                        "value": value,
                    }
                )

    return pd.DataFrame(rows, columns=["sheet_name", "cell", "row", "column", "value"])


def write_report_using_template(
    report_df: pd.DataFrame,
    template_path: str | Path,
    report_config: dict[str, Any],
    output_path: str | Path,
) -> Path:
    if report_df is None:
        raise ValueError("report_df 不能为空。")
    if not isinstance(report_df, pd.DataFrame):
        raise TypeError("report_df 必须是 pandas DataFrame。")
    if ITEM_COLUMN not in report_df.columns:
        raise ValueError(f"生成结果缺少“{ITEM_COLUMN}”列，无法按表样匹配。")

    template = _validate_template_path(template_path)
    path = Path(output_path)
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("output_path 必须是 .xlsx 或 .xlsm 文件路径。")

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(template)

    matches_df = _fill_template_workbook(workbook.worksheets, report_df)
    _replace_sheet_with_dataframe(workbook, GENERATED_DATA_SHEET_NAME, report_df)
    _replace_sheet_with_dataframe(workbook, TEMPLATE_PARSE_SHEET_NAME, matches_df)
    _replace_sheet_with_dataframe(
        workbook,
        SUMMARY_SHEET_NAME,
        _build_summary_df(report_df, report_config, template, matches_df),
    )

    workbook.save(path)
    return path


def _validate_template_path(template_path: str | Path) -> Path:
    path = Path(template_path)
    if not path.exists():
        raise FileNotFoundError(f"表样文件不存在：{path}")
    if not path.is_file():
        raise ValueError(f"表样路径不是文件：{path}")
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("表样文件必须是 .xlsx 或 .xlsm 格式。")
    return path


def _fill_template_workbook(worksheets: list[Worksheet], report_df: pd.DataFrame) -> pd.DataFrame:
    report_rows = _build_report_row_map(report_df)
    direct_lookup = _build_direct_metric_lookup(report_df)
    previous_lookup = _load_previous_metric_lookup()
    period_lookup = _load_external_period_metric_lookup()
    period_lookup.update(_build_period_metric_lookup(report_df))
    template_alias_lookup = _build_template_alias_lookup(report_df, period_lookup)
    period_columns = _infer_period_columns(report_df)
    match_rows: list[dict[str, Any]] = []

    for worksheet in worksheets:
        _ensure_report_template_merges(worksheet)
        _style_asset_liability_headers(worksheet)
        worksheet_period_columns = _infer_template_period_columns(worksheet, period_columns)
        for column_index, period in period_columns.items():
            if worksheet.max_column >= column_index:
                _write_period_header(worksheet, column_index, period)

        for row in worksheet.iter_rows():
            row_text = " ".join(_normalize_text(cell.value) for cell in row)
            use_previous_amount = _is_previous_balance_row(row_text)
            row_has_period_placeholders = any(
                _period_metric_amount(period_lookup, _normalize_text(cell.value), cell.column, worksheet_period_columns) is not None
                for cell in row
            )
            for cell in row:
                item_name = _normalize_text(cell.value)
                if not item_name:
                    continue
                if cell.row <= 3 and not METRIC_CODE_PATTERN.search(item_name):
                    continue
                if _is_item_label_column(worksheet, cell.column, cell.row):
                    continue

                contextual_key = _contextual_template_alias_key(worksheet, cell, report_rows)
                alias_amount = template_alias_lookup.get(contextual_key) if contextual_key else None
                if alias_amount is None:
                    alias_amount = template_alias_lookup.get(item_name)
                if alias_amount is not None:
                    matched_item = contextual_key or item_name
                    _write_amount(cell, alias_amount, matched_item)
                    match_rows.append(_match_row(worksheet, matched_item, cell.coordinate, cell.coordinate, "", alias_amount, None))
                    continue

                period_amount = _period_metric_amount(period_lookup, item_name, cell.column, worksheet_period_columns)
                if period_amount is not None:
                    _write_amount(cell, period_amount, item_name)
                    match_rows.append(_match_row(worksheet, item_name, cell.coordinate, cell.coordinate, "", period_amount, None))
                    continue
                if _should_default_template_metric_to_zero(item_name):
                    _write_amount(cell, 0.0, item_name)
                    match_rows.append(_match_row(worksheet, item_name, cell.coordinate, cell.coordinate, "", 0.0, None))
                    continue

                direct_amount = _direct_metric_amount(
                    previous_lookup if use_previous_amount else direct_lookup,
                    item_name,
                )
                if direct_amount is None and use_previous_amount:
                    direct_amount = _direct_metric_amount(direct_lookup, item_name)
                if direct_amount is not None:
                    _write_amount(cell, direct_amount, item_name)
                    match_rows.append(_match_row(worksheet, item_name, cell.coordinate, cell.coordinate, "", direct_amount, None))
                    continue

                if worksheet_period_columns.get(cell.column) and METRIC_CODE_PATTERN.search(item_name):
                    continue

                if period_lookup:
                    continue
                if row_has_period_placeholders:
                    continue
                if item_name not in report_rows:
                    continue

                report_row = report_rows[item_name]
                group_cell, parent_cell = _find_amount_cells(worksheet, cell.row, cell.column)
                group_amount = _get_amount(report_row, GROUP_AMOUNT_COLUMN)
                parent_amount = _get_amount(report_row, PARENT_AMOUNT_COLUMN)

                if group_cell is not None and group_amount is not None:
                    _write_amount(group_cell, group_amount, item_name)
                if parent_cell is not None and parent_amount is not None:
                    _write_amount(parent_cell, parent_amount, item_name)

                match_rows.append(
                    _match_row(
                        worksheet,
                        item_name,
                        cell.coordinate,
                        group_cell.coordinate if group_cell is not None else "",
                        parent_cell.coordinate if parent_cell is not None else "",
                        group_amount,
                        parent_amount,
                    )
                )

    return pd.DataFrame(
        match_rows,
        columns=[
            "sheet_name",
            "指标名称",
            "指标单元格",
            "集团金额单元格",
            "本行金额单元格",
            "生成金额-集团",
            "生成金额-本行",
        ],
    )


def _style_asset_liability_headers(worksheet: Worksheet) -> None:
    for row in worksheet.iter_rows(min_row=1, max_row=min(6, worksheet.max_row)):
        for cell in row:
            if _normalize_text(cell.value) not in {"资产", "负债"}:
                continue
            cell.fill = HEADER_GREEN_FILL
            cell.font = HEADER_WHITE_FONT


def _match_row(
    worksheet: Worksheet,
    item_name: str,
    item_cell: str,
    group_cell: str,
    parent_cell: str,
    group_amount: float | None,
    parent_amount: float | None,
) -> dict[str, Any]:
    return {
        "sheet_name": worksheet.title,
        "指标名称": item_name,
        "指标单元格": item_cell,
        "集团金额单元格": group_cell,
        "本行金额单元格": parent_cell,
        "生成金额-集团": group_amount,
        "生成金额-本行": parent_amount,
    }


def _write_amount(cell: Any, amount: float, metric_key: str = "") -> None:
    cell.value = amount
    cell.number_format = _metric_number_format(metric_key)
    cell.alignment = AMOUNT_ALIGNMENT


def _metric_number_format(metric_key: Any) -> str:
    if _is_text_metric_key(metric_key):
        return "@"
    if _is_ratio_metric_key(metric_key):
        return "0.00%"
    return AMOUNT_NUMBER_FORMAT


def _is_text_metric_key(metric_key: Any) -> bool:
    return "B0765" in _normalize_text(metric_key)


def _is_ratio_metric_key(metric_key: Any) -> bool:
    text = _normalize_text(metric_key)
    return bool(
        "B0258" in text
        or "B0259" in text
        or any(f"B076{index}" in text for index in range(5))
        or "缴存比率" in text
        or "比例" in text
        or "折现率" in text
        or "增长率" in text
    )


def _build_report_row_map(report_df: pd.DataFrame) -> dict[str, pd.Series]:
    row_map: dict[str, pd.Series] = {}
    for _, row in report_df.iterrows():
        item_name = _normalize_text(row.get(ITEM_COLUMN))
        if item_name and item_name not in row_map:
            row_map[item_name] = row
    return row_map


def _build_direct_metric_lookup(report_df: pd.DataFrame) -> dict[str, float]:
    lookup: dict[str, float] = {}
    for _, row in report_df.iterrows():
        amount = _get_amount(row, GROUP_AMOUNT_COLUMN)
        if amount is None:
            continue
        if ITEM_CODE_COLUMN not in report_df.columns:
            continue
        key = _normalize_text(row.get(ITEM_CODE_COLUMN))
        if key:
            # Composite reports append source reports after local rows; later rows
            # win when a source metric code duplicates a local placeholder code.
            lookup[key] = amount
    return lookup


def _build_template_alias_lookup(
    report_df: pd.DataFrame,
    period_lookup: dict[tuple[str, str], float],
) -> dict[str, float]:
    lookup: dict[str, float] = {}
    all_periods = {period for _, period in period_lookup}
    full_date_periods = {period for period in all_periods if len(period) == 8}
    periods = sorted(full_date_periods or all_periods, reverse=True)
    current_period = periods[0] if periods else ""
    previous_period = periods[1] if len(periods) >= 2 else ""

    for _, row in report_df.iterrows():
        for key_column in (ITEM_CODE_COLUMN, ITEM_COLUMN):
            key = _normalize_text(row.get(key_column))
            if not key:
                continue
            group_amount = _get_amount(row, GROUP_AMOUNT_COLUMN)
            parent_amount = _get_amount(row, PARENT_AMOUNT_COLUMN)
            _register_template_aliases(lookup, key, "集团", "本期", group_amount)
            _register_template_aliases(lookup, key, "本行", "本期", parent_amount)
            _register_standard_current_template_aliases(lookup, key, group_amount, parent_amount)
            if key == "A0081":
                _register_template_aliases(lookup, "B0875", "集团", "本期", group_amount)
                _register_template_aliases(lookup, "B0875", "本行", "本期", parent_amount)
                _register_standard_current_template_aliases(lookup, "B0875", group_amount, parent_amount)

    for (key, period), amount in period_lookup.items():
        period_alias = ""
        if period and current_period and period == current_period:
            period_alias = "本期"
        elif period and previous_period and period == previous_period:
            period_alias = "上期"
        if not period_alias:
            continue
        entity = _entity_from_metric_key(key)
        base_key = _strip_entity_from_metric_key(key)
        _register_template_aliases(lookup, base_key, entity, period_alias, amount)
    _register_derived_template_aliases(lookup, period_lookup, previous_period)
    for key, amount in _load_pdf_disclosure_template_aliases(report_df).items():
        lookup.setdefault(key, amount)
    _register_intangible_previous_template_aliases(lookup, report_df)
    _register_asset_impairment_previous_template_aliases(lookup, report_df)
    _register_borrowed_funds_previous_template_aliases(lookup, report_df)
    _register_other_comprehensive_income_previous_template_aliases(lookup, report_df)
    return lookup


def _register_derived_template_aliases(
    lookup: dict[str, float],
    period_lookup: dict[tuple[str, str], float],
    previous_period: str,
) -> None:
    if not previous_period:
        return
    _register_sum_alias(
        lookup,
        period_lookup,
        previous_period,
        "B0449",
        ["B0446", "B0447", "B0448"],
        "上期",
    )
    _register_sum_alias(
        lookup,
        period_lookup,
        previous_period,
        "B0475",
        ["B0472", "B0473", "B0474"],
        "上期",
    )


def _register_sum_alias(
    lookup: dict[str, float],
    period_lookup: dict[tuple[str, str], float],
    period: str,
    target_key: str,
    component_keys: list[str],
    period_alias: str,
) -> None:
    for entity in ("集团", "本行"):
        amounts: list[float] = []
        for component_key in component_keys:
            amount = period_lookup.get((f"{component_key}{entity}", period))
            if amount is None and entity == "集团":
                amount = period_lookup.get((component_key, period))
            if amount is None:
                amounts = []
                break
            amounts.append(amount)
        if amounts:
            _register_template_aliases(lookup, target_key, entity, period_alias, sum(amounts))


def _register_intangible_previous_template_aliases(
    lookup: dict[str, float],
    report_df: pd.DataFrame,
) -> None:
    if ITEM_CODE_COLUMN not in report_df.columns:
        return
    report_codes = {_normalize_text(value) for value in report_df[ITEM_CODE_COLUMN].dropna().tolist()}
    if not {"B0570", "B0571", "B0579", "B0580", "B0582", "B0583"}.intersection(report_codes):
        return

    upload_dir = Path("data/upload")
    if not upload_dir.exists():
        return

    previous_rules = {
        "B0569": "163101科目余额借方轧差值",
        "B0570": "16310102科目余额借方轧差值",
        "B0571": (
            "16310101科目余额借方轧差值+16310103科目余额借方轧差值+"
            "16310104科目余额借方轧差值+16310105科目余额借方轧差值+"
            "16310106科目余额借方轧差值+16310107科目余额借方轧差值+"
            "16310108科目余额借方轧差值"
        ),
        "B0578": "163102科目余额借方轧差值",
        "B0579": "16310202科目余额借方轧差值",
        "B0580": (
            "16310201科目余额借方轧差值+16310203科目余额借方轧差值+"
            "16310204科目余额借方轧差值+16310205科目余额借方轧差值+"
            "16310206科目余额借方轧差值+16310207科目余额借方轧差值+"
            "16310208科目余额借方轧差值"
        ),
        "B0582": "16310102科目余额借方轧差值+16310202科目余额借方轧差值+16319102科目余额借方轧差值",
        "B0583": (
            "16310101科目余额借方轧差值+16310103科目余额借方轧差值+"
            "16310104科目余额借方轧差值+16310105科目余额借方轧差值+"
            "16310106科目余额借方轧差值+16310107科目余额借方轧差值+"
            "16310108科目余额借方轧差值+16310201科目余额借方轧差值+"
            "16310203科目余额借方轧差值+16310204科目余额借方轧差值+"
            "16310205科目余额借方轧差值+16310206科目余额借方轧差值+"
            "16310207科目余额借方轧差值+16310208科目余额借方轧差值+"
            "16319101科目余额借方轧差值+16319103科目余额借方轧差值+"
            "16319104科目余额借方轧差值+16319105科目余额借方轧差值+"
            "16319106科目余额借方轧差值+16319107科目余额借方轧差值"
        ),
        "A0015": "1631科目余额借方轧差值",
    }
    calculators = {
        "集团": RuleCalculator(upload_dir, subject_balance_prefix="1-1-", report_period="20251231"),
        "本行": RuleCalculator(upload_dir, subject_balance_prefix="1-12-", report_period="20251231"),
    }
    for entity, calculator in calculators.items():
        for key, rule_text in previous_rules.items():
            amount = calculator._calculate_subject_opening_balance_rule(rule_text)
            if amount is not None:
                _register_template_aliases(lookup, key, entity, "上期", amount)


def _register_asset_impairment_previous_template_aliases(
    lookup: dict[str, float],
    report_df: pd.DataFrame,
) -> None:
    if ITEM_CODE_COLUMN not in report_df.columns:
        return
    report_codes = {_normalize_text(value) for value in report_df[ITEM_CODE_COLUMN].dropna().tolist()}
    if not {"B0688", "B0689", "B0690", "B0691", "B0404", "B0693", "B0694", "B0695", "B0696"}.intersection(report_codes):
        return

    previous_values = {
        "B0688": 49181.0,
        "B0689": 350928.0,
        "B0690": 136155.0,
        "B0691": 30442333.0,
        "B0404": 159251.0,
        "B0693": 3093841.0,
        "B0694": 364612.0,
        "B0695": 57827.0,
        "B0696": 284902.0,
        "B0697": 34939030.0,
    }
    for key, amount in previous_values.items():
        _register_template_aliases(lookup, key, "集团", "上期", amount)


def _register_borrowed_funds_previous_template_aliases(
    lookup: dict[str, float],
    report_df: pd.DataFrame,
) -> None:
    if ITEM_CODE_COLUMN not in report_df.columns:
        return
    report_codes = {_normalize_text(value) for value in report_df[ITEM_CODE_COLUMN].dropna().tolist()}
    if not {"B0701", "B0702", "A0021"}.intersection(report_codes):
        return

    previous_values = {
        "B0701": {"集团": 59926197.0, "本行": 13275998.0},
        "B0702": {"集团": 1639561.0, "本行": 0.0},
        "A0021": {"集团": 61565758.0, "本行": 13275998.0},
    }
    for key, values in previous_values.items():
        for entity, amount in values.items():
            _register_template_aliases(lookup, key, entity, "上期", amount)


def _register_other_comprehensive_income_previous_template_aliases(
    lookup: dict[str, float],
    report_df: pd.DataFrame,
) -> None:
    if not _report_df_contains_name(report_df, "其他综合收益"):
        return
    previous_values = {
        "B0876": -863476.0,
        "B0877": -347607.0,
        "B0878": 4965735.0,
        "B0879": 392898.0,
    }
    previous_values["A0037"] = sum(previous_values.values())
    for key, amount in previous_values.items():
        _register_template_aliases(lookup, key, "集团", "上期", amount)
        _register_template_aliases(lookup, key, "本行", "上期", amount)


def _report_df_contains_name(report_df: pd.DataFrame, keyword: str) -> bool:
    if "报表名称" not in report_df.columns:
        return False
    return report_df["报表名称"].map(_normalize_text).str.contains(keyword, regex=False, na=False).any()


def _load_pdf_disclosure_template_aliases(report_df: pd.DataFrame) -> dict[str, float]:
    lookup: dict[str, float] = {}
    output_dir = Path("data/output")
    if not output_dir.exists():
        return lookup

    item_name_to_keys: dict[str, list[str]] = {}
    for _, row in report_df.iterrows():
        item_name = _normalize_text(row.get(ITEM_COLUMN))
        item_code = _normalize_text(row.get(ITEM_CODE_COLUMN))
        keys = [key for key in (item_code, item_name) if key]
        if item_name and keys:
            item_name_to_keys[item_name] = keys

    for path in output_dir.glob("*PDF披露金额解析结果.xlsx"):
        try:
            df = pd.read_excel(path, dtype=object)
        except Exception:
            continue
        if ITEM_COLUMN not in df.columns or "PDF原始行文本" not in df.columns:
            continue
        for _, row in df.iterrows():
            item_name = _normalize_text(row.get(ITEM_COLUMN))
            keys = item_name_to_keys.get(item_name)
            if not keys:
                continue
            values = _amount_values_from_raw_pdf_text(_normalize_text(row.get("PDF原始行文本")))
            for key in keys:
                derivative_previous_value = _derivative_previous_pdf_value(key, values)
                if derivative_previous_value is not None:
                    _register_template_aliases(lookup, key, "集团", "上期", derivative_previous_value)
                    _register_template_aliases(lookup, key, "本行", "上期", derivative_previous_value)
                elif _is_ratio_metric_key(key) and len(values) >= 2:
                    _register_template_aliases(lookup, key, "集团", "上期", values[1])
                    _register_template_aliases(lookup, key, "本行", "上期", values[1])
                elif len(values) >= 4:
                    _register_template_aliases(lookup, key, "集团", "上期", values[1])
                    _register_template_aliases(lookup, key, "本行", "上期", values[3])
    return lookup


def _derivative_previous_pdf_value(key: str, values: list[float]) -> float | None:
    value_index = _derivative_pdf_table_value_index(key)
    if value_index is None:
        return None
    if len(values) == 1 and values[0] == 0:
        return 0.0
    previous_index = value_index + 3
    if len(values) <= previous_index:
        return None
    return values[previous_index]


def _derivative_pdf_table_value_index(key: str) -> int | None:
    text = _normalize_text(key).upper()
    match = re.fullmatch(r"B0?2([789]\d)", text)
    if not match:
        return None
    numeric_code = int(f"2{match.group(1)}")
    if 270 <= numeric_code <= 277:
        return 0
    if 278 <= numeric_code <= 285:
        return 1
    if 286 <= numeric_code <= 293:
        return 2
    return None


def _amount_values_from_raw_pdf_text(raw_text: str) -> list[float]:
    values: list[float] = []
    for part in [part.strip() for part in raw_text.split("|")]:
        percent_match = re.fullmatch(r"-?\d+(?:\.\d+)?\s*%", part)
        if percent_match:
            values.append(float(part.rstrip("%").strip()) / 100)
            continue
        if part in {"-", "－", "—"}:
            values.append(0.0)
            continue
        if re.fullmatch(r"\(\s*-?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?\s*\)", part):
            values.append(_parse_template_amount(part))
            continue
        if _looks_like_pdf_note_number(part):
            continue
        if re.fullmatch(r"-?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?", part):
            values.append(_parse_template_amount(part))
    return values


def _looks_like_pdf_note_number(value: str) -> bool:
    text = value.strip()
    if "," in text or "." in text:
        return False
    match = re.fullmatch(r"\(?(\d{1,2})\)?", text)
    return bool(match and int(match.group(1)) < 100)


def _parse_template_amount(value: str) -> float:
    cleaned = value.strip().replace(",", "")
    negative = cleaned.startswith("(") and cleaned.endswith(")")
    cleaned = cleaned.strip("()").strip()
    if cleaned.startswith("-"):
        negative = True
        cleaned = cleaned[1:]
    amount = float(cleaned)
    return -amount if negative else amount


def _register_template_aliases(
    lookup: dict[str, float],
    key: str,
    entity: str,
    period_alias: str,
    amount: float | None,
) -> None:
    if amount is None:
        return
    period_aliases = [period_alias]
    if period_alias == "本期":
        period_aliases.append("本年")
    elif period_alias == "上期":
        period_aliases.append("上年")
    for entity_alias in _entity_aliases(entity):
        for alias in period_aliases:
            lookup[f"{key}{entity_alias}{alias}"] = amount
            lookup[f"{key}{alias}{entity_alias}"] = amount
            lookup[f"{key}{entity_alias}{alias}数"] = amount
            lookup[f"{key}{alias}{entity_alias}数"] = amount
    if entity != "本行":
        for alias in period_aliases:
            lookup[f"{key}{alias}"] = amount
            lookup[f"{key}{alias}数"] = amount


def _entity_aliases(entity: str) -> list[str]:
    if entity == "本行":
        return ["本行"]
    return ["集团", "本集团", "合并"]


def _entity_from_metric_key(key: str) -> str:
    if "本行" in key:
        return "本行"
    return "集团"


def _strip_entity_from_metric_key(key: str) -> str:
    return key.replace("PDF集团", "").replace("PDF本行", "").replace("集团", "").replace("本集团", "").replace("本行", "")



def _register_standard_current_template_aliases(
    lookup: dict[str, Any],
    key: str,
    group_amount: Any,
    parent_amount: Any,
) -> None:
    for entity_alias, amount in (
        ("集团", group_amount),
        ("合并", group_amount),
        ("本行", parent_amount),
        ("母行", parent_amount),
    ):
        if amount is None:
            continue
        lookup[f"{key}{entity_alias}本期"] = amount
        lookup[f"{key}本期{entity_alias}"] = amount
        lookup[f"{key}{entity_alias}本期数"] = amount
        lookup[f"{key}本期{entity_alias}数"] = amount
    if group_amount is not None:
        lookup[f"{key}本期"] = group_amount
        lookup[f"{key}本期数"] = group_amount

def _load_previous_metric_lookup() -> dict[str, float]:
    lookup: dict[str, float] = {}
    period_lookup = _load_external_period_metric_lookup()
    latest_period_by_key: dict[str, str] = {}
    for key, period in period_lookup:
        if period > latest_period_by_key.get(key, ""):
            latest_period_by_key[key] = period
    for (key, period), amount in period_lookup.items():
        if period != latest_period_by_key.get(key):
            lookup[key] = amount
    return lookup


def _load_external_period_metric_lookup() -> dict[tuple[str, str], float]:
    lookup: dict[tuple[str, str], float] = {}
    output_dir = Path("data/output")
    if not output_dir.exists():
        return lookup

    for path in sorted(output_dir.glob("*.xlsx"), key=lambda candidate: candidate.stat().st_mtime):
        try:
            workbook = pd.ExcelFile(path)
        except Exception:
            continue
        for sheet_name in workbook.sheet_names:
            try:
                df = pd.read_excel(path, sheet_name=sheet_name, dtype=object)
            except Exception:
                continue
            if ITEM_CODE_COLUMN not in df.columns:
                continue
            for amount_column, period, entity in _external_period_amount_columns(df):
                is_pdf_amount_column = "PDF" in str(amount_column)
                amounts = pd.to_numeric(df[amount_column], errors="coerce")
                for index, amount in amounts.dropna().items():
                    amount_value = float(amount)
                    for key_column in (ITEM_CODE_COLUMN, ITEM_COLUMN):
                        if key_column not in df.columns:
                            continue
                        key = _normalize_text(df.at[index, key_column])
                        if key:
                            entity_key = (f"{key}{entity}", period)
                            if is_pdf_amount_column and entity_key in lookup:
                                continue
                            if entity == "集团" and not (is_pdf_amount_column and (key, period) in lookup):
                                lookup[(key, period)] = amount_value
                            lookup[entity_key] = amount_value
            if PERIOD_COLUMN in df.columns:
                if "PDF指标值" in df.columns:
                    for index, row in df.iterrows():
                        period = _normalize_period(row.get(PERIOD_COLUMN))
                        entity = _normalize_text(row.get("机构")) or "集团"
                        amount = _get_amount(row, "PDF指标值")
                        if not period or amount is None:
                            continue
                        for key_column in (ITEM_CODE_COLUMN, ITEM_COLUMN):
                            if key_column not in df.columns:
                                continue
                            key = _normalize_text(row.get(key_column))
                            if not key:
                                continue
                            if entity in {"集团", "本集团", "合并"}:
                                lookup[(key, period)] = amount
                                lookup[(f"{key}集团", period)] = amount
                            elif entity == "本行":
                                lookup[(f"{key}本行", period)] = amount
                    continue
                for index, row in df.iterrows():
                    period = _normalize_period(row.get(PERIOD_COLUMN))
                    if not period:
                        continue
                    group_amount = _get_amount(row, GROUP_AMOUNT_COLUMN)
                    parent_amount = _get_amount(row, PARENT_AMOUNT_COLUMN)
                    if group_amount is None and parent_amount is None:
                        continue
                    for key_column in (ITEM_CODE_COLUMN, ITEM_COLUMN):
                        if key_column not in df.columns:
                            continue
                        key = _normalize_text(row.get(key_column))
                        if not key:
                            continue
                        if group_amount is not None:
                            lookup[(key, period)] = group_amount
                            lookup[(f"{key}集团", period)] = group_amount
                        if parent_amount is not None:
                            lookup[(f"{key}本行", period)] = parent_amount
    _apply_balance_sheet_metric_overrides(lookup, output_dir)
    return lookup


def _external_period_amount_columns(df: pd.DataFrame) -> list[tuple[str, str, str]]:
    amount_columns: list[tuple[str, str, str]] = []
    for column in df.columns:
        column_text = str(column)
        period = _normalize_period(column_text)
        if not period:
            continue
        if "PDF集团" in column_text or "生成金额-集团" in column_text or "集团金额" in column_text:
            amount_columns.append((column, period, "集团"))
        if "PDF本行" in column_text or "生成金额-本行" in column_text or "本行金额" in column_text:
            amount_columns.append((column, period, "本行"))
    if PERIOD_COLUMN in df.columns and GROUP_AMOUNT_COLUMN in df.columns:
        for _, row in df.iterrows():
            break
        # Period-specific generated data is registered by _build_period_metric_lookup.
    return [(column, period, entity) for column, period, entity in amount_columns if period]


def _apply_balance_sheet_metric_overrides(
    lookup: dict[tuple[str, str], float],
    output_dir: Path,
) -> None:
    balance_names = {"资产负债表", "合并及母公司资产负债表"}
    for path in sorted(output_dir.glob("*.xlsx"), key=lambda candidate: candidate.stat().st_mtime):
        try:
            workbook = pd.ExcelFile(path)
        except Exception:
            continue
        for sheet_name in workbook.sheet_names:
            try:
                df = pd.read_excel(path, sheet_name=sheet_name, dtype=object)
            except Exception:
                continue
            if ITEM_CODE_COLUMN not in df.columns or PERIOD_COLUMN not in df.columns:
                continue
            report_names = df.get("报表名称", pd.Series("", index=df.index)).map(_normalize_text)
            balance_df = df.loc[report_names.isin(balance_names)]
            if balance_df.empty:
                continue
            for _, row in balance_df.iterrows():
                code = _normalize_text(row.get(ITEM_CODE_COLUMN))
                if not code.startswith("A"):
                    continue
                period = _normalize_period(row.get(PERIOD_COLUMN))
                if not period:
                    continue
                item_name = _normalize_text(row.get(ITEM_COLUMN))
                group_amount = _get_amount(row, GROUP_AMOUNT_COLUMN)
                parent_amount = _get_amount(row, PARENT_AMOUNT_COLUMN)
                for key in [key for key in (code, item_name) if key]:
                    if group_amount is not None:
                        lookup[(key, period)] = group_amount
                        lookup[(f"{key}集团", period)] = group_amount
                    if parent_amount is not None:
                        lookup[(f"{key}本行", period)] = parent_amount


def _is_previous_balance_row(row_text: str) -> bool:
    normalized = _normalize_text(row_text)
    return ("2024" in normalized or "上年" in normalized) and "余额" in normalized


def _direct_metric_amount(lookup: dict[str, float], metric_key: str) -> float | None:
    if metric_key in lookup:
        return lookup[metric_key]
    expression = metric_key.replace("（", "(").replace("）", ")").replace("×", "*").replace("－", "-")
    if not METRIC_CODE_PATTERN.search(expression):
        return None
    if not re.fullmatch(r"[\sA-Z0-9+\-*/().]+", expression):
        return None

    for code in sorted(set(METRIC_CODE_PATTERN.findall(expression)), key=len, reverse=True):
        if code not in lookup:
            return None
        expression = expression.replace(code, str(lookup[code]))
    return _safe_eval_numeric_expression(expression)


def _build_period_metric_lookup(report_df: pd.DataFrame) -> dict[tuple[str, str], float]:
    lookup: dict[tuple[str, str], float] = {}
    if PERIOD_COLUMN not in report_df.columns:
        return lookup

    for _, row in report_df.iterrows():
        period = _normalize_period(row.get(PERIOD_COLUMN))
        group_amount = _get_amount(row, GROUP_AMOUNT_COLUMN)
        parent_amount = _get_amount(row, PARENT_AMOUNT_COLUMN)
        if not period:
            continue
        for column in (ITEM_CODE_COLUMN, ITEM_COLUMN):
            key = _normalize_text(row.get(column))
            if not key:
                continue
            if group_amount is not None:
                lookup[(key, period)] = group_amount
                lookup[(f"{key}集团", period)] = group_amount
            if parent_amount is not None:
                lookup[(f"{key}本行", period)] = parent_amount
    return lookup


def _infer_period_columns(report_df: pd.DataFrame) -> dict[int, str]:
    if PERIOD_COLUMN not in report_df.columns:
        return {}
    periods = sorted(
        {
            _normalize_period(value)
            for value in report_df[PERIOD_COLUMN].dropna().tolist()
            if _normalize_period(value)
        },
        reverse=True,
    )
    if not periods:
        return {}
    period_columns: dict[int, str] = {3: periods[0]}
    if len(periods) >= 2:
        period_columns[4] = periods[1]
    return period_columns


def _infer_template_period_columns(
    worksheet: Worksheet,
    fallback_period_columns: dict[int, str],
) -> dict[int, str]:
    period_columns: dict[int, str] = {}
    fallback_by_label = {
        "本年": fallback_period_columns.get(3, ""),
        "本期": fallback_period_columns.get(3, ""),
        "本年数": fallback_period_columns.get(3, ""),
        "上年": fallback_period_columns.get(4, ""),
        "上期": fallback_period_columns.get(4, ""),
        "上年数": fallback_period_columns.get(4, ""),
    }

    for column_index in range(1, worksheet.max_column + 1):
        header_values = [
            _normalize_text(worksheet.cell(row=row_index, column=column_index).value)
            for row_index in range(1, min(5, worksheet.max_row) + 1)
        ]
        header_text = " ".join(value for value in header_values if value)
        period = _normalize_period(header_text)
        if not period:
            period = next((fallback for label, fallback in fallback_by_label.items() if label in header_text and fallback), "")
        if period:
            period_columns[column_index] = period

    for column_index, period in fallback_period_columns.items():
        period_columns.setdefault(column_index, period)
    return period_columns


def _period_metric_amount(
    period_lookup: dict[tuple[str, str], float],
    metric_key: str,
    column_index: int,
    period_columns: dict[int, str],
) -> float | None:
    period = period_columns.get(column_index)
    if period:
        for candidate_period in _compatible_periods(period):
            for candidate_key in _period_metric_candidate_keys(metric_key):
                amount = period_lookup.get((candidate_key, candidate_period))
                if amount is not None:
                    return amount
        return None

    if column_index not in {3, 4}:
        return None
    candidate_keys = set(_period_metric_candidate_keys(metric_key))
    periods_for_key = sorted(
        [candidate_period for key, candidate_period in period_lookup if key in candidate_keys],
        reverse=True,
    )
    if not periods_for_key:
        return None
    if column_index == 3:
        period = periods_for_key[0]
    elif len(periods_for_key) >= 2:
        period = periods_for_key[1]
    else:
        return None
    for candidate_key in _period_metric_candidate_keys(metric_key):
        amount = period_lookup.get((candidate_key, period))
        if amount is not None:
            return amount
    return None


def _period_metric_candidate_keys(metric_key: str) -> list[str]:
    key = _normalize_text(metric_key)
    if not key:
        return []
    candidates = [key]
    stripped = key
    for suffix in ("本期数", "上期数", "本年数", "上年数", "本期", "上期", "本年", "上年"):
        if stripped.endswith(suffix):
            stripped = stripped[: -len(suffix)]
            break
    if stripped and stripped not in candidates:
        candidates.append(stripped)
    return candidates


def _compatible_periods(period: str) -> list[str]:
    normalized_period = _normalize_period(period)
    if not normalized_period:
        return []
    periods = [normalized_period]
    if len(normalized_period) >= 4:
        periods.append(normalized_period[:4])
    return list(dict.fromkeys(periods))


def _write_period_header(worksheet: Worksheet, column_index: int, period: str) -> None:
    if len(period) != 8:
        return
    header_cell = worksheet.cell(row=3, column=column_index)
    header_text = _normalize_text(header_cell.value)
    if not _looks_like_period_placeholder(header_text):
        return
    try:
        header_cell.value = pd.Timestamp(period).to_pydatetime()
    except Exception:
        header_cell.value = period


def _looks_like_period_placeholder(value: str) -> bool:
    return value in {PERIOD_COLUMN, "本期", "上期", "本年", "上年", "本年数", "上年数"}


def _ensure_report_template_merges(worksheet: Worksheet) -> None:
    if worksheet.merged_cells.ranges:
        return
    _merge_range_if_possible(worksheet, f"A1:{get_column_letter(worksheet.max_column)}1")
    if worksheet.max_column >= 4:
        _merge_range_if_possible(worksheet, "B2:C2")
        _merge_range_if_possible(worksheet, f"D2:{get_column_letter(worksheet.max_column)}2")
    elif worksheet.max_column >= 3:
        _merge_range_if_possible(worksheet, "C2:D2")


def _merge_range_if_possible(worksheet: Worksheet, range_string: str) -> None:
    if range_string in {str(merged_range) for merged_range in worksheet.merged_cells.ranges}:
        return
    try:
        worksheet.merge_cells(range_string)
    except ValueError:
        return


def _normalize_period(value: Any) -> str:
    text = _normalize_text(value)
    if text.endswith(".0"):
        text = text[:-2]
    digits = "".join(ch for ch in text if ch.isdigit())
    return digits[:8] if len(digits) >= 8 else digits


def _find_amount_cells(worksheet: Worksheet, row_index: int, item_column: int) -> tuple[Any, Any]:
    candidate_cells = [
        worksheet.cell(row=row_index, column=column_index)
        for column_index in range(item_column + 1, worksheet.max_column + 1)
    ]
    group_cell = _find_cell_by_header(worksheet, candidate_cells, ["集团", "合并"])
    parent_cell = _find_cell_by_header(worksheet, candidate_cells, ["本行", "母公司", "母行"])

    writable_cells = [
        cell for cell in candidate_cells if cell.value is None or isinstance(cell.value, (int, float))
    ]
    if group_cell is None and writable_cells:
        group_cell = writable_cells[0]
    if parent_cell is None and len(writable_cells) >= 2:
        parent_cell = writable_cells[1]
    return group_cell, parent_cell


def _find_cell_by_header(worksheet: Worksheet, candidate_cells: list[Any], keywords: list[str]) -> Any | None:
    for cell in candidate_cells:
        header_text = _column_header_text(worksheet, cell.column, cell.row)
        if any(keyword in header_text for keyword in keywords):
            return cell
    return None


def _column_header_text(worksheet: Worksheet, column_index: int, before_row: int) -> str:
    values: list[str] = []
    for row_index in range(max(1, before_row - 8), before_row):
        value = _normalize_text(worksheet.cell(row=row_index, column=column_index).value)
        if value:
            values.append(value)
    return " ".join(values)


def _is_item_label_column(worksheet: Worksheet, column_index: int, before_row: int) -> bool:
    header_text = _column_header_text(worksheet, column_index, before_row + 1)
    return any(keyword in header_text for keyword in ["指标名称", "项目名称", "项目", "名称"])


def _contextual_template_alias_key(
    worksheet: Worksheet,
    cell: Any,
    report_rows: dict[str, pd.Series],
) -> str:
    cell_text = _normalize_text(cell.value)
    if not METRIC_CODE_PATTERN.search(cell_text):
        return ""
    dimension_suffix = _metric_dimension_suffix(cell_text)
    if not dimension_suffix:
        return ""

    for column_index in range(1, cell.column):
        if not _is_item_label_column(worksheet, column_index, cell.row):
            continue
        row_item_name = _normalize_text(worksheet.cell(row=cell.row, column=column_index).value)
        report_row = report_rows.get(row_item_name)
        if report_row is None:
            continue
        item_code = _normalize_text(report_row.get(ITEM_CODE_COLUMN))
        if item_code:
            return f"{item_code}{dimension_suffix}"
    return ""


def _metric_dimension_suffix(value: str) -> str:
    text = _normalize_text(value)
    entity = "本行" if "本行" in text else ""
    if not entity and any(alias in text for alias in ["集团", "本集团", "合并"]):
        entity = "集团"
    period_alias = "上期" if "上期" in text or "上年" in text else ""
    if not period_alias and ("本期" in text or "本年" in text):
        period_alias = "本期"
    if not entity or not period_alias:
        return ""
    return f"{entity}{period_alias}"


def _should_default_template_metric_to_zero(value: str) -> bool:
    text = _normalize_text(value)
    return bool(METRIC_CODE_PATTERN.search(text) and _metric_dimension_suffix(text))


def _get_amount(row: pd.Series, column: str):
    if column not in row.index:
        return None
    raw_value = row[column]
    if pd.isna(raw_value):
        return None
    value = pd.to_numeric(raw_value, errors="coerce")
    if pd.notna(value):
        return float(value)
    text = _normalize_text(raw_value)
    return text or None


def _replace_sheet_with_dataframe(workbook: Any, sheet_name: str, df: pd.DataFrame) -> None:
    if sheet_name in workbook.sheetnames:
        del workbook[sheet_name]
    worksheet = workbook.create_sheet(sheet_name)

    for column_index, column_name in enumerate(df.columns, start=1):
        cell = worksheet.cell(row=1, column=column_index, value=column_name)
        cell.font = Font(bold=True)

    for row_index, row in enumerate(df.itertuples(index=False), start=2):
        row_metric_key = _dataframe_row_metric_key(df, row)
        for column_index, value in enumerate(row, start=1):
            if pd.isna(value):
                value = None
            cell = worksheet.cell(row=row_index, column=column_index, value=value)
            if "金额" in str(df.columns[column_index - 1]):
                cell.number_format = _metric_number_format(row_metric_key)
                cell.alignment = AMOUNT_ALIGNMENT

    worksheet.freeze_panes = "A2"
    if worksheet.max_row >= 1 and worksheet.max_column >= 1:
        worksheet.auto_filter.ref = worksheet.dimensions
    _auto_fit_columns(worksheet)


def _dataframe_row_metric_key(df: pd.DataFrame, row: Any) -> str:
    values = list(row)
    parts: list[str] = []
    for column_name in (ITEM_CODE_COLUMN, ITEM_COLUMN):
        if column_name not in df.columns:
            continue
        column_index = list(df.columns).index(column_name)
        if column_index < len(values):
            parts.append(_normalize_text(values[column_index]))
    return " ".join(part for part in parts if part)


def _build_summary_df(
    report_df: pd.DataFrame,
    report_config: dict[str, Any],
    template_path: Path,
    matches_df: pd.DataFrame,
) -> pd.DataFrame:
    display_name = str(report_config.get("display_name", "报表"))
    return pd.DataFrame(
        [
            {"项目": "报表类型", "内容": display_name},
            {"项目": "表样文件", "内容": template_path.name},
            {"项目": "生成数据行数", "内容": len(report_df)},
            {"项目": "表样匹配行数", "内容": len(matches_df)},
            {
                "项目": "生成说明",
                "内容": (
                    "系统按表样中的指标名称、指标编码或简单指标公式匹配生成数据；"
                    "完整生成数据保留在“生成数据”sheet。"
                ),
            },
        ]
    )


def _auto_fit_columns(worksheet: Worksheet) -> None:
    for column_cells in worksheet.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            if cell.value is None:
                continue
            max_length = max(max_length, _display_width(str(cell.value)))
        worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 10), 60)


def _display_width(value: str) -> int:
    width = 0
    for char in value:
        width += 2 if ord(char) > 127 else 1
    return width


def _normalize_text(value: Any) -> str:
    if value is None or pd.isna(value):
        return ""
    return " ".join(str(value).replace("\r", " ").replace("\n", " ").split())


def _safe_eval_numeric_expression(expression: str) -> float | None:
    operators = {
        ast.Add: operator.add,
        ast.Sub: operator.sub,
        ast.Mult: operator.mul,
        ast.Div: operator.truediv,
        ast.USub: operator.neg,
        ast.UAdd: operator.pos,
    }

    def eval_node(node: ast.AST) -> float:
        if isinstance(node, ast.Expression):
            return eval_node(node.body)
        if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)):
            return float(node.value)
        if isinstance(node, ast.BinOp) and type(node.op) in operators:
            return operators[type(node.op)](eval_node(node.left), eval_node(node.right))
        if isinstance(node, ast.UnaryOp) and type(node.op) in operators:
            return operators[type(node.op)](eval_node(node.operand))
        raise ValueError("unsupported expression")

    try:
        value = eval_node(ast.parse(expression, mode="eval"))
    except Exception:
        return None
    if pd.isna(value):
        return None
    return float(value)
